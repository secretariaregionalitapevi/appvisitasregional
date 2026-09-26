"""Leitura, conciliação e importação idempotente da agenda de visitas."""

from __future__ import annotations

import base64
import difflib
import hashlib
import io
import json
import re
import unicodedata
from datetime import date, datetime, time, timedelta
from pathlib import Path

import requests
from django.conf import settings
from openpyxl import load_workbook


MAX_IMPORT_BYTES = 15 * 1024 * 1024
ALLOWED_EXTENSIONS = {".xlsx", ".xlsm", ".pdf", ".png", ".jpg", ".jpeg", ".webp"}


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return " ".join(re.sub(r"[^A-Za-z0-9]+", " ", text).upper().split())


def normalize_address(value):
    # Endereços do cadastro podem começar com "[latitude, longitude]".
    # Coordenadas não fazem parte da identidade da casa e o primeiro "23"
    # não pode ser confundido com o número do imóvel.
    value = re.sub(r"^\s*\[\s*-?\d+(?:\.\d+)?\s*,\s*-?\d+(?:\.\d+)?\s*\]\s*", "", str(value or ""))
    text = normalize(value)
    replacements = {
        "RUA": "R", "AVENIDA": "AV", "NUMERO": "N", "NO": "N",
        "DOUTOR": "DR", "JARDIM": "JD", "VILA": "VL",
    }
    return " ".join(replacements.get(part, part) for part in text.split())


def host_name(value):
    text = " ".join(str(value or "").split())
    text = re.sub(r"^REUNI[AÃ]O\s+(?:DE\s+)?EVANGELIZA(?:ÇÃO|CAO)\s*", "", text, flags=re.I)
    text = re.sub(r"^IRM(?:Ã|A)O?\s+", "", text, flags=re.I)
    return text.strip().title()


def infer_category(value):
    value = normalize(value)
    return "RE" if "EVANGEL" in value or value == "RE" else "RF"


def parse_date(value, year_hint=None):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    match = re.fullmatch(r"(\d{1,2})/(\d{1,2})", text)
    if match and year_hint:
        return date(int(year_hint), int(match.group(2)), int(match.group(1)))
    raise ValueError(f"Data inválida: {text or 'vazia'}")


def parse_time(value):
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, (int, float)) and 0 <= value < 1:
        seconds = round(value * 86400)
        return time((seconds // 3600) % 24, (seconds % 3600) // 60)
    text = str(value or "").strip().replace("h", ":")
    for fmt in ("%H:%M", "%H:%M:%S", "%H"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            pass
    raise ValueError(f"Horário inválido: {text or 'vazio'}")


def _header_key(value):
    value = normalize(value)
    aliases = {
        "DATA": "date", "DIA": "date", "HRS": "time", "HORA": "time", "HORARIO": "time",
        "NA CASA DA IRMA O": "host", "NA CASA DO IRMA O": "host", "NA CASA DA IRMA": "host",
        "NOME": "host", "IRMA O": "host", "IRMA": "host", "VISITADO": "host",
        "ENDERECO REFERENCIA": "address", "ENDERECO": "address", "REFERENCIA": "address",
        "ATENDE": "attendant", "RESPONSAVEL": "attendant", "ATENDENTE": "attendant",
        "TIPO": "category", "CATEGORIA": "category", "REUNIAO": "category",
    }
    if value in aliases:
        return aliases[value]
    if "CASA" in value and "IRMA" in value:
        return "host"
    if "ENDERECO" in value:
        return "address"
    return None


def parse_workbook(content):
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    rows = []
    for sheet in workbook.worksheets:
        header = None
        year_hint = None
        for cells in sheet.iter_rows(values_only=True):
            values = list(cells)
            joined = " ".join(str(value or "") for value in values)
            year_match = re.search(r"\b(20\d{2})\b", joined)
            if year_match:
                year_hint = int(year_match.group(1))
            mapped = {key: index for index, value in enumerate(values) if (key := _header_key(value))}
            if "date" in mapped and "host" in mapped:
                header = mapped
                continue
            if not header:
                continue
            try:
                item_date = parse_date(values[header["date"]], year_hint)
            except (ValueError, IndexError):
                continue
            try:
                item_time = parse_time(values[header.get("time", -1)])
            except (ValueError, IndexError):
                item_time = time(20, 0)
            def cell(key, default=""):
                position = header.get(key)
                return values[position] if position is not None and position < len(values) else default

            host_raw = cell("host")
            category_raw = cell("category", host_raw)
            rows.append({
                "date": item_date.isoformat(), "time": item_time.strftime("%H:%M"),
                "host": host_name(host_raw), "address": str(cell("address") or "").strip(),
                "category": infer_category(category_raw or host_raw),
                "attendant": " ".join(str(cell("attendant") or "").split()).title(),
            })
    return rows


DOCUMENT_SCHEMA = {
    "type": "object", "additionalProperties": False,
    "properties": {
        "year": {"type": ["integer", "null"]},
        "events": {"type": "array", "items": {
            "type": "object", "additionalProperties": False,
            "properties": {
                "date": {"type": "string"}, "time": {"type": "string"},
                "host": {"type": "string"}, "address": {"type": "string"},
                "category": {"type": "string", "enum": ["RF", "RE"]},
                "attendant": {"type": "string"},
            },
            "required": ["date", "time", "host", "address", "category", "attendant"],
        }},
    },
    "required": ["year", "events"],
}


def parse_visual_document(content, filename, content_type):
    api_key = getattr(settings, "OPENAI_API_KEY", "")
    if not api_key:
        raise RuntimeError("A leitura de imagem/PDF requer OPENAI_API_KEY configurada no servidor.")
    encoded = base64.b64encode(content).decode("ascii")
    is_pdf = Path(filename).suffix.lower() == ".pdf"
    file_input = ({"type": "input_file", "filename": filename,
                   "file_data": f"data:application/pdf;base64,{encoded}"} if is_pdf else
                  {"type": "input_image", "image_url": f"data:{content_type};base64,{encoded}"})
    prompt = (
        "Extraia todas as linhas da agenda de reuniões. Preserve data, horário, nome da pessoa visitada, "
        "endereço completo e responsável pelo atendimento. Use RF para reunião familiar e RE para "
        "reunião de evangelização. Quando o ano aparecer no cabeçalho, aplique-o a todas as datas. "
        "Não invente linhas nem complete nomes ilegíveis. Datas devem sair em AAAA-MM-DD e horas em HH:MM."
    )
    response = requests.post(
        "https://api.openai.com/v1/responses",
        headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
        json={
            "model": getattr(settings, "OPENAI_DOCUMENT_MODEL", "gpt-4.1-mini"),
            "input": [{"role": "user", "content": [{"type": "input_text", "text": prompt}, file_input]}],
            "text": {"format": {"type": "json_schema", "name": "agenda_visitas", "strict": True,
                                "schema": DOCUMENT_SCHEMA}},
        }, timeout=120,
    )
    response.raise_for_status()
    payload = response.json()
    output_text = payload.get("output_text")
    if not output_text:
        output_text = next((part.get("text") for item in payload.get("output", [])
                            for part in item.get("content", []) if part.get("type") == "output_text"), None)
    if not output_text:
        raise ValueError("O documento não retornou dados legíveis.")
    return json.loads(output_text).get("events", [])


def parse_document(upload):
    filename = Path(upload.name or "documento").name
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Formato não aceito. Envie Excel (.xlsx), PDF, PNG, JPG ou WEBP.")
    content = upload.read(MAX_IMPORT_BYTES + 1)
    if not content:
        raise ValueError("O arquivo enviado está vazio.")
    if len(content) > MAX_IMPORT_BYTES:
        raise ValueError("O arquivo excede o limite de 15 MB.")
    rows = parse_workbook(content) if extension in {".xlsx", ".xlsm"} else parse_visual_document(
        content, filename, upload.content_type or "image/jpeg")
    if not rows:
        raise ValueError("Nenhuma reunião foi identificada no documento.")
    return rows, hashlib.sha256(content).hexdigest(), filename


def _number(value):
    match = re.search(r"\b(\d+[A-Z]?)\b", normalize_address(value))
    return match.group(1) if match else ""


def match_member(row, members):
    wanted_name = normalize(row.get("host"))
    wanted_address = normalize_address(row.get("address"))
    wanted_number = _number(row.get("address"))
    exact = [member for member in members if normalize(member.get("nome")) == wanted_name]
    if len(exact) == 1:
        return exact[0], "exact"
    if len(exact) > 1:
        by_address = [member for member in exact if wanted_number and _number(member.get("endereco")) == wanted_number]
        return (by_address[0], "name_address") if len(by_address) == 1 else (None, "ambiguous")
    first_name = wanted_name.split()[0] if wanted_name else ""
    same_first_name = [member for member in members
                       if normalize(member.get("nome")).split()[:1] == [first_name]]
    if len(same_first_name) == 1:
        return same_first_name[0], "first_name"
    by_number = [member for member in same_first_name
                 if wanted_number and _number(member.get("endereco")) == wanted_number]
    if len(by_number) == 1:
        return by_number[0], "name_address"
    scored = []
    for member in members:
        member_name = normalize(member.get("nome"))
        name_score = difflib.SequenceMatcher(None, wanted_name, member_name).ratio()
        address_score = difflib.SequenceMatcher(None, wanted_address, normalize_address(member.get("endereco"))).ratio()
        number_bonus = .18 if wanted_number and wanted_number == _number(member.get("endereco")) else 0
        scored.append((name_score * .72 + address_score * .28 + number_bonus, member))
    scored.sort(key=lambda item: item[0], reverse=True)
    if scored and scored[0][0] >= .78 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= .08):
        return scored[0][1], "fuzzy"
    return None, "not_found"


def import_key(common, row, member_id):
    identity = member_id or normalize_address(row.get("address")) or normalize(row.get("host"))
    raw = "|".join([normalize(common), str(row.get("date")), infer_category(row.get("category")), str(identity)])
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def sector_from_address(address, fallback="Não Definido"):
    normalized = normalize(address)
    match = re.search(r"\b(?:VILA|JARDIM|PARQUE|BAIRRO|CIDADE)\s+[A-Z0-9 ]+", normalized)
    return match.group(0).title() if match else fallback


def prepare_rows(raw_rows, common, members, teams, existing):
    existing_keys = {str(item.get("import_key") or "") for item in existing if item.get("import_key")}
    prepared = []
    seen = set()
    for index, raw in enumerate(raw_rows, start=1):
        try:
            row_date = parse_date(raw.get("date")).isoformat()
            row_time = parse_time(raw.get("time") or "20:00").strftime("%H:%M")
        except ValueError as exc:
            prepared.append({**raw, "line": index, "state": "error", "message": str(exc)})
            continue
        row = {**raw, "line": index, "date": row_date, "time": row_time,
               "host": host_name(raw.get("host")), "category": infer_category(raw.get("category"))}
        member, match_kind = match_member(row, members)
        if not member and match_kind != "ambiguous" and row["host"] and len(normalize_address(row.get("address"))) >= 10 and _number(row.get("address")):
            import uuid
            member = {"id": str(uuid.uuid5(uuid.NAMESPACE_URL, "agenda-member|" + normalize(common) + "|" + normalize(row["host"]) + "|" + normalize_address(row["address"]))),
                      "nome": row["host"], "endereco": row["address"]}
            row["create_member"] = True
        if not member:
            message = "Nome corresponde a mais de um cadastro." if match_kind == "ambiguous" else "Pessoa não localizada com segurança no cadastro."
            prepared.append({**row, "state": "error", "message": message})
            continue
        key = import_key(common, row, member["id"])
        duplicate = key in existing_keys or key in seen or any(
            str(item.get("status")) != "Cancelada"
            and str(item.get("data_inicio") or "")[:10] == row_date
            and infer_category(item.get("categoria")) == row["category"]
            and (str(item.get("irmandade_id") or "") == str(member["id"])
                 or (normalize_address(item.get("endereco_visitado")) and
                     normalize_address(item.get("endereco_visitado")) == normalize_address(row.get("address"))))
            for item in existing
        )
        if duplicate:
            seen.add(key)
            prepared.append({
                **row, "member_id": member["id"], "matched_name": member.get("nome"),
                "matched_address": member.get("endereco"), "import_key": key,
                "state": "duplicate", "message": "Reunião já existente; não será duplicada.",
                "match_kind": match_kind,
            })
            continue
        seen.add(key)
        prepared.append({
            **row, "member_id": member["id"], "matched_name": member.get("nome"),
            "matched_address": member.get("endereco"), "sector": member.get("setor") or sector_from_address(row.get("address")),
            "import_key": key, "state": "ready",
            "message": ("Novo cadastro de irmandade será criado ao confirmar." if row.get("create_member") else
                        ("Conferir conciliação aproximada." if match_kind in {"fuzzy", "first_name"} else "Pronto para importar.")),
            "match_kind": match_kind,
        })
    return prepared


def row_payload(row, source_name, source_hash):
    start = datetime.fromisoformat(f"{row['date']}T{row['time']}")
    notes = f"[[responsavel_atendimento:{row.get('attendant', '')}]]"
    return {
        "irmandade_id": row["member_id"], "titulo": row["matched_name"],
        "data_inicio": start.strftime("%Y-%m-%dT%H:%M:00-03:00"),
        "data_fim": (start + timedelta(hours=1)).strftime("%Y-%m-%dT%H:%M:00-03:00"),
        "equipe_responsavel": None, "equipe_id": None,
        "equipe_tipo": None, "categoria": row["category"], "status": "Marcada",
        "setor": row["sector"], "endereco_visitado": row.get("address") or row.get("matched_address"),
        "observacoes": notes, "import_key": row["import_key"],
    }
