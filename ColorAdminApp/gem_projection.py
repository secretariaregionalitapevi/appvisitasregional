"""Inteligência de formação e projeção da futura orquestra do GEM."""
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta
import io
import math
import re

import requests
from django.conf import settings
from django.core.cache import cache
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .access_control import common_catalog, scope_details, service_headers, user_scope
from .gem import _fetch_last_activity_dates, _norm, _visible_students, can_open_module, is_graduated, operational_activity_from_last_activity
from .sam_program import assess_program, target_for_level
from .sam_history_sync import SOURCE_CONFIG

ATTENDANCE_TARGET, ABSENCE_ALERT, LESSONS_PER_SEMESTER, COURSE_LESSONS = 85, 3, 15, 60


def _actor(request):
    profile = request.session.get("user_profile") or {}
    return profile.get("full_name") or profile.get("nome") or profile.get("name") or profile.get("email") or "Usuário"


def page(request):
    if not can_open_module(request):
        return render(request, "pages/403.html", {"message": "Seu perfil não possui acesso à pasta GEM."}, status=403)
    return render(request, "pages/gem_projection.html", {"scope": scope_details(user_scope(request)), "report_user": _actor(request)})


def _fetch_all(table, select, order=None):
    key = f"gem:projection:{table}:v1"
    cached = cache.get(key)
    if cached is not None:
        return cached
    size = 1000
    params = {"select": select, "offset": 0, "limit": size}
    if order: params["order"] = order
    response = requests.get(f"{settings.SUPABASE_URL}/rest/v1/{table}", headers=service_headers("count=exact"), params=params, timeout=30)
    response.raise_for_status(); first = response.json()
    try: total = int(response.headers.get("Content-Range", "0/0").rsplit("/", 1)[-1])
    except ValueError: total = len(first)
    def fetch(offset):
        query = {"select": select, "offset": offset, "limit": size}
        if order: query["order"] = order
        current = requests.get(f"{settings.SUPABASE_URL}/rest/v1/{table}", headers=service_headers(), params=query, timeout=30)
        current.raise_for_status(); return offset, current.json()
    offsets = list(range(size, total, size))
    with ThreadPoolExecutor(max_workers=min(8, len(offsets) or 1)) as pool:
        pages = sorted(pool.map(fetch, offsets), key=lambda item: item[0]) if offsets else []
    rows = first + [row for _, values in pages for row in values]
    cache.set(key, rows, 300); return rows


def _as_date(value):
    try: return datetime.fromisoformat(str(value or "")[:10]).date()
    except ValueError: return None


def _projection_date(total_calls, dates):
    remaining = max(0, COURSE_LESSONS - total_calls)
    if not remaining: return None
    valid = sorted(value for value in dates if value)
    weeks = max(1, (valid[-1] - valid[0]).days / 7) if len(valid) > 1 else 1
    observed = min(2.0, max(0.5, total_calls / weeks))
    return date.today() + timedelta(weeks=math.ceil(remaining / observed))


def _target_label(level):
    return {"rjm": "RJM", "culto": "Culto oficial", "oficializacao": "Oficialização", "concluido": "Formação concluída"}.get(target_for_level(level), "RJM")


def _common_name_key(value):
    return re.sub(r"^BR-\d+(?:-\d+)*\s*-\s*", "", _norm(value)).strip()


def _official_common_map():
    """Relaciona nomes legados ao rótulo oficial com código BR, sem arriscar homônimos."""
    result = {}
    for row in common_catalog():
        label = str(row.get("comum") or "").strip()
        key = _common_name_key(label)
        if key and label:
            result.setdefault(key, []).append(label)
    return {key: labels[0] for key, labels in result.items() if len(labels) == 1}


def _classification(progress, frequency, absences, level):
    ready = progress >= 100 and frequency is not None and frequency >= ATTENDANCE_TARGET and absences <= ABSENCE_ALERT
    inconsistent = (progress >= 100 and not ready) or ("CULTO OFICIAL" in _norm(level) and progress == 0)
    status = "APTO PARA AVALIAÇÃO" if ready else "ACOMPANHAMENTO PRIORITÁRIO" if absences > ABSENCE_ALERT or (frequency is not None and frequency < ATTENDANCE_TARGET) else "CONFERIR CADASTRO" if inconsistent else "EM EVOLUÇÃO"
    return ready, inconsistent, status


def _stage_key(row):
    target = target_for_level(row.get("nivel"))
    return {"rjm": "entrada_rjm", "culto": "rjm_culto", "oficializacao": "culto_oficializacao"}.get(target, "concluido")


def _stage_label(key):
    return {"entrada_rjm": "Candidato(a) → RJM", "rjm_culto": "RJM → Culto oficial", "culto_oficializacao": "Culto oficial → Oficialização", "concluido": "Formação concluída"}.get(key, key)


def _stage_summary(rows):
    summaries = []
    for key in ("entrada_rjm", "rjm_culto", "culto_oficializacao"):
        items = [row for row in rows if row.get("etapa_key") == key]
        bands = {"0_24": 0, "25_49": 0, "50_74": 0, "75_99": 0, "100": 0}
        for row in items:
            value = row.get("programa") or 0
            band = "100" if value >= 100 else "75_99" if value >= 75 else "50_74" if value >= 50 else "25_49" if value >= 25 else "0_24"
            bands[band] += 1
        summaries.append({"key": key, "label": _stage_label(key), "total": len(items), "media": round(sum(row.get("programa") or 0 for row in items) / len(items)) if items else 0, "aptos": sum(row.get("apto") is True for row in items), "proximos": sum(75 <= (row.get("programa") or 0) < 100 for row in items), "bands": bands})
    return summaries


def _enrich_program_assessment(rows):
    """Aplica o parecer M09 real somente ao recorte devolvido ou exportado."""
    ids = [str(row["id"]) for row in rows if row.get("id")]
    if not ids:
        return rows
    records = defaultdict(lambda: {"msa": [], "metodo": [], "hinario": []})
    tasks = []
    for source in ("msa", "metodo", "hinario"):
        table, fields = SOURCE_CONFIG[source]
        for start in range(0, len(ids), 100):
            tasks.append((source, table, fields, ids[start:start + 100]))
    def fetch(task):
        source, table, fields, chunk = task
        response = requests.get(f"{settings.SUPABASE_URL}/rest/v1/{table}", headers=service_headers(), params={"select": "aluno_id," + ",".join(fields), "aluno_id": "in.(" + ",".join(chunk) + ")", "limit": 10000}, timeout=30)
        response.raise_for_status()
        return source, response.json()
    with ThreadPoolExecutor(max_workers=min(8, len(tasks))) as pool:
        for source, values in pool.map(fetch, tasks):
            for value in values:
                records[str(value.get("aluno_id"))][source].append(value)
    activity_dates = _fetch_last_activity_dates(ids)
    for row in rows:
        assessment = assess_program(row, records[str(row["id"])])
        row["m09_percentual"] = assessment["completion_percent"]
        row["m09_areas_atendidas"] = sum(item["status"] == "ok" for item in assessment["requirements"])
        row["m09_areas_total"] = len(assessment["requirements"])
        row["m09_elegivel"] = assessment["eligible"]
        row["ultimo_lancamento"] = activity_dates.get(str(row["id"]))
        row["status_historico"] = operational_activity_from_last_activity(row["ultimo_lancamento"])["operational_status"]
        dates = [_as_date(item.get("data_aula")) for item in records[str(row["id"])]["msa"]]
        inaugural = min((value for value in dates if value), default=None)
        row["aula_inaugural"] = inaugural.isoformat() if inaugural else None
        row["plano_percentual"] = min(100, round((row.get("aulas") or 0) * 100 / COURSE_LESSONS))
        if not assessment["eligible"] and row.get("apto"):
            row["apto"] = False; row["inconsistencia"] = True; row["status"] = "CONFERIR PROGRAMA M09"
        frequency = row.get("frequencia")
        if inaugural and frequency:
            duration_days = round(730 * max(1, ATTENDANCE_TARGET / frequency))
            estimated = inaugural + timedelta(days=duration_days)
            if estimated < date.today() and row["plano_percentual"] < 100:
                estimated = date.today() + timedelta(days=round((COURSE_LESSONS - (row.get("aulas") or 0)) * 730 / COURSE_LESSONS))
            row["projecao_conclusao"] = estimated.isoformat()
            row["semestre"] = min(4, max(1, ((date.today() - inaugural).days // 183) + 1))
        elif not frequency:
            row["projecao_conclusao"] = None
    return rows


def _build_rows(request):
    _, students = _visible_students(request); students = [row for row in students if not is_graduated(row)]
    visible = {str(row["id"]): row for row in students}
    official_commons = _official_common_map()
    detailed = bool(request.GET.get("q") or request.GET.get("etapa") or request.GET.get("export") == "1")
    classes = []
    if detailed:
        class_response = requests.get(f"{settings.SUPABASE_URL}/rest/v1/sam_gem_classes", headers=service_headers(), params={"select": "id,data_aula", "order": "data_aula.desc", "limit": 1000}, timeout=30)
        class_response.raise_for_status(); classes = class_response.json()
    calls = []
    class_ids = [str(row["id"]) for row in classes]
    for start in range(0, len(class_ids), 100):
        ids = class_ids[start:start + 100]
        response = requests.get(f"{settings.SUPABASE_URL}/rest/v1/sam_gem_attendance", headers=service_headers(), params={"select": "aluno_id,aula_id,presente", "aula_id": f"in.({','.join(ids)})", "limit": 10000}, timeout=30)
        response.raise_for_status(); calls.extend(response.json())
    class_dates = {str(row.get("id")): _as_date(row.get("data_aula")) for row in classes}
    history = defaultdict(list)
    for call in calls:
        student_id = str(call.get("aluno_id") or "")
        if student_id in visible: history[student_id].append((call.get("presente") is True, class_dates.get(str(call.get("aula_id")))))
    leaders = {}
    try:
        if not detailed:
            raise LookupError
        for exam in _fetch_all("musica_exames_individuais", "comum,encarregado_local,data_exame", "data_exame.desc"):
            key = _norm(exam.get("comum"))
            if key and exam.get("encarregado_local") and key not in leaders: leaders[key] = exam["encarregado_local"]
    except (requests.RequestException, LookupError): pass
    result = []
    for student_id, student in visible.items():
        student = dict(student)
        source_common = student.get("comum_congregacao") or ""
        student["comum_congregacao"] = official_commons.get(_common_name_key(source_common), source_common)
        items = history.get(student_id, []); total = len(items); present = sum(value for value, _ in items); absences = total - present
        frequency = round(present * 100 / total) if total else None; progress = int(student.get("programa_minimo_percentual") or 0)
        ready, inconsistent, status = _classification(progress, frequency, absences, student.get("nivel"))
        assessment = {"completion_percent": progress}
        projected = _projection_date(total, [value for _, value in items])
        result.append({"id": student_id, "nome": student.get("nome_aluno"), "instrumento": student.get("instrumento") or "A definir", "comum": student.get("comum_congregacao") or "Não informada", "municipio": student.get("municipio") or "Não informado", "nivel": student.get("nivel") or "Não informado", "proxima_etapa": _target_label(student.get("nivel")), "programa": progress, "comprovacao_m09": assessment["completion_percent"], "aulas": total, "presencas": present, "faltas_sem_justificativa": absences, "frequencia": frequency, "semestre": min(4, total // LESSONS_PER_SEMESTER + 1) if total else 1, "projecao_conclusao": projected.isoformat() if projected else None, "status": status, "apto": ready, "inconsistencia": inconsistent, "encarregado_local": leaders.get(_norm(student.get("comum_congregacao"))) or "Não cadastrado"})
    for row in result:
        source = visible.get(str(row.get("id"))) or {}
        row["plano_percentual"] = min(100, round((row.get("aulas") or 0) * 100 / COURSE_LESSONS))
        row["m09_percentual"] = None
        row["m09_areas_atendidas"] = None
        row["m09_areas_total"] = 3
        row["m09_elegivel"] = None
        row["aula_inaugural"] = None
        row["status_historico"] = source.get("operational_status") or "SEM HISTORICO"
        row["ultimo_lancamento"] = source.get("last_activity_at")
        row["etapa_key"] = _stage_key(row)
        row["etapa_label"] = _stage_label(row["etapa_key"])
    return result


def _filtered(request, rows=None):
    rows = _build_rows(request) if rows is None else list(rows); query = _norm(request.GET.get("q")); status = _norm(request.GET.get("status")); city = _norm(request.GET.get("municipio")); common = _norm(request.GET.get("comum")); instrument = _norm(request.GET.get("instrumento")); stage = request.GET.get("etapa", "").strip()
    if query: rows = [row for row in rows if query in _norm(" ".join(str(row.get(k) or "") for k in ("nome", "instrumento", "comum")))]
    if status: rows = [row for row in rows if _norm(row["status"]) == status]
    if city: rows = [row for row in rows if _norm(row["municipio"]) == city]
    if common: rows = [row for row in rows if _norm(row["comum"]) == common]
    if instrument: rows = [row for row in rows if _norm(row["instrumento"]) == instrument]
    if stage: rows = [row for row in rows if row.get("etapa_key") == stage]
    return sorted(rows, key=lambda row: (not row["apto"], row["status"], _norm(row["nome"])))


def api_dashboard(request):
    if not can_open_module(request): return JsonResponse({"error": "Acesso negado."}, status=403)
    try:
        all_rows = _build_rows(request)
        original_query = request.GET
        summary_query = request.GET.copy(); summary_query.pop("etapa", None)
        request.GET = summary_query; summary_rows = _filtered(request, all_rows)
        request.GET = original_query; rows = _filtered(request, all_rows)
        totals = {"alunos": len(all_rows), "aptos": sum(row["apto"] for row in all_rows), "prioritarios": sum(row["status"] == "ACOMPANHAMENTO PRIORITÁRIO" for row in all_rows), "inconsistencias": sum(row["inconsistencia"] for row in all_rows), "sem_frequencia": sum(row["frequencia"] is None for row in all_rows)}
        detailed = bool(request.GET.get("q") or request.GET.get("etapa"))
        visible_rows = rows if request.GET.get("export") == "1" else rows[:80] if detailed else rows[:300]
        if detailed or request.GET.get("export") == "1":
            _enrich_program_assessment(visible_rows)
        return JsonResponse({"totals": totals, "stages": _stage_summary(summary_rows), "items": visible_rows, "result_count": len(rows), "municipalities": sorted({row["municipio"] for row in all_rows}), "commons": sorted({row["comum"] for row in all_rows}), "instruments": sorted({row["instrumento"] for row in all_rows}), "criteria": {"frequencia": ATTENDANCE_TARGET, "faltas": ABSENCE_ALERT, "aulas_semestre": LESSONS_PER_SEMESTER, "aulas_curso": COURSE_LESSONS}})
    except requests.RequestException: return JsonResponse({"error": "Não foi possível consultar a base sincronizada do SAM."}, status=502)


def export_excel(request):
    if not can_open_module(request): return HttpResponse(status=403)
    try: rows = _enrich_program_assessment(_filtered(request))
    except requests.RequestException: return HttpResponse("Não foi possível consultar a base sincronizada.", status=502)
    wb = Workbook(); ws = wb.active; ws.title = "Projeção"
    headers = ["ALUNO", "INSTRUMENTO", "COMUM", "MUNICÍPIO", "ENCARREGADO LOCAL", "NÍVEL ATUAL", "PRÓXIMA ETAPA", "PLANO DE AULAS", "COBERTURA MSA", "PROGRAMA M09", "ÁREAS M09", "FREQUÊNCIA", "FALTAS S/ JUST.", "SEMESTRE", "PROJEÇÃO", "STATUS HISTÓRICO", "ÚLTIMO LANÇAMENTO", "SITUAÇÃO"]; last = get_column_letter(len(headers))
    for index, (label, size) in enumerate((("CONGREGAÇÃO CRISTÃ NO BRASIL", 15), ("Regional Itapevi - São Paulo", 10), ("GRUPO DE ESTUDOS MUSICAIS", 12), ("Projeção da Orquestra", 11)), 1):
        ws.merge_cells(start_row=index, start_column=1, end_row=index, end_column=len(headers)); cell = ws.cell(index, 1, label); cell.font = Font(bold=index in {1, 3, 4}, size=size, color="FFFFFF" if index == 1 else "1E4B7A"); cell.fill = PatternFill("solid", fgColor="1E4B7A" if index == 1 else "EAF2F8"); cell.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(headers)); ws.cell(5, 1, f"Emissão: {datetime.now():%d/%m/%Y %H:%M} · Responsável: {_actor(request)} · {len(rows)} aluno(s)")
    for col, value in enumerate(headers, 1):
        cell = ws.cell(7, col, value); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1E4B7A"); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, row in enumerate(rows, 8):
        projection = _as_date(row["projecao_conclusao"]); last_activity = _as_date(row.get("ultimo_lancamento")); values = [row["nome"], row["instrumento"], row["comum"], row["municipio"], row["encarregado_local"], row["nivel"], row["proxima_etapa"], f'{row["plano_percentual"]}%', f'{row["programa"]}%', f'{row["m09_percentual"]}%', f'{row["m09_areas_atendidas"]} de {row["m09_areas_total"]}', "Sem dados" if row["frequencia"] is None else f'{row["frequencia"]}%', row["faltas_sem_justificativa"], row["semestre"], projection.strftime("%d/%m/%Y") if projection else "Sem base", row.get("status_historico"), last_activity.strftime("%d/%m/%Y") if last_activity else "Sem lançamento", row["status"]]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value); cell.fill = PatternFill("solid", fgColor="F4F6F8" if row_index % 2 == 0 else "FFFFFF"); cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col, width in enumerate([30,18,34,18,24,20,20,14,14,14,13,12,14,10,14,16,16,28], 1): ws.column_dimensions[get_column_letter(col)].width = width
    thin = Side(style="thin", color="D8DEE5")
    for sheet_row in ws.iter_rows(min_row=7):
        for cell in sheet_row: cell.border = Border(bottom=thin)
    ws.freeze_panes = "A8"; ws.auto_filter.ref = f"A7:{last}{max(7, 7 + len(rows))}"; ws.sheet_view.showGridLines = False; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    summary = wb.create_sheet("Resumo por etapa", 0)
    summary.merge_cells("A1:H1"); summary["A1"] = "CONGREGAÇÃO CRISTÃ NO BRASIL"; summary["A1"].font = Font(bold=True, size=15, color="FFFFFF"); summary["A1"].fill = PatternFill("solid", fgColor="1E4B7A"); summary["A1"].alignment = Alignment(horizontal="center")
    summary.merge_cells("A2:H2"); summary["A2"] = "Regional Itapevi - São Paulo"; summary["A2"].alignment = Alignment(horizontal="center")
    summary.merge_cells("A3:H3"); summary["A3"] = "GRUPO DE ESTUDOS MUSICAIS · FUNIL DAS PRÓXIMAS MUDANÇAS"; summary["A3"].font = Font(bold=True, color="1E4B7A"); summary["A3"].alignment = Alignment(horizontal="center")
    summary.merge_cells("A4:H4"); summary["A4"] = f"Emissão: {datetime.now():%d/%m/%Y %H:%M} · Responsável: {_actor(request)} · Recorte: filtros selecionados"
    summary_headers = ["PRÓXIMA MUDANÇA", "TOTAL", "MÉDIA MSA", "0–24%", "25–49%", "50–74%", "75–99%", "100% / APTOS M09"]
    for col, value in enumerate(summary_headers, 1):
        cell = summary.cell(6, col, value); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1E4B7A"); cell.alignment = Alignment(horizontal="center")
    for row_index, item in enumerate(_stage_summary(rows), 7):
        values = [item["label"], item["total"], f'{item["media"]}%', item["bands"]["0_24"], item["bands"]["25_49"], item["bands"]["50_74"], item["bands"]["75_99"], f'{item["bands"]["100"]} / {item["aptos"]}']
        for col, value in enumerate(values, 1):
            cell = summary.cell(row_index, col, value); cell.alignment = Alignment(horizontal="left" if col == 1 else "center"); cell.fill = PatternFill("solid", fgColor="F4F6F8" if row_index % 2 == 0 else "FFFFFF")
    summary.column_dimensions["A"].width = 36
    for col in "BCDEFGH": summary.column_dimensions[col].width = 15
    summary.freeze_panes = "A7"; summary.auto_filter.ref = "A6:H9"; summary.sheet_view.showGridLines = False; summary.page_setup.orientation = "landscape"; summary.page_setup.fitToWidth = 1; summary.sheet_properties.pageSetUpPr.fitToPage = True; summary.print_title_rows = "1:6"
    output = io.BytesIO(); wb.save(output); output.seek(0); response = HttpResponse(output.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); response["Content-Disposition"] = f'attachment; filename="Projecao_Orquestra_GEM_{datetime.now():%d-%m-%Y_%H-%M}.xlsx"'; return response
