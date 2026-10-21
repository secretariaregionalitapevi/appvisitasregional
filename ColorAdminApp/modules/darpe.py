"""Cadastro e eventos do D.A.R.P.E com autorização e escopo territorial no servidor."""
import io
import json
import logging
import re
import unicodedata
from datetime import date, datetime
from urllib.parse import quote
from uuid import uuid4

import requests
from django.conf import settings
from django.http import JsonResponse
from django.http import HttpResponse
from django.shortcuts import render
from django.views.decorators.cache import never_cache
from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .access_control import _norm, can_access, common_catalog, filter_rows, scope_details, service_headers, user_scope
from .module_access import MODULE_DARPE, can_access_module
from .music_events_identity import CommonResolver
from .musicalizacao import REGIONAL_MUNICIPALITIES


RESOURCES = {
    "cadastros": {
        "table": "darpe_cadastro", "order": "nome_completo.asc",
        "fields": {
            "cpf", "nome_completo", "instrumento_oficializado", "rg", "data_nascimento", "sexo",
            "estado_civil", "comum_congregacao", "data_batismo", "logradouro", "numero", "complemento",
            "bairro", "cidade", "uf", "cep", "ddd", "telefone", "darpe_ministerio", "darpe_atendente",
            "darpe_colaborador", "darpe_encarregado", "darpe_musico", "darpe_organista",
            "darpe_musico_organista", "darpe_grupo_canto",
            "declaracao_local", "declaracao_data", "nome_ministerio", "nome_encarregado_orquestra",
            "nome_voluntario", "status",
        },
        "location": lambda row: {"comum": row.get("comum_congregacao"), "cidade": row.get("cidade")},
    },
    "eventos": {
        "table": "darpe_eventos", "order": "data_evento.desc",
        "fields": {"titulo", "tipo", "data_evento", "hora_inicio", "hora_fim", "comum", "municipio",
                   "local_evento", "endereco", "status", "observacoes"},
        "location": lambda row: {"comum": row.get("comum"), "municipio": row.get("municipio")},
    },
    "participantes": {
        "table": "darpe_evento_participantes", "order": "created_at.desc",
        "fields": {"evento_id", "cadastro_id", "funcao", "situacao", "observacoes"},
        "location": lambda row: {},
    },
}

logger = logging.getLogger(__name__)


def _write_cadastro(method, *, reference=None, **kwargs):
    """Compatibilidade entre as colunas musicais antigas e as separadas."""
    data = kwargs['json']
    def adapt(payload):
        payload = dict(payload)
        if reference is not None:
            for field in ('darpe_musico', 'darpe_organista', 'darpe_musico_organista'):
                if field not in reference:
                    payload.pop(field, None)
        return payload
    kwargs['json'] = [adapt(row) for row in data] if isinstance(data, list) else adapt(data)
    response = method(_url(RESOURCES['cadastros']), **kwargs)
    if response.status_code == 400:
        try:
            error = response.json()
        except ValueError:
            error = {}
        # Falhas de schema ocorrem antes da escrita; somente essas permitem repetição.
        if error.get('code') == 'PGRST204':
            message = error.get('message', '')
            missing = next((field for field in ('darpe_musico', 'darpe_organista', 'darpe_musico_organista')
                            if f"'{field}'" in message), None)
            if missing:
                removed = {missing} if missing == 'darpe_musico_organista' else {'darpe_musico', 'darpe_organista'}
                def strip(payload):
                    return {key: value for key, value in payload.items() if key not in removed}
                kwargs['json'] = [strip(row) for row in kwargs['json']] if isinstance(data, list) else strip(kwargs['json'])
                response = method(_url(RESOURCES['cadastros']), **kwargs)
    return response


def _normalize_payload(payload, resource):
    payload = {key: (value.strip() or None) if isinstance(value, str) else value
               for key, value in payload.items()}
    dates = {'data_nascimento': 'Data de nascimento', 'data_batismo': 'Data de batismo',
             'declaracao_data': 'Data da declaração', 'data_evento': 'Data do evento'}
    for key, label in dates.items():
        if payload.get(key):
            try:
                date.fromisoformat(payload[key])
            except (ValueError, TypeError):
                raise ValueError(f'{label}: informe uma data válida.')
    if resource == 'cadastros':
        for key, size in (('cpf', 11), ('cep', 8)):
            if payload.get(key):
                digits = re.sub(r'\D', '', str(payload[key]))
                if len(digits) != size:
                    raise ValueError(f'{key.upper()}: informe {size} dígitos ou deixe em branco.')
                payload[key] = digits
        if payload.get('uf'):
            payload['uf'] = payload['uf'].upper()
            if not re.fullmatch('[A-Z]{2}', payload['uf']):
                raise ValueError('UF: informe a sigla com duas letras ou deixe em branco.')
        if 'status' in payload:
            payload['status'] = payload['status'] or 'ATIVO'
    return payload


def _actor(request):
    profile = request.session.get("user_profile") or {}
    return {
        "id": str(request.session.get("user_id") or profile.get("user_id") or ""),
        "name": str(profile.get("full_name") or profile.get("nome") or profile.get("email") or "Usuário autenticado").strip(),
    }


def _denied():
    return JsonResponse({"error": "Seu perfil não possui acesso à pasta D.A.R.P.E."}, status=403)


def _url(config):
    return f"{settings.SUPABASE_URL}/rest/v1/{config['table']}"


def _phone(value):
    digits = "".join(char for char in str(value or "") if char.isdigit())
    if len(digits) == 11:
        return f"({digits[:2]}) {digits[2]} {digits[3:7]}-{digits[7:]}"
    if len(digits) == 10:
        return f"({digits[:2]}) {digits[2:6]}-{digits[6:]}"
    return str(value or "")


def _with_canonical_commons(rows):
    """Acrescenta o código BR do catálogo sem modificar o valor legado salvo."""
    try:
        resolver = CommonResolver(common_catalog())
    except (requests.RequestException, ValueError):
        resolver = None
    enriched = []
    for original in rows:
        row = dict(original)
        canonical = resolver.resolve(row.get("comum_congregacao"), row.get("cidade")) if resolver else None
        row["comum_exibicao"] = (canonical or {}).get("comum") or row.get("comum_congregacao")
        enriched.append(row)
    try:
        from .darpe_sam import enrich_roles
        enriched = enrich_roles(enriched)
    except (requests.RequestException, ValueError):
        logger.warning('Não foi possível conferir o âmbito dos encarregados no SAM.')
    return enriched


def _canonical_city(value):
    """Retorna uma única grafia oficial para cada município regional."""
    cities = {_norm(city): city for city in REGIONAL_MUNICIPALITIES}
    return cities.get(_norm(value))


def _regional_rows(rows):
    normalized = []
    for original in rows:
        city = _canonical_city(original.get("cidade"))
        if city:
            normalized.append({**original, "cidade": city})
    return normalized


def _import_clean(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    value = str(value).strip()
    return value or None


def _import_yes(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char)).strip().upper()
    return value in {"SIM", "S", "TRUE", "1", "X"}


def _musical_flags(instrument, sex, selected):
    """Converte a coluna legada em funções mutuamente exclusivas."""
    if not selected:
        return False, False
    organist = "ORGAO" in _norm(instrument) or _norm(sex) == "FEMININO"
    return not organist, organist


def _import_key(row):
    return _norm(row.get("nome_completo")), _norm(row.get("comum_congregacao"))


def _resolve_import_common(resolver, common, city):
    canonical = resolver.resolve(common, city)
    if canonical or _norm(common).startswith("BR-"):
        return canonical

    def name_key(value, municipality):
        value = _norm(value)
        for short, full in (("JD", "JARDIM"), ("VL", "VILA"), ("DR", "DOUTOR")):
            value = re.sub(rf"\b{short}\.?\s+", full + " ", value)
        value = re.sub(r"\s+-\s+" + re.escape(_norm(municipality)) + r"$", "", value)
        if value in {"CENTRAL DE " + _norm(municipality), _norm(municipality) + " - CENTRAL"}:
            return "CENTRAL"
        # A primeira Rosemary consta no catálogo sem o numeral I.
        if _norm(municipality) == "ITAPEVI" and value == "JARDIM ROSEMARY I":
            return "JARDIM ROSEMARY"
        return value

    matches = []
    for candidate in resolver.by_code.values():
        municipality = candidate.get("cidade")
        if city and _norm(city) != _norm(municipality):
            continue
        description = re.sub(r"^BR-\d+-\d+\s*-\s*", "", candidate["comum"])
        if name_key(common, municipality) == name_key(description, municipality):
            matches.append(candidate)
    return matches[0] if len(matches) == 1 else None


def _read_imported_cadastros(upload):
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook["Cadastro"]
    except (KeyError, OSError, ValueError) as exc:
        raise ValueError("Selecione a planilha oficial com a aba Cadastro.") from exc
    resolver = CommonResolver(common_catalog())
    actor_id = "importacao-darpe-planilha"
    actor_name = "Importação do mapeamento regional D.A.R.P.E"
    rows = []
    unresolved = []
    for line, values in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
        if not _import_clean(values[1]):
            continue
        common = _import_clean(values[7]) or "NÃO INFORMADA"
        city = _import_clean(values[13])
        canonical = _resolve_import_common(resolver, common, city)
        if canonical:
            common, city = canonical["comum"], canonical["cidade"]
        city = _canonical_city(city)
        if not city:
            unresolved.append(f"linha {line}: {common}")
        musician, organist = _musical_flags(values[2], values[5], _import_yes(values[22]))
        rows.append({
            "cpf": _import_clean(values[0]), "nome_completo": _import_clean(values[1]),
            "instrumento_oficializado": _import_clean(values[2]), "rg": _import_clean(values[3]),
            "data_nascimento": _import_clean(values[4]), "sexo": _import_clean(values[5]),
            "estado_civil": _import_clean(values[6]), "comum_congregacao": common,
            "data_batismo": _import_clean(values[8]), "logradouro": _import_clean(values[9]),
            "numero": _import_clean(values[10]), "complemento": _import_clean(values[11]),
            "bairro": _import_clean(values[12]), "cidade": city,
            "uf": (_import_clean(values[14]) or "").upper() or None, "cep": _import_clean(values[15]),
            "ddd": _import_clean(values[16]), "telefone": _import_clean(values[17]),
            "darpe_ministerio": _import_yes(values[18]), "darpe_atendente": _import_yes(values[19]),
            "darpe_colaborador": _import_yes(values[20]), "darpe_encarregado": _import_yes(values[21]),
            "darpe_musico": musician, "darpe_organista": organist,
            "darpe_musico_organista": musician or organist, "darpe_grupo_canto": _import_yes(values[23]),
            "declaracao_local": _import_clean(values[24]), "declaracao_data": _import_clean(values[25]),
            "nome_ministerio": _import_clean(values[26]),
            "nome_encarregado_orquestra": _import_clean(values[27]),
            "nome_voluntario": _import_clean(values[28]), "status": "ATIVO",
            "created_by_user_id": actor_id, "created_by_name": actor_name,
            "updated_by_user_id": actor_id, "updated_by_name": actor_name,
        })
    keys = [_import_key(row) for row in rows]
    workbook.close()
    if not rows:
        raise ValueError("A aba Cadastro não possui registros para importar.")
    if len(keys) != len(set(keys)):
        raise ValueError(f"A planilha possui {len(keys) - len(set(keys))} cadastro(s) repetido(s).")
    if unresolved:
        raise ValueError(f"Existem {len(unresolved)} cadastro(s) sem município regional confirmado. "
                         + "; ".join(unresolved) + ". Informe o município ou o código BR da comum na planilha.")
    return rows


def _read_all_rows(config, select="*", filters=None):
    """Le todas as paginas do PostgREST, inclusive acima de 1.000 linhas."""
    rows, expected = [], None
    while True:
        response = requests.get(_url(config), headers=service_headers("count=exact"),
                                params={"select": select, "order": "id.asc", "offset": len(rows), "limit": 1000,
                                        **(filters or {})},
                                timeout=30)
        response.raise_for_status()
        page = response.json()
        rows.extend(page)
        total = str(response.headers.get("Content-Range", "")).rsplit("/", 1)[-1]
        if total.isdigit():
            if expected is not None and expected != int(total):
                raise ValueError('A base mudou durante a leitura. Atualize a conferência e tente novamente.')
            expected = int(total)
            if not page and len(rows) < expected:
                raise ValueError('A leitura da base ficou incompleta. Tente novamente.')
        if not page or (total.isdigit() and len(rows) >= int(total)):
            return rows
        if not total.isdigit() and len(page) < 1000:
            return rows


def import_cadastros(request):
    if not can_access_module(request, MODULE_DARPE):
        return _denied()
    if request.method != "POST":
        return JsonResponse({"error": "Método não permitido."}, status=405)
    upload = request.FILES.get("arquivo")
    if not upload:
        return JsonResponse({"error": "Selecione a planilha oficial do D.A.R.P.E."}, status=400)
    try:
        rows = _read_imported_cadastros(upload)
        url = _url(RESOURCES["cadastros"])
        current = _read_all_rows(RESOURCES["cadastros"])
        reference = current[0] if current else None
        existing = {}
        for row in current:
            existing.setdefault(_import_key(row), []).append(row)
        pending, keep = [], set()
        for row in rows:
            matches = existing.get(_import_key(row)) or []
            previous = matches[0] if matches else None
            record_id = previous["id"] if previous else str(uuid4())
            keep.add(record_id)
            if previous:
                if not row.get('instrumento_oficializado') and previous.get('instrumento_oficializado'):
                    row = {**row, 'instrumento_oficializado': previous['instrumento_oficializado']}
                elif row.get('instrumento_oficializado') and previous.get('instrumento_oficializado'):
                    from .darpe_sam import instrument_key
                    if instrument_key(row['instrumento_oficializado']) == instrument_key(previous['instrumento_oficializado']):
                        row = {**row, 'instrumento_oficializado': previous['instrumento_oficializado']}
                row = {**row, **{key: previous[key] for key in
                       ("created_by_user_id", "created_by_name") if key in previous}}
                fields = [key for key in row if not key.startswith(("created_by_", "updated_by_"))
                          and key in previous]
                if all(row[key] == previous.get(key) for key in fields):
                    continue
            pending.append({**row, "id": record_id})
        # Upsert por UUID preserva os vinculos de participantes e a autoria original.
        for start in range(0, len(pending), 500):
            response = _write_cadastro(requests.post, reference=reference,
                                      headers=service_headers("resolution=merge-duplicates,return=minimal"),
                                      params={"on_conflict": "id"}, json=pending[start:start + 500], timeout=60)
            response.raise_for_status()
        # Limpeza somente depois de TODOS os lotes gravados com sucesso.
        stale = [row["id"] for row in current if row["id"] not in keep]
        for start in range(0, len(stale), 100):
            response = requests.delete(url, headers=service_headers("return=minimal"),
                                       params={"id": "in.(" + ",".join(stale[start:start + 100]) + ")"}, timeout=30)
            response.raise_for_status()
        return JsonResponse({"message": "Importação concluída.", "total": len(rows)})
    except ValueError as exc:
        return JsonResponse({"error": str(exc)}, status=400)
    except requests.RequestException:
        return JsonResponse({"error": "Não foi possível concluir a importação no banco de dados."}, status=502)


def export_cadastros_excel(request):
    if not can_access_module(request, MODULE_DARPE):
        return _denied()
    config = RESOURCES["cadastros"]
    try:
        response = requests.get(_url(config), headers=service_headers(),
                                params={"select": "*", "order": config["order"]}, timeout=30)
        response.raise_for_status()
        rows = filter_rows(user_scope(request), [{**row, **config["location"](row)} for row in response.json()])
        rows = _with_canonical_commons(_regional_rows(rows))
    except requests.RequestException:
        return JsonResponse({"error": "Não foi possível consultar os cadastros para exportação."}, status=502)
    cities = {_norm(value) for value in request.GET.getlist("cidade") if str(value).strip()}
    commons = {str(value).strip().upper() for value in request.GET.getlist("comum") if str(value).strip()}
    search = str(request.GET.get("pesquisa") or "").strip().upper()
    rows = [row for row in rows if (not cities or _norm(row.get("cidade")) in cities)
            and (not commons or str(row.get("comum_congregacao") or "").upper() in commons)
            and (not search or search in json.dumps(row, ensure_ascii=False, default=str).upper())]
    profile = request.session.get("user_profile") or {}
    actor = profile.get("full_name") or profile.get("nome") or profile.get("email") or "Usuário do sistema"
    now = datetime.now()
    wb = Workbook(); ws = wb.active; ws.title = "CADASTROS"
    headers = ["Nome completo", "Comum", "Município", "Telefone", "Instrumento", "Funções", "Status"]
    last = get_column_letter(len(headers)); navy, pale = "1E4B7A", "EAF2F8"; thin = Side(style="thin", color="CCD5DD")
    for number, text, size, color, fill in ((1,"CONGREGAÇÃO CRISTÃ NO BRASIL",15,"FFFFFF",navy),(2,"Regional Itapevi - São Paulo",10,"FFFFFF",navy),(3,"D.A.R.P.E",12,navy,pale)):
        ws.merge_cells(f"A{number}:{last}{number}"); cell=ws.cell(number,1,text); cell.font=Font(size=size,bold=number!=2,color=color); cell.fill=PatternFill("solid",fgColor=fill); cell.alignment=Alignment(horizontal="center")
    city_scope = ", ".join(request.GET.getlist("cidade")) or "Todos os municípios"
    common_scope = ", ".join(request.GET.getlist("comum")) or "Todas as comuns"
    ws.merge_cells("A4:D4"); ws["A4"] = f"Cadastro regional · Municípios: {city_scope} · Comuns: {common_scope}"
    ws.merge_cells(f"E4:{last}4"); ws["E4"] = f"Emissão: {now:%d/%m/%Y %H:%M} · Responsável: {actor}"; ws["E4"].alignment=Alignment(horizontal="right")
    for cell in ws[4]: cell.font=Font(size=9,bold=True,color="536A7D")
    for column,label in enumerate(headers,1):
        cell=ws.cell(6,column,label); cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=navy); cell.alignment=Alignment(horizontal="center"); cell.border=Border(bottom=thin)
    role_fields=(("darpe_ministerio","Ministério"),("darpe_atendente","Atendente"),("darpe_colaborador","Colaborador"),("darpe_encarregado","Encarregado"),("darpe_musico_organista","Músico / Organista"),("darpe_grupo_canto","Grupo de Canto"))
    for index,row in enumerate(sorted(rows,key=lambda x:(str(x.get("cidade") or ""),str(x.get("comum_congregacao") or ""),str(x.get("nome_completo") or ""))),7):
        functions=", ".join(label for key,label in role_fields if row.get(key))
        values=[row.get("nome_completo"),row.get("comum_exibicao") or row.get("comum_congregacao"),row.get("cidade"),_phone("".join(filter(None,[str(row.get("ddd") or ""),str(row.get("telefone") or "")]))),row.get("instrumento_oficializado"),functions,row.get("status")]
        for column,value in enumerate(values,1):
            cell=ws.cell(index,column,value or "")
            cell.border=Border(bottom=thin)
            cell.alignment=Alignment(vertical="center",wrap_text=column in {1,2,6})
            if index%2==0:
                cell.fill=PatternFill("solid",fgColor="F4F6F8")
    for index,width in enumerate([34,42,22,20,24,38,14],1): ws.column_dimensions[get_column_letter(index)].width=width
    ws.freeze_panes="A7"; ws.auto_filter.ref=f"A6:{last}{max(6,ws.max_row)}"; ws.sheet_view.showGridLines=False; ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0; ws.sheet_properties.pageSetUpPr.fitToPage=True; ws.print_title_rows="1:6"
    stream=io.BytesIO(); wb.save(stream); result=HttpResponse(stream.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    result["Content-Disposition"] = f"attachment; filename*=UTF-8''{quote(f'Cadastro_DARPE_{now:%d-%m-%Y_%H-%M}.xlsx')}"
    return result


@never_cache
def page(request, section="dashboard"):
    if not can_access_module(request, MODULE_DARPE):
        return render(request, "pages/403.html", {"message": "Seu perfil não possui acesso à pasta D.A.R.P.E."}, status=403)
    if section not in {"dashboard", "cadastros", "eventos"}:
        section = "dashboard"
    return render(request, "pages/darpe.html", {"section": section, "scope": scope_details(user_scope(request))})


@never_cache
def api_resource(request, resource, record_id=None):
    if not can_access_module(request, MODULE_DARPE):
        return _denied()
    config = RESOURCES.get(resource)
    if not config:
        return JsonResponse({"error": "Recurso inválido."}, status=404)
    scope = user_scope(request)
    try:
        if request.method == "GET" and not record_id:
            rows = _read_all_rows(config) if resource == "cadastros" else None
            if rows is None:
                response = requests.get(_url(config), headers=service_headers(),
                                        params={"select": "*", "order": config["order"]}, timeout=15)
                response.raise_for_status()
                rows = response.json()
            if resource != "participantes":
                rows = filter_rows(scope, [{**row, **config["location"](row)} for row in rows])
            if resource == "cadastros":
                rows = _with_canonical_commons(_regional_rows(rows))
            return JsonResponse({"items": rows, "scope": scope_details(scope)})

        current = None
        if record_id:
            response = requests.get(_url(config), headers=service_headers(),
                                    params={"select": "*", "id": f"eq.{record_id}", "limit": 1}, timeout=15)
            response.raise_for_status()
            rows = response.json()
            current = rows[0] if rows else None
            if not current:
                return JsonResponse({"error": "Registro não encontrado."}, status=404)
            if resource != "participantes" and not can_access(scope, config["location"](current)):
                return _denied()

        if request.method in {"POST", "PATCH"}:
            payload = {key: value for key, value in json.loads(request.body or "{}").items() if key in config["fields"]}
            if not payload:
                return JsonResponse({"error": "Informe os dados do registro."}, status=400)
            payload = _normalize_payload(payload, resource)
            candidate = {**(current or {}), **payload}
            if resource == "cadastros" and candidate.get("darpe_musico") and candidate.get("darpe_organista"):
                return JsonResponse({"error": "Selecione Músico ou Organista, nunca as duas funções."}, status=400)
            if resource == "cadastros" and candidate.get("cidade"):
                canonical_city = _canonical_city(candidate["cidade"])
                if canonical_city:
                    payload["cidade"] = canonical_city
                    candidate["cidade"] = canonical_city
            if resource != "participantes" and not can_access(scope, config["location"](candidate)):
                return _denied()
            if resource == "cadastros" and not str(candidate.get("nome_completo") or "").strip():
                return JsonResponse({"error": "Informe o nome completo."}, status=400)
            if resource == "cadastros":
                if not candidate.get('comum_congregacao'):
                    return JsonResponse({'error': 'Informe a comum congregação.'}, status=400)
                if any(key in payload for key in ('darpe_musico', 'darpe_organista')):
                    payload['darpe_musico_organista'] = bool(candidate.get('darpe_musico') or candidate.get('darpe_organista'))
            actor = _actor(request)
            if not actor["id"]:
                return JsonResponse({"error": "A sessão não possui autoria válida."}, status=401)
            payload.update({"updated_by_user_id": actor["id"], "updated_by_name": actor["name"]})
            if not record_id:
                payload.update({"created_by_user_id": actor["id"], "created_by_name": actor["name"]})
            method = requests.patch if record_id else requests.post
            options = dict(headers=service_headers("return=representation"),
                           params={"id": f"eq.{record_id}"} if record_id else None, json=payload, timeout=15)
            response = (_write_cadastro(method, reference=current, **options) if resource == 'cadastros'
                        else method(_url(config), **options))
            response.raise_for_status()
            return JsonResponse({"item": response.json()[0]}, status=200 if record_id else 201)

        if request.method == "DELETE" and record_id:
            response = requests.delete(_url(config), headers=service_headers("return=minimal"),
                                       params={"id": f"eq.{record_id}"}, timeout=15)
            response.raise_for_status()
            return JsonResponse({}, status=204)
        return JsonResponse({"error": "Método não permitido."}, status=405)
    except json.JSONDecodeError:
        return JsonResponse({"error": "JSON inválido."}, status=400)
    except ValueError as exc:
        return JsonResponse({'error': str(exc)}, status=400)
    except requests.HTTPError as exc:
        response = exc.response
        try:
            code = response.json().get('code') if response is not None else None
        except ValueError:
            code = None
        logger.warning('DARPE %s %s: HTTP %s, código %s', request.method, resource,
                       response.status_code if response is not None else None, code)
        if code == '23505':
            return JsonResponse({'error': 'Já existe um cadastro com este CPF ou com o mesmo nome e comum.'}, status=409)
        if code in {'22007', '22008', '22P02', '23514', '23502'}:
            return JsonResponse({'error': 'Confira os campos obrigatórios, datas, CPF, CEP, UF e situação do registro.'}, status=400)
        return JsonResponse({'error': 'O banco recusou a operação do D.A.R.P.E. Tente novamente ou informe o suporte.'}, status=502)
    except requests.RequestException:
        return JsonResponse({"error": "Não foi possível acessar os dados do D.A.R.P.E."}, status=502)
