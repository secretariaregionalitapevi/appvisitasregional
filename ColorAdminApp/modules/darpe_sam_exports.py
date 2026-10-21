"""Relatórios da conferência DARPE/SAM no padrão institucional do projeto."""
from collections import Counter
from math import ceil
from datetime import datetime
from io import BytesIO
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

import requests
from django.http import HttpResponse, JsonResponse
from django.views.decorators.cache import never_cache
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .darpe import _actor, _denied
from .darpe_sam import LABELS, filtered, read_comparison
from .module_access import MODULE_DARPE, can_access_module

NAVY, PALE = '1E4B7A', 'EAF2F8'
TITLE = 'Conferência DARPE × SAM'
HEADERS = ['Nome completo', 'Município', 'Comum', 'Instrumento SAM', 'Oficialização SAM',
           'Instrumento DARPE', 'Situação DARPE', 'Correspondência', 'Resultado']


def report_values(row):
    result = LABELS[row['situacao']]
    if row.get('sam_nome') and row['sam_nome'] != row['nome']:
        result += ' · Nome SAM: ' + row['sam_nome']
    level = row.get('sam_nivel') or ('Não se aplica' if row['situacao'] in {'ministerio', 'sem_funcao_musical'} else 'Não localizada')
    if row.get('cargo_sam'):
        level += ' · ' + row['cargo_sam']
    if row['situacao'] in {'ambiguo', 'nome_divergente'} and row.get('sam_nome'):
        level = 'A confirmar: ' + level
    return [row['nome'], row['cidade'], row['comum'], row.get('sam_instrumento') or '—',
            level, row.get('darpe_instrumento') or 'Não informado',
            row.get('darpe_status') or 'Não informado', row.get('correspondencia') or 'Não encontrada',
            result]


def report_scope(params):
    cities = ', '.join(params.getlist('cidade')) or 'Todos os municípios'
    parts = [cities, params.get('comum') or 'Todas as comuns',
             LABELS.get(params.get('situacao'), 'Todos os resultados'),
             params.get('darpe_status') or 'Ativos e inativos']
    if params.get('q'):
        parts.append('Pesquisa: ' + params['q'])
    return ' · '.join(parts)


def build_excel(rows, scope, actor, now):
    book = Workbook()
    sheet = book.active
    sheet.title = 'CONFERÊNCIA DARPE SAM'
    for number, label, size, background, foreground in (
        (1, 'CONGREGAÇÃO CRISTÃ NO BRASIL', 15, NAVY, 'FFFFFF'),
        (2, 'Regional Itapevi - São Paulo', 10, NAVY, 'FFFFFF'),
        (3, 'D.A.R.P.E · ' + TITLE, 12, PALE, NAVY),
    ):
        sheet.merge_cells(start_row=number, start_column=1, end_row=number, end_column=9)
        cell = sheet.cell(number, 1, label)
        cell.font = Font(size=size, bold=number != 2, color=foreground)
        cell.fill = PatternFill('solid', fgColor=background)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        sheet.row_dimensions[number].height = 25 if number == 1 else 21
    sheet.merge_cells('A4:E4')
    sheet['A4'] = 'Recorte: ' + scope
    sheet.merge_cells('F4:I4')
    sheet['F4'] = f'Emissão: {now:%d/%m/%Y %H:%M} · Responsável: {actor}'
    sheet.row_dimensions[4].height = 45
    for cell in (sheet['A4'], sheet['F4']):
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.font = Font(size=9, color='536A7D')
    counts = Counter(row['correspondencia'] for row in rows)
    sheet.merge_cells('A5:I5')
    sheet['A5'] = (f'{len(rows)} cadastros DARPE · Correspondência confirmada: {counts["Confirmada"]}'
                   f' · Não confirmada: {counts["Não confirmada"]} · Não encontrada: {counts["Não encontrada"]}'
                   f' · Conferência não aplicável: {counts["Não se aplica"]}')
    sheet['A5'].font = Font(size=9, bold=True, color=NAVY)
    border = Border(bottom=Side(style='thin', color='CCD5DD'))
    for col, label in enumerate(HEADERS, 1):
        cell = sheet.cell(6, col, label)
        cell.fill = PatternFill('solid', fgColor=NAVY)
        cell.font = Font(size=9, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    sheet.row_dimensions[6].height = 30
    widths = [36, 24, 48, 23, 24, 23, 17, 24, 38]
    for number, row in enumerate(rows, 7):
        values = report_values(row)
        for col, value in enumerate(values, 1):
            cell = sheet.cell(number, col, str(value))
            # Nomes e localidades são dados, inclusive quando começam por '='.
            cell.data_type = 's'
            cell.font = Font(size=9, color='263238')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.border = border
            if number % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='F4F6F8')
        sheet.row_dimensions[number].height = max(38, 13 * max(ceil(len(str(value)) / (width - 3)) for value, width in zip(values, widths)) + 8)
        status = sheet.cell(number, 7)
        status.fill = PatternFill('solid', fgColor='D9F5E8' if row.get('darpe_status') == 'ATIVO' else 'E9EDF2')
        status.font = Font(size=9, bold=True, color='08724D' if row.get('darpe_status') == 'ATIVO' else '586577')
        status.alignment = Alignment(horizontal='center', vertical='center')
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = 'A7'
    sheet.auto_filter.ref = f'A6:I{max(6, sheet.max_row)}'
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = '1:6'
    sheet.print_area = f'A1:I{sheet.max_row}'
    sheet.oddFooter.center.text = 'D.A.R.P.E · Página &P de &N'
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def build_pdf(rows, scope, actor, now):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    page_width, page_height = landscape(A4)
    usable = page_width - 56
    normal = ParagraphStyle('normal', fontName='Helvetica', fontSize=7.2, leading=9.4, textColor='#263238')
    header = ParagraphStyle('header', parent=normal, fontName='Helvetica-Bold', textColor=colors.white, alignment=TA_CENTER)
    context = ParagraphStyle('context', parent=normal, fontSize=8, leading=10, textColor='#1e4b7a')
    metadata = ParagraphStyle('metadata', parent=normal, fontSize=7, leading=9, alignment=TA_RIGHT)
    def p(value, style=normal):
        return Paragraph(escape(str(value)), style)

    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self.states = []

        def showPage(self):
            self.states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self.states)
            for state in self.states:
                self.__dict__.update(state)
                self.saveState()
                center = page_width / 2
                self.setFillColor(colors.HexColor('#263238'))
                for y, text, font, size, color in (
                    (page_height - 29, 'CONGREGAÇÃO CRISTÃ NO BRASIL', 'Helvetica-Bold', 15, '#263238'),
                    (page_height - 43, 'Regional Itapevi - São Paulo', 'Helvetica', 9, '#263238'),
                    (page_height - 60, 'D.A.R.P.E', 'Helvetica-Bold', 12, '#1e4b7a'),
                    (page_height - 74, TITLE, 'Helvetica', 9, '#263238'),
                ):
                    self.setFont(font, size)
                    self.setFillColor(colors.HexColor(color))
                    self.drawCentredString(center, y, text)
                details = Paragraph(f'<b>Página {self._pageNumber} de {total}</b><br/>'
                                    f'Emissão: {now:%d/%m/%Y %H:%M}<br/>'
                                    f'Recorte: {escape(scope)}<br/>Responsável: {escape(actor)}', metadata)
                _, height = details.wrap(210, 90)
                details.drawOn(self, page_width - 238, page_height - 16 - height)
                self.setFont('Helvetica', 7)
                self.setFillColor(colors.HexColor('#6b7280'))
                self.drawString(28, 19, f'D.A.R.P.E · {len(rows)} cadastros · Conferência SAM')
                self.drawRightString(page_width - 28, 19, f'Página {self._pageNumber} de {total}')
                self.restoreState()
                super().showPage()
            super().save()

    data = [[p(label, header) for label in HEADERS]]
    for row in rows:
        data.append([p(value) for value in report_values(row)])
    widths = [120, 62, 128, 67, 71, 68, 45, 90, usable - 651]
    table = LongTable(data, colWidths=widths, repeatRows=1, splitByRow=1, hAlign='LEFT')
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e4b7a')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f4f6f8')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), .35, colors.HexColor('#d8dee5')),
        ('LINEBELOW', (0, 0), (-1, -1), .45, colors.HexColor('#c8d1da')),
        ('LEFTPADDING', (0, 0), (-1, -1), 4), ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 5), ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    band = Table([[p(f'RECORTE DO RELATÓRIO · {scope} · {len(rows)} CADASTROS', context)]], colWidths=[usable])
    band.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#eaf2f8')),
                             ('TOPPADDING', (0, 0), (-1, -1), 8), ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                            topMargin=100, bottomMargin=34, title=TITLE, author=actor)
    doc.build([band, Spacer(1, 8), table], canvasmaker=NumberedCanvas)
    return stream.getvalue()


@never_cache
def export(request, format):
    if not can_access_module(request, MODULE_DARPE):
        return _denied()
    if request.method != 'GET':
        return JsonResponse({'error': 'Método não permitido.'}, status=405)
    try:
        rows, _ = read_comparison(request)
        rows = filtered(rows, request.GET)
        if not rows:
            return JsonResponse({'error': 'Nenhum cadastro encontrado nos filtros selecionados.'}, status=400)
        actor = _actor(request)['name']
        now = datetime.now(ZoneInfo('America/Sao_Paulo'))
        scope = report_scope(request.GET)
        data = build_pdf(rows, scope, actor, now) if format == 'pdf' else build_excel(rows, scope, actor, now)
        response = HttpResponse(data, content_type='application/pdf' if format == 'pdf' else
                                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="Conferencia_DARPE_SAM_{now:%d-%m-%Y_%H-%M}.{format}"'
        return response
    except requests.RequestException:
        return JsonResponse({'error': 'Não foi possível consultar os cadastros para exportação.'}, status=502)
    except ValueError:
        return JsonResponse({'error': 'A base mudou durante a consulta. Atualize os dados e tente novamente.'}, status=400)
