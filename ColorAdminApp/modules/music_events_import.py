"""Leitura da aba nominal AppCont. Não usa resumos nem presume uma base SAM."""
import hashlib
import io
import unicodedata
import zipfile
from collections import Counter
from datetime import datetime
from uuid import UUID

from openpyxl import load_workbook


def norm(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFKD', str(value or ''))
                           if not unicodedata.combining(c)).upper().split())


def category(cargo, instrument, family):
    role, inst = norm(cargo), norm(instrument)
    if inst in {'ORGAO', 'ORGANISTA'} or 'ORGANISTA' in role or (
            norm(family) == 'TECLADO' and inst != 'ACORDEON'):
        return 'ORGANISTA'
    if inst or role in {'MUSICO', 'MUSICISTA'}:
        return 'MUSICO'
    if any(word in role for word in ('PORTEIR', 'BOMBEIR', 'MEDIC', 'ENFERMEIR', 'APOIO')):
        return 'APOIO'
    if role and role != 'IRMANDADE':
        return 'MINISTERIO'
    return 'IRMANDADE'


def parse_workbook(content):
    if len(content) > 15 * 1024 * 1024:
        raise ValueError('O arquivo deve ter até 15 MB.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Planilha muito grande após descompactação.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError('Envie um arquivo Excel .xlsx válido.') from exc
    try:
        candidates = []
        for sheet in workbook:
            for index, row in enumerate(sheet.iter_rows(max_row=8, values_only=False), 1):
                header = [norm(cell.value) for cell in row]
                required = {'UUID', 'NOME COMPLETO', 'COMUM', 'CIDADE', 'CARGO', 'LOCAL_ENSAIO', 'DATA_ENSAIO'}
                if required.issubset(header):
                    candidates.append((sheet, index, header))
        if len(candidates) != 1:
            raise ValueError('É necessária uma única aba nominal com UUID, nome, comum, cidade, cargo, local e data do ensaio.')
        sheet, header_row, header = candidates[0]
        records, duplicates, dates, errors = {}, [], Counter(), []
        locations, categories, families = Counter(), Counter(), Counter()
        raw_count = 0
        for line, cells in enumerate(sheet.iter_rows(min_row=header_row + 1), header_row + 1):
            if line > 20010:
                raise ValueError('Limite de 20 mil linhas por importação.')
            source = {key: cells[i].value for i, key in enumerate(header) if key and i < len(cells)}
            if not source.get('UUID') and not source.get('NOME COMPLETO'):
                continue
            raw_count += 1
            try:
                if any(cell.data_type == 'f' for cell in cells):
                    raise ValueError('fórmula na linha nominal')
                identity = str(UUID(str(source.get('UUID') or '').strip()))
                name = str(source.get('NOME COMPLETO') or '').strip()
                local = str(source.get('LOCAL_ENSAIO') or '').strip()
                if not name or not local:
                    raise ValueError('nome/local ausente')
                moment = source.get('DATA_ENSAIO')
                if isinstance(moment, datetime):
                    day = moment.date().isoformat()
                else:
                    text = str(moment or '').strip()
                    day = datetime.strptime(text[:10], '%d/%m/%Y' if '/' in text else '%Y-%m-%d').date().isoformat()
                dates[day] += 1
                role = str(source.get('CARGO') or '').strip()
                instrument = str(source.get('INSTRUMENTO') or '').strip()
                family = norm(source.get('NAIPE_INSTRUMENTO')) or 'OUTROS'
                kind = category(role, instrument, family)
                if kind == 'ORGANISTA':
                    family = 'TECLADO'
                elif family not in {'CORDAS', 'MADEIRAS', 'METAIS', 'TECLADO', 'OUTROS'}:
                    family = 'OUTROS'
                record = {'origem_registro_id': identity, 'nome': name[:250], 'local': local[:180],
                          'comum_informada': str(source.get('COMUM') or '').strip()[:250],
                          'municipio_informado': str(source.get('CIDADE') or '').strip()[:180],
                          'categoria': kind, 'familia': family, 'cargo': role[:180],
                          'instrumento': instrument[:180], 'nivel': str(source.get('NIVEL') or '')[:100],
                          'registrado_por_legado': str(source.get('REGISTRADO_POR') or '')[:250],
                          'data_original': day, 'data_hora_original': moment.isoformat() if isinstance(moment, datetime) else str(moment),
                          'observacoes': str(source.get('ANOTACOES') or '')[:4000]}
                if identity in records:
                    previous = records[identity]
                    if norm(previous['nome']) != norm(record['nome']) or norm(previous['local']) != norm(record['local']):
                        raise ValueError('UUID repetido com pessoa ou local diferente; corrija o arquivo')
                    duplicates.append({'linha': line, 'uuid': identity})
                records[identity] = record  # Regra explícita: última ocorrência do UUID.
            except (ValueError, TypeError) as exc:
                errors.append({'linha': line, 'motivo': str(exc)})
        if errors:
            raise ValueError('Linhas inválidas: ' + '; '.join(f"{e['linha']}: {e['motivo']}" for e in errors[:12]))
        if not records:
            raise ValueError('A aba nominal está vazia.')
        for record in records.values():
            locations[record['local']] += 1
            categories[record['categoria']] += 1
            families[record['familia']] += 1
        names = Counter((norm(row['nome']),norm(row['comum_informada'])) for row in records.values())
        return {'sha256': hashlib.sha256(content).hexdigest(), 'sheet': sheet.title,
                'raw_count': raw_count, 'count': len(records), 'duplicates': len(duplicates),
                'possible_duplicates': sum(value-1 for value in names.values() if value>1),
                'dates': dict(dates), 'suggested_date': dates.most_common(1)[0][0],
                'locations': dict(locations), 'categories': dict(categories), 'families': dict(families),
                'records': list(records.values())}
    finally:
        workbook.close()
