"""Exportações do MESMO recorte do painel, conforme REPORT_EXPORT_STANDARD."""
import io
import math
import re
from datetime import datetime
from html import escape
from pathlib import Path
from zoneinfo import ZoneInfo

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, Image, PageBreak, Flowable
from reportlab.pdfgen.canvas import Canvas
from reportlab.graphics.shapes import Drawing, Rect, String

NAVY='1E4B7A'; PALE='EAF2F8'
HEADERS=['Município','Comum','Encarregado: situação / local','Músicos','Organistas','Total musical',
         'Base músicos','Base organistas','% Músicos','% Organistas']


def vertical_label(text,height):
    drawing=Drawing(14,height)
    drawing.add(String(7,height/2,text,fontName='Helvetica-Bold',fontSize=8,
        textAnchor='middle',angle=90,fillColor=colors.black))
    return drawing


def report_note(data):
    if data.get('report')=='encarregados':
        return 'Presença confirmada pela correspondência entre o encarregado cadastrado e a lista nominal do ensaio.'
    if data.get('registry',{}).get('available') and not data['base_complete']:
        return 'Base SAM atual: Ativo, Alerta e Concluído, todos os níveis. Percentual = presentes da comum ÷ cadastro SAM atual; referência válida na data de emissão.'
    return 'Sem referência cadastral para este recorte.' if not data['base_complete'] else 'Percentuais sobre a base histórica SAM e presenças elegíveis vinculadas.'


def scope_label(data):
    f=data['filters']
    events=(', '.join(e['titulo'] for e in data['events']) if len(data['events'])<=3 else f"{len(data['events'])} eventos no período") or 'Sem eventos'
    parts=[events,', '.join(f.get('municipios',[])) or ('Regional Itapevi' if f.get('origem','regional')=='regional' else 'Regional e visitantes' if f.get('origem')=='todos' else 'Visitantes')]
    if f.get('comuns'):parts.append('Comuns: '+', '.join(f['comuns']))
    if f.get('ano'):parts.append('Ano '+f['ano'])
    if f.get('mes'):parts.append('Mês '+f['mes'])
    for key,label in [('inicio','De'),('fim','Até')]:
        if f.get(key):parts.append(label+' '+datetime.strptime(f[key],'%Y-%m-%d').strftime('%d/%m/%Y'))
    if f.get('local'):parts.append('Local: '+next((x['nome'] for x in data['options']['locations'] if x['id']==f['local']),'Selecionado'))
    return ' · '.join(parts)


def sheets(data):
    current=not data['base_complete'] and data.get('registry',{}).get('available')
    frequency=[]
    def subtotal(rows,label):
        m=sum(r['MUSICO'] for r in rows);o=sum(r['ORGANISTA'] for r in rows)
        validated=data['base_complete'] and any(r['musicos_base'] is not None for r in rows)
        bm=sum(r.get('cadastro_musicos') or 0 for r in rows) if current else sum(r['musicos_base'] or 0 for r in rows) if validated else None
        bo=sum(r.get('cadastro_organistas') or 0 for r in rows) if current else sum(r['organistas_base'] or 0 for r in rows) if validated else None
        pm=(m/bm if current else sum(r.get('_musicos_elegiveis',0) for r in rows)/bm) if bm else None
        po=(o/bo if current else sum(r.get('_organistas_elegiveis',0) for r in rows)/bo) if bo else None
        return [label,'SUBTOTAL' if label!='TOTAL GERAL' else '','',m,o,m+o,bm,bo,pm,po]
    city_rows=[]
    for r in data['rows']:
        if city_rows and city_rows[-1]['municipio']!=r['municipio']:
            frequency.append(subtotal(city_rows,city_rows[-1]['municipio']));city_rows=[]
        city_rows.append(r)
        bm=r.get('cadastro_musicos') if current else r.get('musicos_base')
        bo=r.get('cadastro_organistas') if current else r.get('organistas_base')
        pm=(r['MUSICO']/bm if bm else None) if current else (r['percentual_musicos']/100 if r['percentual_musicos'] is not None else None)
        po=(r['ORGANISTA']/bo if bo else None) if current else (r['percentual_organistas']/100 if r['percentual_organistas'] is not None else None)
        leader=(r['encarregado_nome']+'\n'+r.get('encarregado_situacao_label',r['encarregado'])) if r.get('encarregado_nome') else r['encarregado']
        frequency.append([r['municipio'],r['comum'],leader,r['MUSICO'],r['ORGANISTA'],r['MUSICO']+r['ORGANISTA'],
                          bm,bo,pm,po])
    if city_rows:frequency.append(subtotal(city_rows,city_rows[-1]['municipio']))
    if data['rows']:frequency.append(subtotal(data['rows'],'TOTAL GERAL'))
    base_label='Base SAM atual' if current else 'Base SAM do evento'
    frequency_headers=['Município','Comum','Encarregado Local','Músicos presentes','Organistas presentes','Total Geral',
        base_label+' · músicos',base_label+' · organistas','% músicos','% organistas']
    frequency_widths=[26,60,30,18,20,18,21,23,16,17]
    conductors=data.get('conductors',{})
    conductor_rows=[]
    for r in conductors.get('rows',[]):
        present_names=(' / '.join(r.get('nomes') or []) if r['situacao']=='PRESENTE' else '') or r.get('encarregado_nome') or 'Não identificado'
        status=r.get('presenca_lista_label', 'NJ')
        conductor_rows.append([datetime.strptime(str(r['data'])[:10],'%Y-%m-%d').strftime('%d/%m/%Y'),r['municipio'],
            r['comum']+'\nEncarregado: '+present_names,status,r.get('participacao','') if r.get('consta_lista') and r.get('locais_validos') else '',
            ' / '.join(r.get('locais_validos',[]))])
    reports=[('Frequência por comum',frequency_headers,frequency,frequency_widths),
            ('Encarregados por comum',['Data','Município de origem','Comum / encarregado','Presença','Participação','Local do ensaio'],
             conductor_rows,[16,26,70,22,30,42]),
            ('Encarregados por município',['Município','PRESENTE','JUSTIFICOU','NJ','SE','EF'],
             [[r['municipio']]+[sum(item.get('presenca_situacao')==status for item in conductors.get('rows',[]) if item['municipio']==r['municipio'])
                                for status in ['PRESENTE','JUSTIFICOU','NAO_JUSTIFICOU','SEM_ENCARREGADO','ENFERMO']]
              for r in conductors.get('municipalities',[])],[50,25,25,25,25,25]),
            ('Resumo por local',['Local','Músicos','Organistas','Ministério','Apoio','Irmandade','Total'],
             [[x['nome'],x['MUSICO'],x['ORGANISTA'],x['MINISTERIO'],x['APOIO'],x['IRMANDADE'],x['total']] for x in data['locations']], [44,19,19,19,19,19,19]),
            ('Instrumentos',['Músicos por instrumento / Organistas','Participantes presentes'],list(sorted(data['instruments'].items())),[95,45]),
            ('Famílias',['Família / naipe','Presentes'],list(sorted(data['families'].items())),[95,45]),
            ('Cargos',['Cargo (dimensão não aditiva)','Presentes'],list(sorted(data['roles'].items())),[95,45])]
    if data.get('report')=='encarregados':return reports[1:2]
    if data.get('report')=='comuns':return reports[:1]
    if data.get('report')=='composicao':return reports[-3:]
    return reports


def executive_story(data,width,style,head,context):
    """Síntese de uma página; tabelas nominais ficam em exportações próprias."""
    fmt=lambda v:f'{int(v or 0):,}'.replace(',','.')
    cell=lambda v:Paragraph(escape(str(v)),style)
    heading=ParagraphStyle('executive-heading',parent=style,fontSize=10,leading=13,fontName='Helvetica-Bold',textColor=colors.HexColor('#1e4b7a'))
    t=data['totals'];musical=t.get('MUSICO',0)+t.get('ORGANISTA',0)
    kpis=[]
    for label,value in [('Músicos',t.get('MUSICO',0)),('Organistas',t.get('ORGANISTA',0)),('Total musical',musical),('Público geral',t.get('total',0))]:
        kpis.append(Paragraph(escape(label)+f'<br/><font size="22" color="#1e4b7a"><b>{fmt(value)}</b></font>',ParagraphStyle('kpi',parent=style,fontSize=9,leading=27)))
    cards=Table([kpis],colWidths=[width/4]*4)
    cards.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#eaf2f8')),('BOX',(0,0),(-1,-1),.4,colors.HexColor('#d8dee5')),('LEFTPADDING',(0,0),(-1,-1),13),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
    rows=data.get('regional_overview') if data['filters'].get('origem','regional')=='regional' else data['municipalities']
    rows=rows or data['municipalities']
    visible=list(rows[:7])
    if len(rows)>7:
        visible.append({'municipio':'Outros municípios',**{key:sum(r.get(key,0) for r in rows[7:]) for key in ['MUSICO','ORGANISTA','total','encarregados']}})
    body=[[Paragraph(x,head) for x in ['Município de origem','Músicos','Organistas','Total musical','Encarregados locais']]]
    for r in visible:body.append([cell(r['municipio']),cell(fmt(r['MUSICO'])),cell(fmt(r['ORGANISTA'])),cell(fmt(r['MUSICO']+r['ORGANISTA'])),cell(fmt(r.get('encarregados',0)))])
    body.append([cell('TOTAL DO RECORTE'),cell(fmt(t.get('MUSICO'))),cell(fmt(t.get('ORGANISTA'))),cell(fmt(musical)),cell(fmt(t.get('encarregados')))])
    table=Table(body,colWidths=[width*.4]+[width*.15]*4,repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#1e4b7a')),('ROWBACKGROUNDS',(0,1),(-1,-2),[colors.white,colors.HexColor('#f4f6f8')]),('BACKGROUND',(0,-1),(-1,-1),colors.HexColor('#eaf2f8')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#d8dee5')),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('VALIGN',(0,0),(-1,-1),'MIDDLE')]))
    def bars(title,items,tone):
        drawing=Drawing(width/2-10,110);drawing.add(String(0,99,title,fontName='Helvetica-Bold',fontSize=10,fillColor=colors.HexColor('#1e4b7a')))
        maximum=max([v for _,v in items]+[1]);bar_width=width/2-185
        for i,(label,value) in enumerate(items):
            y=77-i*21
            drawing.add(String(0,y,str(label),fontName='Helvetica',fontSize=8,fillColor=colors.HexColor('#36546b')))
            drawing.add(Rect(110,y-1,bar_width,7,fillColor=colors.HexColor('#edf2f6'),strokeColor=None))
            drawing.add(Rect(110,y-1,bar_width*value/maximum,7,fillColor=colors.HexColor(tone),strokeColor=None))
            drawing.add(String(width/2-18,y,fmt(value),textAnchor='end',fontSize=9,fontName='Helvetica-Bold',fillColor=colors.HexColor('#1e4b7a')))
        return drawing
    families=data.get('musician_families',{})
    family_items=[(label,families.get(key,0)) for key,label in [('CORDAS','Cordas'),('MADEIRAS','Madeiras'),('METAIS','Metais')]]
    others=t.get('MUSICO',0)-sum(value for _,value in family_items)
    if others:family_items.append(('Outros / não informado',others))
    composition=bars('Composição dos músicos',family_items,'#2f8de4')
    staffing=bars('Participação administrativa',[('Encarregados locais',t.get('encarregados',0)),('Ministério',t.get('ministerio_cargos',0)),('Apoio',t.get('APOIO',0)),('Irmandade',t.get('IRMANDADE',0))],'#12a594')
    charts=Table([[composition,staffing]],colWidths=[width/2]*2)
    charts.setStyle(TableStyle([('LEFTPADDING',(0,0),(-1,-1),0),('TOPPADDING',(0,0),(-1,-1),0),('VALIGN',(0,0),(-1,-1),'TOP')]))
    comparable=data.get('evolution',{}).get('comparable',False)
    message='Evolução: consulte a série do município no painel; apenas encontros equivalentes são comparados.' if comparable else 'Evolução ainda não mensurável: faltam dois encontros equivalentes no recorte.'
    if comparable:
        history=data['evolution'];points=history['points']
        def change(field):
            value=history['metrics'][field]['change']
            return 'sem base anterior' if value is None else f'{value:+.1f}%'.replace('.',',')
        date=lambda v:datetime.strptime(str(v)[:10],'%Y-%m-%d').strftime('%d/%m/%Y')
        message=f"Comparação do mesmo circuito ({date(points[-2]['date'])} a {date(points[-1]['date'])}): músicos {change('musicians')}; organistas {change('organists')}."
    return [Paragraph(escape(scope_label(data)),context),Spacer(1,12),cards,Spacer(1,13),Paragraph('Participação por município de origem',heading),Spacer(1,6),table,Spacer(1,14),charts,
            Paragraph(escape(message),style),Spacer(1,5),Paragraph('Encarregados locais e ministério são cargos e podem estar incluídos no total musical. Não somar novamente ao público geral.',style)]


def build_presence_model_pdf(data, actor='Secretaria Musical', now=None):
    if all(event.get('tipo') == 'ENSAIO_REGIONAL' for event in data.get('events', [])):
        from .music_events_annual_reports import build_model_pdf
        return build_model_pdf(data, actor, now)
    return _build_legacy_presence_pdf(data)


def _build_legacy_presence_pdf(data):
    """Reproduz a grade do relatório unificado usado pela Secretaria Musical."""
    stream=io.BytesIO();page_width,page_height=landscape(A4)
    doc=SimpleDocTemplate(stream,pagesize=landscape(A4),leftMargin=30,rightMargin=30,topMargin=26,bottomMargin=24)
    normal=ParagraphStyle('legacy-cell',fontName='Helvetica',fontSize=7.2,leading=8.2,textColor=colors.black)
    center=ParagraphStyle('legacy-center',parent=normal,alignment=1)
    bold=ParagraphStyle('legacy-bold',parent=center,fontName='Helvetica-Bold')
    event=data.get('events',[{}])[0] if data.get('events') else {}
    kind='ENSAIO REGIONAL' if event.get('tipo')=='ENSAIO_REGIONAL' else 'REUNIÃO TÉCNICA'
    date=datetime.strptime(str(event.get('data_evento'))[:10],'%Y-%m-%d').strftime('%d/%m/%Y') if event.get('data_evento') else ''
    places=[x['nome'] for x in data.get('locations',[])][:4]
    place_title=' E '.join(dict.fromkeys(str(x).upper() for x in places)) or 'REGIONAL ITAPEVI'
    title=Paragraph(f'<b>RELATÓRIO UNIFICADO {kind}</b><br/><b>{escape(place_title)}</b><br/><b>{date}</b>',
        ParagraphStyle('legacy-title',fontName='Helvetica-Bold',fontSize=13,leading=16,alignment=1))
    logo_path=Path(__file__).parents[2]/'static'/'img'/'logo'/'ccb.png'
    logo=Image(str(logo_path),width=58,height=58) if logo_path.exists() else Paragraph('<b>CCB</b>',bold)
    header=Table([[logo,title]],colWidths=[165,page_width-225],rowHeights=[66])
    header.setStyle(TableStyle([('BOX',(0,0),(-1,-1),.7,colors.black),('LINEAFTER',(0,0),(0,0),.7,colors.black),
        ('ALIGN',(0,0),(0,0),'CENTER'),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),4),
        ('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),3),('BOTTOMPADDING',(0,0),(-1,-1),3)]))
    current=not data['base_complete'] and data.get('registry',{}).get('available')
    def values(row):
        bm=row.get('cadastro_musicos') if current else row.get('musicos_base')
        bo=row.get('cadastro_organistas') if current else row.get('organistas_base')
        pm=(row['MUSICO']/bm if bm else None) if current else (row.get('percentual_musicos')/100 if row.get('percentual_musicos') is not None else None)
        po=(row['ORGANISTA']/bo if bo else None) if current else (row.get('percentual_organistas')/100 if row.get('percentual_organistas') is not None else None)
        common=re.sub(r'^BR-\d+-\d+\s*-\s*','',str(row['comum']),flags=re.I)
        leader=row.get('encarregado') or row.get('encarregado_situacao_label') or 'Não informado'
        return [common,leader,row['MUSICO'],row['ORGANISTA'],row['MUSICO']+row['ORGANISTA'],bm,bo,pm,po]
    rows=[(r['municipio'],values(r)) for r in data.get('rows',[])]
    # O cabeçalho institucional ocupa 66 pt; 24 linhas preservam a tabela inteira
    # dentro da área útil do A4 paisagem em todas as páginas.
    chunks=[rows[i:i+24] for i in range(0,len(rows),24)] or [[]]
    palette={'CAUCAIA DO ALTO':'#F8DC87','COTIA':'#BCD7ED','ITAPEVI':'#BFE3DA','JANDIRA':'#DDD0EA',
             'PIRAPORA DO BOM JESUS':'#F4C7A8','SANTANA DE PARNAÍBA':'#F6D8E2','VARGEM GRANDE PAULISTA':'#D9E6AE'}
    story=[]
    for page_index,chunk in enumerate(chunks):
        if page_index:story.append(PageBreak())
        story.extend([header,Spacer(1,0)])
        table_rows=[[Paragraph('',bold),Paragraph('',bold),Paragraph('PRESENTES NO ENSAIO',bold),'','','',Paragraph('NÚMEROS POR LOCALIDADE',bold),'','',''],
                    ['',Paragraph('Comum',center),Paragraph('Encarregado Local',center),Paragraph('Músicos',center),Paragraph('Organistas',center),Paragraph('Total Geral',center),Paragraph('Músicos',center),Paragraph('Organistas',center),Paragraph('% Músicos',center),Paragraph('% Organistas',center)]]
        for city,row in chunk:
            table_rows.append(['',Paragraph(escape(str(row[0])),normal),Paragraph(escape(str(row[1])),bold),
                Paragraph(str(row[2]),center),Paragraph(str(row[3]),center),Paragraph(f'<b>{row[4]}</b>',center),
                Paragraph('—' if row[5] is None else str(row[5]),center),Paragraph('—' if row[6] is None else str(row[6]),center),
                Paragraph('—' if row[7] is None else f'{row[7]*100:.0f}%',center),Paragraph('—' if row[8] is None else f'{row[8]*100:.0f}%',center)])
        widths=[18,162,118,42,48,52,61,61,57,64]
        table=Table(table_rows,colWidths=widths,rowHeights=[14,15]+[18]*len(chunk),repeatRows=2)
        styles=[('SPAN',(0,0),(1,0)),('SPAN',(2,0),(5,0)),('SPAN',(6,0),(9,0)),
            ('BACKGROUND',(2,0),(5,0),colors.HexColor('#c9dff0')),('BACKGROUND',(6,0),(9,0),colors.HexColor('#cfe6bd')),
            ('GRID',(0,0),(-1,-1),.45,colors.black),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
            ('LEFTPADDING',(0,0),(-1,-1),2),('RIGHTPADDING',(0,0),(-1,-1),2),
            ('TOPPADDING',(0,0),(-1,-1),1),('BOTTOMPADDING',(0,0),(-1,-1),1)]
        start=0
        while start<len(chunk):
            city=chunk[start][0];end=start
            while end+1<len(chunk) and chunk[end+1][0]==city:end+=1
            r0=start+2;r1=end+2
            styles.extend([('SPAN',(0,r0),(0,r1)),('BACKGROUND',(0,r0),(0,r1),colors.HexColor(palette.get(city,'#E5E7EB'))),
                ('ALIGN',(0,r0),(0,r1),'CENTER')])
            label=vertical_label(city.replace(' DO ALTO','').replace(' DO BOM JESUS',''),max(10,18*(end-start+1)-4))
            table_rows[r0][0]=label;table._cellvalues[r0][0]=label
            start=end+1
        for i,(_,row) in enumerate(chunk,2):
            styles.append(('BACKGROUND',(1,i),(7,i),colors.HexColor('#d9d9d9') if i%2==0 else colors.white))
            for col,value in [(8,row[7]),(9,row[8])]:
                if value is not None:
                    shade=max(0,min(1,value));styles.append(('BACKGROUND',(col,i),(col,i),colors.Color(0.88-.30*shade,0.96-.18*shade,0.91-.22*shade)))
        table.setStyle(TableStyle(styles));story.append(table)
    doc.build(story)
    return stream.getvalue()


def build_excel(data, actor, now):
    book=Workbook();book.remove(book.active)
    for name,headers,rows,widths in sheets(data):
        ws=book.create_sheet(name);end=get_column_letter(len(headers))
        for index,text in enumerate(['CONGREGAÇÃO CRISTÃ NO BRASIL','Regional Itapevi - São Paulo',
                                     'MÚSICA · ENSAIOS REGIONAIS E RTs',name],1):
            ws.merge_cells(start_row=index,start_column=1,end_row=index,end_column=len(headers))
            cell=ws.cell(index,1,text);cell.font=Font(size=15 if index==1 else 10,bold=True,color='FFFFFF' if index<3 else NAVY)
            cell.fill=PatternFill('solid',fgColor=NAVY if index<3 else PALE);cell.alignment=Alignment(horizontal='center')
        for index,text in [(5,f'Emissão: {now:%d/%m/%Y %H:%M} · Responsável: {actor}'),(6,'Recorte: '+scope_label(data)),
                           (7,report_note(data))]:
            ws.merge_cells(start_row=index,start_column=1,end_row=index,end_column=len(headers));ws.cell(index,1,text)
            ws.cell(index,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[index].height=30 if index==6 else 23
        for col,label in enumerate(headers,1):
            cell=ws.cell(9,col,label);cell.fill=PatternFill('solid',fgColor=NAVY);cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(horizontal='center',wrap_text=True)
        ws.row_dimensions[9].height=30
        for index,row in enumerate(rows,10):
            for col,value in enumerate(row,1):
                cell=ws.cell(index,col,value)
                if isinstance(value,str):cell.data_type='s' # nunca executar fórmulas oriundas de dados externos
                cell.font=Font(size=10,color='243746');cell.alignment=Alignment(vertical='center',wrap_text=True)
                cell.border=Border(bottom=Side(style='hair',color='D8DEE5'))
                if index%2==0:cell.fill=PatternFill('solid',fgColor='F4F6F8')
                if name=='Frequência por comum' and col in {9,10}:cell.number_format='0.00%'
                if name=='Frequência por comum' and (row[1]=='SUBTOTAL' or row[0]=='TOTAL GERAL'):
                    cell.fill=PatternFill('solid',fgColor=PALE);cell.font=Font(bold=True,color=NAVY)
                if name=='Encarregados por comum' and col==4:
                    color='148052' if value in {'PRESENTE','JUSTIFICOU'} else 'C92D4D' if value=='NÃO JUSTIFICOU' else NAVY
                    cell.font=Font(bold=True,color=color)
                if name=='Encarregados por comum' and col==5:
                    cell.font=Font(bold=True,color='C26A00' if value=='FORA DO MUNICÍPIO' else '148052')
            ws.row_dimensions[index].height=max(30,13*max(math.ceil(len(str(value or ''))/max(1,widths[i]-3)) for i,value in enumerate(row)))
        for index,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(index)].width=width
        ws.freeze_panes='A10';ws.auto_filter.ref=f'A9:{end}{max(9,ws.max_row)}';ws.sheet_view.showGridLines=False
        ws.print_title_rows='1:9';ws.print_options.horizontalCentered=True;ws.print_area=f'A1:{end}{ws.max_row}'
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.oddFooter.center.text='Página &P de &N';ws.oddFooter.right.text='Regional Itapevi'
    stream=io.BytesIO();book.save(stream);return stream.getvalue()


class NumberedCanvas(Canvas):
    def __init__(self,*args,**kwargs):super().__init__(*args,**kwargs);self.pages=[]
    def showPage(self):self.pages.append(dict(self.__dict__));self._startPage()
    def save(self):
        for page in self.pages:
            self.__dict__.update(page)
            self.setFont('Helvetica',7);self.setFillColor(colors.HexColor('#6b7280'))
            self.drawRightString(landscape(A4)[0]-28,landscape(A4)[1]-24,f'Página {self._pageNumber} de {len(self.pages)}')
            self.drawString(28,18,'Música · Ensaios Regionais e RTs · Regional Itapevi')
            self.drawRightString(landscape(A4)[0]-28,18,f'Página {self._pageNumber} de {len(self.pages)}')
            Canvas.showPage(self)
        Canvas.save(self)


def build_pdf(data,actor,now):
    if data.get('report')=='comuns':
        return build_presence_model_pdf(data,actor,now)
    stream=io.BytesIO();page_width,page_height=landscape(A4);width=page_width-56
    doc=SimpleDocTemplate(stream,pagesize=landscape(A4),leftMargin=28,rightMargin=28,topMargin=100,bottomMargin=34)
    style=ParagraphStyle('cell',fontName='Helvetica',fontSize=7,leading=9,textColor=colors.HexColor('#243746'))
    header_style=ParagraphStyle('head',parent=style,textColor=colors.white,alignment=1,fontName='Helvetica-Bold')
    context=ParagraphStyle('context',parent=style,backColor=colors.HexColor('#eaf2f8'),borderPadding=6)
    def para(value,head=False):
        text=str(value if value is not None else '—');cell_style=header_style if head else style
        if not head and text=='PRESENTE':
            cell_style=ParagraphStyle('present',parent=style,fontName='Helvetica-Bold',textColor=colors.HexColor('#148052'))
        if not head and text in {'JUSTIFICOU','NÃO JUSTIFICOU','NÃO INFORMADO','SEM ENCARREGADO','ENFERMO'}:
            tone='#148052' if text=='JUSTIFICOU' else '#c92d4d' if text=='NÃO JUSTIFICOU' else '#526b7a'
            cell_style=ParagraphStyle('conductor-state',parent=style,fontName='Helvetica-Bold',textColor=colors.HexColor(tone))
        if not head and text in {'NO MUNICÍPIO','FORA DO MUNICÍPIO'}:
            cell_style=ParagraphStyle('attendance-place',parent=style,fontName='Helvetica-Bold',textColor=colors.HexColor('#c26a00' if text.startswith('FORA') else '#148052'))
        if not head and (text.startswith('NJ ·') or text.startswith('J ·')):
            cell_style=ParagraphStyle('status',parent=style,fontName='Helvetica-Bold',textColor=colors.HexColor('#c92d4d' if text.startswith('NJ ·') else '#148052'))
        return Paragraph(escape(text).replace('\n','<br/>'),cell_style)
    def header(canvas,document):
        canvas.saveState();center=page_width/2
        canvas.setFillColor(colors.black);canvas.setFont('Helvetica-Bold',15);canvas.drawCentredString(center,page_height-29,'CONGREGAÇÃO CRISTÃ NO BRASIL')
        canvas.setFont('Helvetica',9);canvas.drawCentredString(center,page_height-43,'Regional Itapevi - São Paulo')
        canvas.setFillColor(colors.HexColor('#1e4b7a'));canvas.setFont('Helvetica-Bold',12);canvas.drawCentredString(center,page_height-60,'ENSAIOS REGIONAIS E RTs')
        title={'executivo':'Resumo executivo de presença','encarregados':'Presença dos Encarregados Locais','comuns':'Frequência musical por comum','composicao':'Composição musical e cargos'}.get(data.get('report'),'Relatório administrativo de frequência')
        canvas.setFont('Helvetica',9);canvas.drawCentredString(center,page_height-74,title)
        f=data['filters'];identity=data.get('report_identity',{})
        recorte=', '.join(f.get('municipios',[])) or ('Regional Itapevi' if f.get('origem','regional')=='regional' else 'Visitantes' if f.get('origem')=='visitantes' else 'Regional e visitantes')
        lines=[f'Emissão: {now:%d/%m/%Y %H:%M}','Recorte: '+recorte,'Comuns: '+(', '.join(f.get('comuns',[])) or 'Todas'),'Responsável: '+actor]
        if identity.get('scope'):lines.append('Escopo: '+identity['scope'])
        if identity.get('common'):lines.append(identity['common'])
        meta=Paragraph('<br/>'.join(escape(x) for x in lines),ParagraphStyle('meta',parent=style,alignment=2,fontSize=7,leading=9))
        _,height=meta.wrap(190,100);meta.drawOn(canvas,page_width-218,page_height-34-height)
        canvas.restoreState()
    story=[Paragraph(escape('RECORTE DO RELATÓRIO · '+scope_label(data)),context),Spacer(1,12),
           Paragraph(escape(report_note(data)),style),Spacer(1,10)]
    if data.get('report')=='executivo':
        doc.build(executive_story(data,width,style,header_style,context),onFirstPage=header,onLaterPages=header,canvasmaker=NumberedCanvas)
        return stream.getvalue()
    for name,headers,rows,_ in sheets(data):
        if name=='Frequência por comum':
            ratios=[.09,.275,.14,.075,.08,.075,.09,.095,.04,.04]
            # Razões normalizadas garantem o uso exato da largura útil.
            widths=[width*x/sum(ratios) for x in ratios]
        elif name=='Encarregados por comum':widths=[width*x for x in [.08,.15,.34,.11,.15,.17]]
        elif len(headers)==2:widths=[width*.8,width*.2]
        else:widths=[width*.28]+[width*.72/(len(headers)-1)]*(len(headers)-1)
        title_style=ParagraphStyle('tabletitle',parent=style,fontName='Helvetica-Bold',fontSize=10,leading=14,textColor=colors.HexColor('#1e4b7a'))
        body=[[Paragraph(escape(name),title_style)]+[para('')]*(len(headers)-1)]
        if name=='Frequência por comum':
            body.append([para('')]*3+[para('PRESENTES NO ENSAIO',True)]+[para('')]*2+
                        [para('NÚMEROS POR LOCALIDADE',True)]+[para('')]*3)
        body.append([para(h,True) for h in headers])
        for row in rows:
            body.append([para(f'{v*100:.2f}%' if name=='Frequência por comum' and i in {8,9} and v is not None else v) for i,v in enumerate(row)])
        if not rows:body.append([para('Sem registros neste recorte')]+[para('')]*(len(headers)-1))
        header_rows=3 if name=='Frequência por comum' else 2
        table=Table(body,colWidths=widths,repeatRows=header_rows,hAlign='LEFT')
        table.setStyle(TableStyle([('SPAN',(0,0),(-1,0)),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eaf2f8')),
            ('BACKGROUND',(0,header_rows-1),(-1,header_rows-1),colors.HexColor('#1e4b7a')),
            ('ROWBACKGROUNDS',(0,header_rows),(-1,-1),[colors.white,colors.HexColor('#f4f6f8')]),
            ('GRID',(0,0),(-1,-1),.35,colors.HexColor('#d8dee5')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
            ('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3),
            ('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
        if name=='Frequência por comum':
            table.setStyle(TableStyle([('SPAN',(0,1),(2,1)),('SPAN',(3,1),(5,1)),('SPAN',(6,1),(9,1)),
                ('BACKGROUND',(0,1),(-1,1),colors.HexColor('#1e4b7a')),('ALIGN',(0,1),(-1,1),'CENTER')]))
            for index,row in enumerate(rows,3):
                if row[1]=='SUBTOTAL' or row[0]=='TOTAL GERAL':
                    table.setStyle(TableStyle([('BACKGROUND',(0,index),(-1,index),colors.HexColor('#eaf2f8'))]))
        if name!='Frequência por comum' and len(rows)<=15:
            story.append(KeepTogether([table,Spacer(1,16)]))
        else:story.extend([table,Spacer(1,16)])
    doc.build(story,onFirstPage=header,onLaterPages=header,canvasmaker=NumberedCanvas)
    return stream.getvalue()


def export_response(data,request,file_format):
    now=datetime.now(ZoneInfo('America/Sao_Paulo'));p=request.session.get('user_profile') or {}
    actor=p.get('full_name') or p.get('nome') or p.get('name') or p.get('email') or 'Usuário'
    from .access_control import user_scope
    scope=user_scope(request)
    data['report_identity']={'scope':' · '.join(filter(None,[{'global':'Global','regional':'Regional','municipal':'Municipal','local':'Local'}.get(scope['level']),scope.get('municipio')])),
                             'common':scope.get('comum') or ''}
    requested=request.GET.get('relatorio')
    data['report']=requested if requested in {'executivo','encarregados','comuns','composicao'} else ('comuns' if file_format=='pdf' else 'completo')
    content=build_excel(data,actor,now) if file_format=='excel' else build_pdf(data,actor,now)
    ext='xlsx' if file_format=='excel' else 'pdf'
    response=HttpResponse(content,content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet' if ext=='xlsx' else 'application/pdf')
    prefix={'executivo':'Resumo_Executivo_Ensaios_RTs','encarregados':'Encarregados_Locais','comuns':'Participacao_por_Comum','composicao':'Composicao_Musical'}.get(data.get('report'),'Ensaios_RTs')
    response['Content-Disposition']=f'attachment; filename="{prefix}_{now:%d-%m-%Y_%H-%M}.{ext}"'
    events = data.get('events', [])
    if file_format == 'pdf' and data['report'] == 'comuns' and events and all(e.get('tipo') == 'ENSAIO_REGIONAL' for e in events):
        from .music_events_filenames import presence_filename, download_header
        cities = data.get('filters', {}).get('municipios') or sorted({r['municipio'] for r in data.get('rows', [])})
        response['Content-Disposition'] = download_header(presence_filename(cities, [e['data_evento'] for e in events]))
    response['Cache-Control']='private, no-store'
    return response
