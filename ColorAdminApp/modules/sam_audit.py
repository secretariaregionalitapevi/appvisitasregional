"""Central de investigacao dos eventos preservados do SAM original."""
import io
import re
import unicodedata
from collections import Counter
from datetime import datetime, timedelta, timezone

import requests
from django.conf import settings
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.template.response import TemplateResponse
from django.utils import timezone as django_timezone
from django.views.decorators.csrf import ensure_csrf_cookie
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .access_control import common_catalog, service_headers, user_scope
from .gem import can_open_module
from .sam_audit_enrichment import enrich_events, search_event_reference_ids
from .sam_audit_parser import parse_references


def _url(table):
    return f"{settings.SUPABASE_URL}/rest/v1/{table}"


def _allowed(request):
    return can_open_module(request) and user_scope(request)["level"] in {"global", "regional"}


@ensure_csrf_cookie
def page(request):
    if not _allowed(request):
        return render(request, "pages/403.html", {"message": "O Historico SAM exige perfil regional ou global."}, status=403)
    profile = request.session.get("user_profile") or {}
    actor = profile.get("full_name") or profile.get("nome") or profile.get("name") or profile.get("email") or "Usuario"
    return TemplateResponse(request, "pages/sam_audit.html", {"report_user": actor})


def _filters(request):
    params = {}
    search = str(request.GET.get("q") or "").strip().replace("*", "")[:120]
    if search:
        safe = search.replace(",", " ").replace("(", " ").replace(")", " ")
        clauses = [f"actor_name_snapshot.ilike.*{safe}*", f"action_raw.ilike.*{safe}*",
                   f"congregation_snapshot.ilike.*{safe}*", f"target_name_snapshot.ilike.*{safe}*"]
        clauses.extend(f"action_raw.ilike.*{reference_id}*" for reference_id in search_event_reference_ids(safe))
        params["or"] = "(" + ",".join(clauses) + ")"
    for query, column in (("actor", "actor_name_snapshot"), ("action", "action_type"),
                          ("congregation", "congregation_snapshot"), ("target", "target_type")):
        value = str(request.GET.get(query) or "").strip()
        if value:
            params[column] = f"eq.{value}"
    start, end = request.GET.get("start"), request.GET.get("end")
    date_filters = []
    if start:
        date_filters.append(f"occurred_at.gte.{start}T00:00:00-03:00")
    if end:
        date_filters.append(f"occurred_at.lte.{end}T23:59:59-03:00")
    if date_filters:
        params["and"] = "(" + ",".join(date_filters) + ")"
    return params


def api_events(request):
    if not _allowed(request):
        return JsonResponse({"error": "O Historico SAM exige perfil regional ou global."}, status=403)
    try:
        page_number = max(1, int(request.GET.get("page") or 1))
        page_size = min(50, max(20, int(request.GET.get("page_size") or 50)))
        start_index = (page_number - 1) * page_size
        location_params = _event_location_query_params(request)
        if location_params is not None and not location_params:
            rows, total = [], 0
        else:
            query_params = {"select": "*", "order": "occurred_at.desc,id.desc", "offset": start_index,
                            "limit": page_size, **_filters(request), **(location_params or {})}
            events_response = requests.get(
                _url("sam_audit_events"), headers=service_headers("count=exact"),
                params=query_params, timeout=(2, 8),
            )
            events_response.raise_for_status()
            rows = [_hydrate_event_refs(row) for row in events_response.json()]
            content_range = events_response.headers.get("Content-Range", "")
            try:
                total = int(content_range.rsplit("/", 1)[1])
            except (IndexError, ValueError):
                total = start_index + len(rows)

        counts = Counter(row.get("action_type") or "outro" for row in rows)
        cutoff = datetime.now(timezone.utc) - timedelta(hours=24)
        recent = 0
        for row in rows:
            try:
                if datetime.fromisoformat(str(row.get("occurred_at")).replace("Z", "+00:00")) >= cutoff:
                    recent += 1
            except (TypeError, ValueError):
                pass
        return JsonResponse({"events": rows, "total": total, "page": page_number, "page_size": page_size,
                             "summary": {"visible": len(rows), "last_24h_visible": recent,
                                         "actors_visible": len({r.get('actor_name_snapshot') for r in rows}),
                                         "action_counts_visible": dict(counts)}, "sync": {}})
    except (requests.RequestException, ValueError):
        return JsonResponse({"error": "Nao foi possivel consultar o Historico SAM. Aplique a migration 028 e sincronize os logs."}, status=503)


def api_event_detail(request, event_id):
    if not _allowed(request):
        return JsonResponse({"error": "O Historico SAM exige perfil regional ou global."}, status=403)
    try:
        response = requests.get(
            _url("sam_audit_events"), headers=service_headers(),
            params={"select": "*", "id": f"eq.{event_id}", "limit": 1}, timeout=(3.05, 10),
        )
        response.raise_for_status()
        rows = response.json()
        if not rows:
            return JsonResponse({"error": "Evento não encontrado."}, status=404)
        rows = [_associate_adjacent_student(_hydrate_event_refs(rows[0]))]
        try:
            rows = enrich_events(rows)
        except requests.RequestException:
            pass
        return JsonResponse({"event": rows[0]})
    except requests.RequestException:
        return JsonResponse({"error": "Não foi possível consultar os detalhes deste evento."}, status=503)


def api_sync_status(request):
    if not _allowed(request):
        return JsonResponse({"error": "O Historico SAM exige perfil regional ou global."}, status=403)
    try:
        response = requests.get(
            _url("sam_audit_sync_cursor"), headers=service_headers(),
            params={"select": "*", "source": "eq.sam_system_logs", "limit": 1}, timeout=(1.5, 2.5),
        )
        response.raise_for_status()
        return JsonResponse({"sync": (response.json() or [{}])[0]})
    except (requests.RequestException, ValueError):
        return JsonResponse({"sync": {}, "available": False})


def _hydrate_event_refs(event):
    event = dict(event)
    refs, target_type, target_source_id = parse_references(event.get("action_raw"))
    refs.update(event.get("parsed_refs") or {})
    event["parsed_refs"] = refs
    event["target_type"] = event.get("target_type") or target_type
    event["target_source_id"] = event.get("target_source_id") or target_source_id
    return event


def _associate_adjacent_student(event, window_seconds=5):
    refs = event.get("parsed_refs") or {}
    if event.get("target_type") != "metodo" or refs.get("student_id"):
        return event
    try:
        occurred_at = datetime.fromisoformat(str(event.get("occurred_at")).replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return event
    actor = event.get("actor_name_snapshot")
    common = event.get("congregation_snapshot")
    if not actor or not common:
        return event

    lower_bound = occurred_at - timedelta(seconds=window_seconds)
    try:
        response = requests.get(
            _url("sam_audit_events"), headers=service_headers(),
            params={"select": "id,occurred_at,action_raw,parsed_refs", "actor_name_snapshot": f"eq.{actor}",
                    "congregation_snapshot": f"eq.{common}", "occurred_at": f"gte.{lower_bound.isoformat()}",
                    "order": "occurred_at.asc", "limit": 50},
            timeout=(2, 5),
        )
        response.raise_for_status()
        candidates = []
        for row in response.json():
            if row.get("id") == event.get("id"):
                continue
            try:
                candidate_time = datetime.fromisoformat(str(row.get("occurred_at")).replace("Z", "+00:00"))
            except (TypeError, ValueError):
                continue
            delta = abs((candidate_time - occurred_at).total_seconds())
            if delta > window_seconds:
                continue
            candidate = _hydrate_event_refs(row)
            student_id = (candidate.get("parsed_refs") or {}).get("student_id")
            if student_id:
                candidates.append((delta, str(student_id), candidate))
    except requests.RequestException:
        return event

    student_ids = {student_id for _, student_id, _ in candidates}
    if len(student_ids) != 1:
        return event
    _, student_id, source_event = min(candidates, key=lambda item: item[0])
    event["parsed_refs"]["student_id"] = int(student_id)
    event["resolution_context"] = {
        "association": "adjacent_log_reference",
        "confidence": "inferred",
        "source_event_id": source_event.get("id"),
        "time_delta_seconds": min(candidates, key=lambda item: item[0])[0],
    }
    return event


def _all_filtered_events(request, maximum=10000):
    rows = []
    filters = _filters(request)
    for offset in range(0, maximum, 1000):
        response = requests.get(_url("sam_audit_events"), headers=service_headers(),
                                params={"select": "*", "order": "occurred_at.desc,id.desc",
                                        "offset": offset, "limit": 1000, **filters}, timeout=(3.05, 45))
        response.raise_for_status()
        batch = response.json()
        for start in range(0, len(batch), 100):
            rows.extend(enrich_events(batch[start:start + 100]))
        if len(batch) < 1000:
            break
    return _filter_by_location(request, rows)


def _scope_label(request):
    parts = []
    labels = (("q", "Busca"), ("action", "Acao"), ("actor", "Operador"),
              ("congregation", "Localidade"), ("start", "De"), ("end", "Ate"))
    for key, label in labels:
        value = str(request.GET.get(key) or "").strip()
        if value:
            parts.append(f"{label}: {value}")
    return " | ".join(parts) or "Todos os eventos preservados"


def _event_report_row(event):
    resolution = event.get("resolution") or {}
    student, lesson, group = resolution.get("student") or {}, resolution.get("class") or {}, resolution.get("group") or {}
    references = ", ".join(f"{key}={value}" for key, value in (event.get("parsed_refs") or {}).items())
    members = "; ".join(member.get("name") or "" for member in group.get("members") or [])
    return [event.get("occurred_at"), event.get("actor_name_snapshot"), event.get("congregation_snapshot"),
            event.get("action_type"), event.get("action_raw"), resolution.get("primary_name"),
            student.get("name"), group.get("name") or lesson.get("name"), group.get("course") or lesson.get("course"),
            lesson.get("source_id"), lesson.get("date"), lesson.get("instructor"), group.get("start_date"),
            group.get("end_date"), members, event.get("device"), event.get("platform"), event.get("ip_address"), references]


def report_data(request):
    if not _allowed(request):
        return JsonResponse({"error": "O Historico SAM exige perfil regional ou global."}, status=403)
    try:
        rows = _all_filtered_events(request)
        profile = request.session.get("user_profile") or {}
        actor = profile.get("full_name") or profile.get("nome") or profile.get("name") or profile.get("email") or "Usuario"
        return JsonResponse({"items": rows, "count": len(rows), "scope": _scope_label(request), "report_user": actor})
    except requests.RequestException:
        return JsonResponse({"error": "Nao foi possivel preparar o relatorio do Historico SAM."}, status=503)


def export_excel(request):
    if not _allowed(request):
        return JsonResponse({"error": "O Historico SAM exige perfil regional ou global."}, status=403)
    try:
        rows = _all_filtered_events(request)
    except requests.RequestException:
        return JsonResponse({"error": "Nao foi possivel preparar a planilha do Historico SAM."}, status=503)
    profile = request.session.get("user_profile") or {}
    actor = profile.get("full_name") or profile.get("nome") or profile.get("name") or profile.get("email") or "Usuario"
    now, scope = datetime.now(), _scope_label(request)
    workbook = Workbook(); sheet = workbook.active; sheet.title = "HISTORICO SAM"
    headers = ["DATA E HORA", "OPERADOR", "LOCALIDADE", "TIPO", "ACAO ORIGINAL", "ENTIDADE IDENTIFICADA",
               "ALUNO", "TURMA", "CURSO", "ID AULA SAM", "DATA DA AULA", "INSTRUTOR", "ABERTURA DA TURMA",
               "FECHAMENTO DA TURMA", "MEMBROS DA TURMA", "DISPOSITIVO", "PLATAFORMA", "IP", "REFERENCIAS SAM"]
    last = get_column_letter(len(headers)); navy, pale = "1E4B7A", "EAF2F8"; thin = Side(style="thin", color="D8DEE5")
    for row_number, (text, size, dark) in enumerate((("CONGREGAÇÃO CRISTÃ NO BRASIL", 15, True),
                                                       ("Regional Itapevi - São Paulo", 10, True),
                                                       ("GRUPO DE ESTUDOS MUSICAIS · SAM", 12, False),
                                                       ("Relat\u00f3rio executivo do Hist\u00f3rico SAM", 10, False)), 1):
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=len(headers))
        cell = sheet.cell(row_number, 1, text); cell.font = Font(size=size, bold=row_number in {1, 3, 4}, color="FFFFFF" if dark else navy)
        cell.fill = PatternFill("solid", fgColor=navy if dark else pale); cell.alignment = Alignment(horizontal="center")
    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=10)
    sheet["A5"] = f"Recorte: {scope} · {len(rows)} evento(s)"
    sheet.merge_cells(start_row=5, start_column=11, end_row=5, end_column=len(headers))
    sheet["K5"] = f"Emissao: {now:%d/%m/%Y %H:%M} · Responsavel: {actor}"; sheet["K5"].alignment = Alignment(horizontal="right")
    for cell in sheet[5]: cell.font = Font(size=9, bold=True, color="536A7D")
    for column, label in enumerate(headers, 1):
        cell = sheet.cell(7, column, label); cell.font = Font(bold=True, color="FFFFFF", size=9)
        cell.fill = PatternFill("solid", fgColor=navy); cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number, event in enumerate(rows, 8):
        values = _event_report_row(event)
        occurred = values[0]
        try: values[0] = django_timezone.localtime(datetime.fromisoformat(str(occurred).replace("Z", "+00:00"))).strftime("%d/%m/%Y %H:%M")
        except (TypeError, ValueError): values[0] = str(occurred or "")
        for date_index in (10, 12, 13):
            value = values[date_index]
            if value:
                try: values[date_index] = datetime.fromisoformat(str(value)[:10]).strftime("%d/%m/%Y")
                except ValueError: pass
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value if value is not None else "")
            cell.fill = PatternFill("solid", fgColor="F4F6F8" if row_number % 2 == 0 else "FFFFFF")
            cell.alignment = Alignment(vertical="top", wrap_text=column in {5, 15, 19}); cell.border = Border(bottom=thin)
    widths = [20, 27, 28, 13, 42, 28, 28, 28, 22, 14, 15, 25, 18, 18, 55, 15, 17, 17, 30]
    for column, width in enumerate(widths, 1): sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A8"; sheet.auto_filter.ref = f"A7:{last}{max(7, sheet.max_row)}"; sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"; sheet.page_setup.fitToWidth = 1; sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True; sheet.print_title_rows = "1:7"
    stream = io.BytesIO(); workbook.save(stream)
    response = HttpResponse(stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Historico_SAM_{now:%d-%m-%Y_%H-%M}.xlsx"'
    return response


def export_pdf(request):
    if not _allowed(request):
        return JsonResponse({"error": "O Historico SAM exige perfil regional ou global."}, status=403)
    try:
        rows = _all_filtered_events(request)
    except requests.RequestException:
        return JsonResponse({"error": "Nao foi possivel preparar o PDF do Historico SAM."}, status=503)
    profile = request.session.get("user_profile") or {}
    actor = profile.get("full_name") or profile.get("nome") or profile.get("name") or profile.get("email") or "Usuario"
    now, scope = datetime.now(), _scope_label(request)
    stream = io.BytesIO(); navy = colors.HexColor("#1e4b7a"); pale = colors.HexColor("#eaf2f8")
    document = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm,
                                 topMargin=35 * mm, bottomMargin=13 * mm, title="Relat\u00f3rio executivo do Hist\u00f3rico SAM")
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("SamCell", parent=styles["BodyText"], fontName="Helvetica", fontSize=6.4,
                                leading=8, textColor=colors.HexColor("#263238"), alignment=TA_LEFT)
    header_style = ParagraphStyle("SamHeader", parent=cell_style, fontName="Helvetica-Bold", fontSize=6.5,
                                  leading=8, textColor=colors.white, alignment=TA_CENTER)
    context_style = ParagraphStyle("SamContext", parent=cell_style, fontName="Helvetica-Bold", fontSize=7.5,
                                   leading=9, textColor=navy)
    def text(value):
        import html
        return Paragraph(html.escape(str(value or "—")).replace("\n", "<br/>"), cell_style)
    data = [[Paragraph(label, header_style) for label in ("DATA/HORA", "OPERADOR / LOCALIDADE", "ACAO ORIGINAL",
                                                           "ALUNO / ENTIDADE", "TURMA / AULA", "ORIGEM")]]
    for event in rows:
        resolution = event.get("resolution") or {}; student = resolution.get("student") or {}
        lesson = resolution.get("class") or {}; group = resolution.get("group") or {}
        occurred = event.get("occurred_at")
        try: occurred = django_timezone.localtime(datetime.fromisoformat(str(occurred).replace("Z", "+00:00"))).strftime("%d/%m/%Y %H:%M")
        except (TypeError, ValueError): occurred = str(occurred or "")
        data.append([text(occurred), text((event.get("actor_name_snapshot") or "—") + "\n" + (event.get("congregation_snapshot") or "")),
                     text((event.get("action_type") or "outro").upper() + "\n" + (event.get("action_raw") or "")),
                     text(resolution.get("primary_name") or student.get("name") or "N\u00e3o identificado"),
                     text("\n".join(filter(None, (group.get("name") or lesson.get("name"),
                                                     f"Aula #{lesson.get('source_id')}" if lesson.get("source_id") else None)))),
                     text(" \\u00b7 ".join(filter(None, (event.get("device"), event.get("platform"), event.get("ip_address")))))])
    table = Table(data, colWidths=[27 * mm, 41 * mm, 69 * mm, 47 * mm, 41 * mm, 52 * mm], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#d8dee5")),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f6f8")]),
                               ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                               ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)]))
    story = [Table([[Paragraph(f"RECORTE DO RELATORIO · {scope} · {len(rows)} EVENTOS", context_style)]],
                   colWidths=[277 * mm], style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), pale),
                                                          ("BOX", (0, 0), (-1, -1), .4, colors.HexColor("#c8d1da")),
                                                          ("LEFTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5),
                                                          ("BOTTOMPADDING", (0, 0), (-1, -1), 5)])), Spacer(1, 3 * mm), table]
    def decorate(canvas, doc):
        width, height = landscape(A4); canvas.saveState(); canvas.setFillColor(colors.HexColor("#263238"))
        canvas.setFont("Helvetica-Bold", 15); canvas.drawCentredString(width / 2, height - 15 * mm, "CONGREGAÇÃO CRISTÃ NO BRASIL")
        canvas.setFont("Helvetica", 9); canvas.drawCentredString(width / 2, height - 20 * mm, "Regional Itapevi - São Paulo")
        canvas.setFillColor(navy); canvas.setFont("Helvetica-Bold", 12); canvas.drawCentredString(width / 2, height - 26 * mm, "GRUPO DE ESTUDOS MUSICAIS · SAM")
        canvas.setFillColor(colors.HexColor("#263238")); canvas.setFont("Helvetica", 9); canvas.drawCentredString(width / 2, height - 31 * mm, "Relat\u00f3rio executivo do Hist\u00f3rico SAM")
        canvas.setFont("Helvetica-Bold", 8); canvas.drawRightString(width - 10 * mm, height - 12 * mm, f"P\u00e1gina {doc.page}")
        canvas.setFont("Helvetica", 7); canvas.drawRightString(width - 10 * mm, height - 17 * mm, f"Emiss\u00e3o: {now:%d/%m/%Y %H:%M}")
        canvas.drawRightString(width - 10 * mm, height - 22 * mm, f"Respons\u00e1vel: {actor[:48]}")
        canvas.drawRightString(width - 10 * mm, height - 27 * mm, f"Recorte: {scope[:48]}")
        canvas.setFillColor(colors.HexColor("#6b7280")); canvas.drawString(10 * mm, 7 * mm, f"GEM · Historico SAM · {len(rows)} eventos")
        canvas.drawRightString(width - 10 * mm, 7 * mm, f"P\u00e1gina {doc.page}"); canvas.restoreState()
    document.build(story, onFirstPage=decorate, onLaterPages=decorate)
    response = HttpResponse(stream.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="Historico_SAM_{now:%d-%m-%Y_%H-%M}.pdf"'
    return response




def _normalized(value):
    return "".join(char for char in str(value or "").upper() if char.isalnum())


def _location_filters(request):
    return ({_normalized(value) for value in request.GET.getlist("municipio") if value},
            {_normalized(value) for value in request.GET.getlist("comum") if value})


def _common_core(common, municipality):
    display_name = str(common or "").strip().split(" - ", 1)[-1]
    normalize_words = lambda value: re.findall(
        r"[A-Z0-9]+",
        "".join(char for char in unicodedata.normalize("NFKD", str(value or "").upper())
                if not unicodedata.combining(char)),
    )
    city_words = set(normalize_words(municipality))
    words = [word for word in normalize_words(display_name)
             if word not in city_words and word not in {"DE", "DA", "DO", "DAS", "DOS", "E"}]
    return " ".join(words)


def _common_snapshot_lookup(catalog):
    cores_by_city = {}
    for item in catalog:
        common, municipality = item.get("comum"), item.get("cidade")
        core = _common_core(common, municipality)
        if core and municipality:
            cores_by_city.setdefault(core, set()).add(_normalized(municipality))
    ambiguous_cores = {core for core, cities in cores_by_city.items() if len(cities) > 1}
    matches = {}
    aliases_by_location = {}
    for item in catalog:
        common = str(item.get("comum") or "").strip()
        municipality = str(item.get("cidade") or "").strip()
        if not common:
            continue
        display_name = common.split(" - ", 1)[-1].strip()
        core = _common_core(common, municipality)
        aliases = {common, display_name}
        if core:
            if core not in ambiguous_cores:
                aliases.add(core)
            if municipality:
                aliases.update({f"{municipality} {core}", f"{core} {municipality}"})
        location = (_normalized(municipality), _normalized(common))
        aliases = {alias for alias in aliases
                   if not (core in ambiguous_cores and _normalized(alias) == _normalized(core))}
        aliases_by_location.setdefault(location, set()).update(aliases)
        for alias in aliases:
            key = _normalized(alias)
            matches.setdefault(key, set()).add(location)
    lookup = {alias: next(iter(locations)) for alias, locations in matches.items() if len(locations) == 1}
    return lookup, aliases_by_location


def _event_location_query_params(request):
    municipalities, commons = _location_filters(request)
    if not municipalities and not commons:
        return None

    catalog = common_catalog()
    selected_commons = {_normalized(value) for value in request.GET.getlist("comum") if value}
    matching_rows = [item for item in catalog if item.get("comum")
                     and (not municipalities or _normalized(item.get("cidade")) in municipalities)
                     and (not selected_commons or _normalized(item.get("comum")) in selected_commons)]
    if not matching_rows:
        return {}

    _, aliases_by_location = _common_snapshot_lookup(catalog)
    aliases = set()
    for item in matching_rows:
        common, municipality = str(item.get("comum") or "").strip(), str(item.get("cidade") or "").strip()
        aliases.update(aliases_by_location.get((_normalized(municipality), _normalized(common)), set()))
    if not aliases:
        return {}

    values = []
    for value in sorted(aliases):
        escaped = value.replace("\\", "\\\\").replace('"', '\\"')
        values.append(f'"{escaped}"')
    return {"congregation_snapshot": "in.(" + ",".join(values) + ")"}


def _event_location(event, snapshot_lookup):
    resolution = event.get("resolution") or {}
    student = resolution.get("student") or {}
    group = resolution.get("group") or {}
    lesson = resolution.get("class") or {}
    common = (student.get("congregation") or group.get("congregation") or lesson.get("congregation")
              or event.get("congregation_snapshot") or "")
    return snapshot_lookup.get(_normalized(common), ("", ""))


def _filter_by_location(request, rows):
    municipalities, commons = _location_filters(request)
    if not municipalities and not commons:
        return rows
    snapshot_lookup, _ = _common_snapshot_lookup(common_catalog())
    return [row for row in rows if (
        (location := _event_location(row, snapshot_lookup)) != ("", "")
        and (not municipalities or location[0] in municipalities)
        and (not commons or location[1] in commons)
    )]



def api_location_options(request):
    """Catálogo de município e comum, idêntico à base usada pelo Painel de Alunos."""
    if not _allowed(request):
        return JsonResponse({"error": "O Histórico SAM exige perfil regional ou global."}, status=403)
    try:
        catalog = common_catalog()
        municipalities = sorted({str(row.get("cidade") or "").strip() for row in catalog if row.get("cidade")})
        commons = sorted(
            (row for row in catalog if row.get("comum")),
            key=lambda row: (str(row.get("comum") or "").casefold(), str(row.get("cidade") or "").casefold()),
        )
        return JsonResponse({
            "municipalities": [{"label": value, "value": value} for value in municipalities],
            "commons": [{"label": str(row.get("comum")).strip(), "value": str(row.get("comum")).strip(),
                         "municipality": str(row.get("cidade") or "").strip()} for row in commons],
        })
    except requests.RequestException:
        return JsonResponse({"error": "Não foi possível carregar municípios e comuns."}, status=503)
