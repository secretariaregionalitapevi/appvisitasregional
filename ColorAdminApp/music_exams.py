"""Gestão individual de exames e testes da Música."""
import csv, hashlib, io, json, re, unicodedata, uuid
from difflib import SequenceMatcher
from datetime import datetime
import requests
from django.conf import settings
from django.core.cache import cache
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .access_control import can_access, service_headers, user_scope
from .gem import _fetch_students, can_open_module

TABLE="musica_exames_individuais"; BUCKET="gem_exam_documents"
AGENDA_CSV="https://docs.google.com/spreadsheets/d/e/2PACX-1vSr59qtjfIiXj8PKems0FQzJulSdyUa70QhH8KaJSnNtbh9U0cB0KrRrxaKz93pc9CSLmTAn8btFc-0/pub?output=csv&gid=0&single=true"
RESULTS={"AGENDADO","APROVADO","RETORNOU AO CICLO","AUSENTE"}
LEGACY_FIELDS=("musicos_culto_oficial","musicos_oficializacao","musicos_troca_instrumento","organistas_rjm","organistas_culto_oficial","organistas_oficializacao","organistas_testes_especiais")
GROUPS=["ITAPEVI","JANDIRA","COTIA / CAUCAIA / VGP","SANTANA / PIRAPORA"]

def norm(value):
    text=unicodedata.normalize("NFKD",str(value or ""))
    return " ".join("".join(c for c in text if not unicodedata.combining(c)).upper().split())
def regional_group(city):
    value=norm(city)
    if "JANDIRA" in value:return GROUPS[1]
    if any(x in value for x in ("COTIA","CAUCAIA","VARGEM GRANDE","VGP")):return GROUPS[2]
    if any(x in value for x in ("SANTANA","PIRAPORA")):return GROUPS[3]
    return GROUPS[0]
def actor(request):
    p=request.session.get("user_profile") or {}
    return p.get("full_name") or p.get("name") or p.get("email") or "Usuário"
def rest_url():return f"{settings.SUPABASE_URL}/rest/v1/{TABLE}"
def denied():return JsonResponse({"error":"Seu perfil não possui acesso à pasta Música."},status=403)
def query_filters(request):
    return {key:request.GET.get(key,"").strip() for key in ("municipio","comum","mes","inicio","fim","lancamento_inicio","lancamento_fim")}
def matches_filters(row,filters,legacy=False):
    if filters["municipio"] and regional_group(row.get("municipio"))!=filters["municipio"]:return False
    if not legacy and filters["comum"] and norm(row.get("comum"))!=norm(filters["comum"]):return False
    exam=str(row.get("data_exame") or "")[:10];launched=str(row.get("created_at") or "")[:10]
    if filters["mes"] and exam[5:7]!=filters["mes"]:return False
    if filters["inicio"] and exam<filters["inicio"]:return False
    if filters["fim"] and exam>filters["fim"]:return False
    if filters["lancamento_inicio"] and launched<filters["lancamento_inicio"]:return False
    if filters["lancamento_fim"] and launched>filters["lancamento_fim"]:return False
    return True
def fetch_rows(request,apply_filters=True):
    response=requests.get(rest_url(),headers=service_headers(),params={"select":"*","order":"data_exame.desc,created_at.desc"},timeout=25);response.raise_for_status()
    scope=user_scope(request);rows=[row for row in response.json() if can_access(scope,{"municipio":row.get("municipio"),"comum":row.get("comum")})]
    return [row for row in rows if matches_filters(row,query_filters(request))] if apply_filters else rows
def fetch_legacy(request):
    select="id,municipio,data_exame,created_at,lancado_por,"+",".join(LEGACY_FIELDS)
    response=requests.get(f"{settings.SUPABASE_URL}/rest/v1/musica_exames_lancamentos",headers=service_headers(),params={"select":select,"order":"data_exame.desc,created_at.desc"},timeout=25);response.raise_for_status()
    scope=user_scope(request);rows=[row for row in response.json() if can_access(scope,{"municipio":row.get("municipio"),"comum":""})]
    return [row for row in rows if matches_filters(row,query_filters(request),True)]
def legacy_summary(rows):
    groups={name:{"grupo":name,**{field:0 for field in LEGACY_FIELDS}} for name in GROUPS}
    for row in rows:
        item=groups[regional_group(row.get("municipio"))]
        for field in LEGACY_FIELDS:item[field]+=int(row.get(field) or 0)
    for item in groups.values():
        item["musicos_total"]=sum(item[x] for x in LEGACY_FIELDS[:3]);item["organistas_total"]=sum(item[x] for x in LEGACY_FIELDS[3:]);item["total"]=item["musicos_total"]+item["organistas_total"]
    totals={field:sum(item[field] for item in groups.values()) for field in LEGACY_FIELDS};totals["musicos_total"]=sum(totals[x] for x in LEGACY_FIELDS[:3]);totals["organistas_total"]=sum(totals[x] for x in LEGACY_FIELDS[3:]);totals["total"]=totals["musicos_total"]+totals["organistas_total"]
    return list(groups.values()),totals
def idempotency_key(payload):
    fields=("aluno_id","nome_aluno","data_exame","tipo_exame","instrumento")
    return hashlib.sha256("|".join(norm(payload.get(x)) for x in fields).encode()).hexdigest()
def visible_students(request):
    scope=user_scope(request)
    return [s for s in _fetch_students() if can_access(scope,{"municipio":s.get("municipio"),"comum":s.get("comum_congregacao")})]

def match_student(students,name):
    wanted=norm(name).replace("�","")
    exact=next((s for s in students if norm(s.get("nome_aluno"))==norm(name)),None)
    if exact:return exact
    ranked=sorted(((SequenceMatcher(None,wanted,norm(s.get("nome_aluno"))).ratio(),s) for s in students),key=lambda item:item[0],reverse=True)
    return ranked[0][1] if ranked and ranked[0][0]>=.94 and (len(ranked)==1 or ranked[0][0]-ranked[1][0]>=.03) else None

def agenda_events():
    cached=cache.get("musica_exam_agenda_events")
    if cached is not None:return cached
    response=requests.get(AGENDA_CSV,timeout=12);response.raise_for_status();events=[]
    for row in csv.DictReader(io.StringIO(response.text.lstrip("\ufeff"))):
        values={norm(k):v for k,v in row.items()}
        if norm(values.get("DEPARTAMENTO"))!="MUSICA":continue
        try:exam_date=datetime.strptime((values.get("DATA") or "").strip(),"%d/%m/%Y").date().isoformat()
        except ValueError:continue
        title=(values.get("EVENTO") or "").strip();hour=(values.get("HORA") or "").strip()
        if not title:continue
        event_id=hashlib.sha256(f"{exam_date}|{hour}|{norm(title)}".encode()).hexdigest()[:24]
        events.append({"id":event_id,"data":exam_date,"hora":hour,"titulo":title,"url":"https://ccbagenda.vercel.app/"})
    events.sort(key=lambda x:(x["data"],x["hora"]),reverse=True);cache.set("musica_exam_agenda_events",events,600);return events

def api_agenda(request):
    if not can_open_module(request):return denied()
    try:return JsonResponse({"items":agenda_events()})
    except requests.RequestException:return JsonResponse({"error":"A agenda não respondeu. Tente novamente."},status=502)
def page(request):
    if not can_open_module(request):return render(request,"pages/403.html",{"message":"Seu perfil não possui acesso à pasta Música."},status=403)
    return render(request,"pages/music_exams.html",{"report_user":actor(request)})
def api_dashboard(request):
    if not can_open_module(request):return denied()
    try:
        items=fetch_rows(request);legacy=fetch_legacy(request);groups,totals=legacy_summary(legacy)
        students=[{"id":s.get("id"),"nome":s.get("nome_aluno"),"registro_msa":s.get("registro_msa"),"municipio":s.get("municipio"),"comum":s.get("comum_congregacao"),"instrumento":s.get("instrumento"),"categoria":s.get("cargo_ministerio")} for s in visible_students(request)]
        launches=[{**row,"grupo":regional_group(row.get("municipio")),"total":sum(int(row.get(f) or 0) for f in LEGACY_FIELDS)} for row in legacy]
        return JsonResponse({"items":items,"groups":groups,"totals":totals,"students":students,"legacy_launches":launches,"filters":query_filters(request)})
    except requests.RequestException:return JsonResponse({"error":"Não foi possível consultar os exames. Verifique a migração 022."},status=502)
def api_record(request,record_id=None):
    if not can_open_module(request):return denied()
    if request.method!="POST" or record_id:return JsonResponse({"error":"Método não permitido."},status=405)
    try:
        data=request.POST.dict() if (request.content_type or "").startswith("multipart/") else json.loads(request.body or "{}")
        student=next((s for s in visible_students(request) if str(s.get("id"))==str(data.get("aluno_id"))),None)
        if not student:return JsonResponse({"error":"Selecione um aluno válido dentro do seu escopo."},status=400)
        event=next((item for item in agenda_events() if item["id"]==data.get("agenda_evento_id")),None)
        if not event:return JsonResponse({"error":"Selecione o evento correspondente na agenda da Música."},status=400)
        result=norm(data.get("resultado") or "AGENDADO").replace("REPROVADO","RETORNOU AO CICLO")
        if result not in RESULTS:return JsonResponse({"error":"Resultado inválido."},status=400)
        payload={"aluno_id":student["id"],"nome_aluno":student.get("nome_aluno"),"registro_msa":student.get("registro_msa"),"municipio":student.get("municipio"),"comum":student.get("comum_congregacao"),"instrumento":student.get("instrumento"),"categoria":student.get("cargo_ministerio"),"data_exame":event["data"],"tipo_exame":norm(data.get("tipo_exame")),"resultado":result,"encarregado_local":data.get("encarregado_local") or None,"observacoes":data.get("observacoes") or None,"lancado_por":actor(request),"origem":"PAINEL","agenda_evento_id":event["id"],"agenda_evento_titulo":event["titulo"],"agenda_evento_data":event["data"],"agenda_evento_hora":event["hora"] or None,"agenda_evento_url":event["url"]}
        if not payload["data_exame"] or not payload["tipo_exame"]:return JsonResponse({"error":"Informe a data e o tipo do exame."},status=400)
        payload["chave_idempotencia"]=idempotency_key(payload);proof=request.FILES.get("prova")
        if proof:
            if proof.size>12*1024*1024:return JsonResponse({"error":"A prova deve ter no máximo 12 MB."},status=400)
            ext=re.sub(r"[^a-z0-9]","",proof.name.rsplit(".",1)[-1].lower())[:8] or "pdf";path=f"{student['id']}/{uuid.uuid4()}.{ext}"
            upload=requests.post(f"{settings.SUPABASE_URL}/storage/v1/object/{BUCKET}/{path}",headers={"apikey":settings.SUPABASE_SERVICE_ROLE_KEY,"Authorization":f"Bearer {settings.SUPABASE_SERVICE_ROLE_KEY}","Content-Type":proof.content_type or "application/octet-stream"},data=proof.read(),timeout=40);upload.raise_for_status();payload.update({"prova_url":path,"prova_nome":proof.name})
        response=requests.post(rest_url()+"?on_conflict=chave_idempotencia",headers=service_headers("resolution=merge-duplicates,return=representation"),json=payload,timeout=25);response.raise_for_status()
        return JsonResponse({"ok":True,"item":(response.json() or [payload])[0]})
    except (ValueError,requests.RequestException) as error:return JsonResponse({"error":f"Não foi possível salvar o exame: {error}"},status=502)

def parse_reference_workbook(file,request):
    sheet=load_workbook(file,data_only=True).worksheets[0];students=visible_students(request);current_type="";current_date=None;parsed=[];unmatched=[]
    for values in sheet.iter_rows(values_only=True):
        cells=[str(v).strip() if v is not None else "" for v in values];joined=" ".join(cells[:6]);normalized=norm(joined);match=re.search(r"(\d{2}/\d{2}/\d{4})",joined)
        if match:current_date=datetime.strptime(match.group(1),"%d/%m/%Y").date().isoformat()
        if "PRE" in normalized and ("OFICIAL" in normalized or "AVALIA" in normalized):current_type="PRÉ-AVALIAÇÃO DE OFICIALIZAÇÃO"
        elif "CULTO OFICIAL" in normalized:current_type="EXAME DE CULTO OFICIAL"
        elif "TROCA DE INSTRUMENT" in normalized:current_type="TROCA DE INSTRUMENTO"
        if len(cells)<6 or not re.fullmatch(r"\d+(?:\.0)?",cells[0]) or not cells[1]:continue
        name,instrument,common,leader,result=cells[1:6];student=match_student(students,name)
        if not student:unmatched.append({"nome":name,"comum":common,"linha":cells[0]});continue
        payload={"aluno_id":student["id"],"nome_aluno":student.get("nome_aluno"),"registro_msa":student.get("registro_msa"),"municipio":student.get("municipio"),"comum":student.get("comum_congregacao") or common,"instrumento":instrument or student.get("instrumento"),"categoria":student.get("cargo_ministerio"),"data_exame":current_date,"tipo_exame":current_type or "APTIDÃO PERIÓDICA","resultado":norm(result) if norm(result) in RESULTS else "AGENDADO","encarregado_local":leader or None,"lancado_por":actor(request),"origem":"IMPORTACAO"}
        payload["chave_idempotencia"]=idempotency_key(payload);parsed.append(payload)
    return parsed,unmatched
def api_import(request):
    if not can_open_module(request):return denied()
    if request.method!="POST" or not request.FILES.get("arquivo"):return JsonResponse({"error":"Envie uma planilha Excel."},status=400)
    try:
        parsed,unmatched=parse_reference_workbook(request.FILES["arquivo"],request);existing={row.get("chave_idempotencia") for row in fetch_rows(request,False)};batch=[row for row in parsed if row["chave_idempotencia"] not in existing];duplicates=len(parsed)-len(batch)
        if batch:
            response=requests.post(rest_url()+"?on_conflict=chave_idempotencia",headers=service_headers("resolution=ignore-duplicates,return=minimal"),json=batch,timeout=40);response.raise_for_status()
        return JsonResponse({"ok":True,"importados":len(batch),"duplicados":duplicates,"nao_encontrados":unmatched})
    except Exception as error:return JsonResponse({"error":f"Não foi possível importar: {error}"},status=400)
def api_participants(request):
    if not can_open_module(request):return denied()
    try:return JsonResponse({"items":[row for row in fetch_rows(request) if norm(row.get("resultado"))=="AGENDADO"]})
    except requests.RequestException:return JsonResponse({"error":"Não foi possível carregar participantes."},status=502)
def document(request,record_id):
    if not can_open_module(request):return denied()
    try:
        response=requests.get(rest_url(),headers=service_headers(),params={"select":"prova_url,prova_nome,municipio,comum","id":f"eq.{record_id}","limit":1},timeout=15);response.raise_for_status();items=response.json()
        if not items or not items[0].get("prova_url") or not can_access(user_scope(request),items[0]):return JsonResponse({"error":"Documento não encontrado."},status=404)
        item=items[0];download=requests.get(f"{settings.SUPABASE_URL}/storage/v1/object/authenticated/{BUCKET}/{item['prova_url']}",headers={"apikey":settings.SUPABASE_SERVICE_ROLE_KEY,"Authorization":f"Bearer {settings.SUPABASE_SERVICE_ROLE_KEY}"},timeout=30);download.raise_for_status();result=HttpResponse(download.content,content_type=download.headers.get("Content-Type","application/octet-stream"));result["Content-Disposition"]=f'inline; filename="{item.get("prova_nome") or "prova"}"';return result
    except requests.RequestException:return JsonResponse({"error":"Não foi possível abrir o documento."},status=502)
def export_excel(request):
    if not can_open_module(request):return denied()
    groups,totals=legacy_summary(fetch_legacy(request));items=fetch_rows(request);wb=Workbook();ws=wb.active;ws.title="Resumo Regional"
    def title(sheet,last_col,report):
        sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last_col);sheet["A1"]="CONGREGAÇÃO CRISTÃ NO BRASIL";sheet["A1"].font=Font(bold=True,color="FFFFFF",size=15);sheet["A1"].fill=PatternFill("solid",fgColor="1E4B7A");sheet["A1"].alignment=Alignment(horizontal="center")
        for row,text in ((2,"Regional Itapevi - São Paulo · MÚSICA"),(3,report),(4,f"Emissão: {datetime.now():%d/%m/%Y %H:%M} · Responsável: {actor(request)}")):sheet.merge_cells(start_row=row,start_column=1,end_row=row,end_column=last_col);sheet.cell(row,1,text)
    title(ws,9,"Relatório consolidado de Exames e Testes")
    headers=["REGIÃO","Músicos · Culto Oficial","Músicos · Oficialização","Músicos · Troca Instrumento","Organistas · RJM","Organistas · Culto Oficial","Organistas · Oficialização","Organistas · Teste Especial","TOTAL GERAL"]
    for column,value in enumerate(headers,1):cell=ws.cell(6,column,value);cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="1E4B7A");cell.alignment=Alignment(horizontal="center",wrap_text=True)
    fields=list(LEGACY_FIELDS)
    for index,item in enumerate(groups,1):
        for column,value in enumerate([item["grupo"]]+[item[x] for x in fields]+[item["total"]],1):ws.cell(6+index,column,value)
    for column,value in enumerate(["TOTAIS GERAIS"]+[totals[x] for x in fields]+[totals["total"]],1):cell=ws.cell(11,column,value);cell.font=Font(bold=True);cell.fill=PatternFill("solid",fgColor="EAF2F8")
    for index,width in enumerate([26,18,18,22,16,19,19,19,14],1):ws.column_dimensions[chr(64+index)].width=width
    ws.freeze_panes="A7";ws.auto_filter.ref="A6:I10";ws.sheet_view.showGridLines=False;ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1
    detail=wb.create_sheet("Registros Individuais");title(detail,10,"Histórico individual de Exames e Testes");detail_headers=["Nº","Aluno(a)","Registro MSA","Município","Comum","Instrumento","Tipo","Data","Resultado","Prova"]
    for column,value in enumerate(detail_headers,1):cell=detail.cell(6,column,value);cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="1E4B7A");cell.alignment=Alignment(horizontal="center")
    for index,item in enumerate(items,1):
        exam_date=datetime.strptime(item["data_exame"],"%Y-%m-%d").strftime("%d/%m/%Y") if item.get("data_exame") else "—"
        for column,value in enumerate([index,item.get("nome_aluno"),item.get("registro_msa"),item.get("municipio"),item.get("comum"),item.get("instrumento"),item.get("tipo_exame"),exam_date,item.get("resultado"),"Sim" if item.get("prova_url") else "Não"],1):detail.cell(6+index,column,value)
    for index,width in enumerate([6,30,14,18,34,20,28,13,20,12],1):detail.column_dimensions[chr(64+index)].width=width
    detail.freeze_panes="A7";detail.auto_filter.ref=f"A6:J{6+len(items)}";detail.sheet_view.showGridLines=False;detail.page_setup.orientation="landscape";detail.page_setup.fitToWidth=1
    stream=io.BytesIO();wb.save(stream);stamp=datetime.now().strftime("%d-%m-%Y_%H-%M");response=HttpResponse(stream.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet");response["Content-Disposition"]=f'attachment; filename="Exames_GEM_{stamp}.xlsx"';return response