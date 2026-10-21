import json
from unittest.mock import Mock, patch
import requests

from django.test import RequestFactory, SimpleTestCase
from django.core.files.uploadedfile import SimpleUploadedFile

from ColorAdminApp.modules.darpe import RESOURCES, _regional_rows, _with_canonical_commons, _resolve_import_common, _read_imported_cadastros, _read_all_rows, _write_cadastro, api_resource, export_cadastros_excel, import_cadastros, page
from ColorAdminApp.modules.music_events_identity import CommonResolver
from ColorAdminApp.module_access import MODULE_DARPE, allowed_modules, primary_module


def request_for(path, method="get", body=None):
    factory = RequestFactory()
    request = (factory.get(path) if method == "get" else
               getattr(factory, method)(path, data=json.dumps(body or {}), content_type="application/json"))
    request.session = {"authenticated": True, "user_id": "user-1", "user_profile": {
        "user_id": "user-1", "role_id": 2, "sector": "D.A.R.P.E", "access_level": "regional",
        "full_name": "Operador DARPE",
    }}
    return request


class DarpeModuleTests(SimpleTestCase):
    @patch('ColorAdminApp.modules.darpe._with_canonical_commons', side_effect=lambda rows: rows)
    @patch('ColorAdminApp.modules.darpe.requests.patch')
    @patch('ColorAdminApp.modules.darpe.requests.get')
    def test_unchecked_role_remains_false_on_fresh_list_read(self, get, update, _enrich):
        stored = {'id': 'test', 'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL',
                  'cidade': 'ITAPEVI', 'darpe_ministerio': True, 'darpe_colaborador': True}
        get.return_value = Mock(headers={'Content-Range': '0-0/1'}, json=lambda: [dict(stored)])
        def persist(*args, **kwargs):
            stored.update(kwargs['json'])
            return Mock(status_code=200, json=lambda: [dict(stored)])
        update.side_effect = persist
        saved = api_resource(request_for('/darpe/api/cadastros/test/', 'patch', {
            'darpe_ministerio': True, 'darpe_colaborador': False,
        }), 'cadastros', 'test')
        self.assertEqual(saved.status_code, 200)
        refreshed = api_resource(request_for('/darpe/api/cadastros/'), 'cadastros')
        self.assertEqual(refreshed.status_code, 200)
        row = json.loads(refreshed.content)['items'][0]
        self.assertTrue(row['darpe_ministerio'])
        self.assertIs(row['darpe_colaborador'], False)
        self.assertIn('no-store', refreshed['Cache-Control'])
        self.assertIn('no-store', saved['Cache-Control'])

    def import_request(self):
        request = request_for('/darpe/api/cadastros/importar/', 'post')
        request.FILES['arquivo'] = SimpleUploadedFile('cadastro.xlsx', b'fixture')
        return request

    @patch('ColorAdminApp.modules.darpe.requests.get')
    def test_reads_all_pages_even_when_server_page_limit_is_lower(self, get):
        get.side_effect = [Mock(headers={'Content-Range': '0-1/3'}, json=lambda: [{'id': '1'}, {'id': '2'}]),
                           Mock(headers={'Content-Range': '2-2/3'}, json=lambda: [{'id': '3'}])]
        self.assertEqual(len(_read_all_rows(RESOURCES['cadastros'])), 3)
        self.assertEqual(get.call_args.kwargs['params']['offset'], 2)

    @patch('ColorAdminApp.modules.darpe.requests.delete')
    @patch('ColorAdminApp.modules.darpe.requests.post')
    @patch('ColorAdminApp.modules.darpe._read_all_rows', return_value=[])
    @patch('ColorAdminApp.modules.darpe._read_imported_cadastros')
    def test_import_scales_with_batches_not_people(self, read, current, post, delete):
        post.return_value.status_code = 201
        for size in (15, 150, 1500, 15000):
            with self.subTest(size=size):
                post.reset_mock()
                read.return_value = [{'nome_completo': f'Pessoa {i}', 'comum_congregacao': 'CENTRAL'}
                                     for i in range(size)]
                response = import_cadastros(self.import_request())
                self.assertEqual(response.status_code, 200)
                self.assertEqual(post.call_count, (size + 499) // 500)
                self.assertEqual(sum(len(call.kwargs['json']) for call in post.call_args_list), size)
                self.assertEqual(post.call_args.kwargs['params'], {'on_conflict': 'id'})
        delete.assert_not_called()

    @patch('ColorAdminApp.modules.darpe.requests.delete')
    @patch('ColorAdminApp.modules.darpe.requests.post')
    @patch('ColorAdminApp.modules.darpe._read_all_rows')
    @patch('ColorAdminApp.modules.darpe._read_imported_cadastros')
    def test_import_skips_unchanged_and_preserves_id_and_creation_author(self, read, current, post, delete):
        row = {'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL', 'telefone': '123',
               'created_by_name': 'Importação'}
        read.return_value = [row]
        current.return_value = [{**row, 'id': 'original', 'created_by_name': 'Autor original'}]
        self.assertEqual(import_cadastros(self.import_request()).status_code, 200)
        post.assert_not_called()
        delete.assert_not_called()
        read.return_value = [{**row, 'telefone': '456'}]
        post.return_value.status_code = 200
        self.assertEqual(import_cadastros(self.import_request()).status_code, 200)
        saved = post.call_args.kwargs['json'][0]
        self.assertEqual(saved['id'], 'original')
        self.assertEqual(saved['created_by_name'], 'Autor original')

    @patch('ColorAdminApp.modules.darpe.requests.post')
    @patch('ColorAdminApp.modules.darpe._read_all_rows')
    @patch('ColorAdminApp.modules.darpe._read_imported_cadastros')
    def test_import_blank_instrument_preserves_sam_enrichment(self, read, current, post):
        row = {'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL', 'instrumento_oficializado': None}
        read.return_value = [row]
        current.return_value = [{**row, 'id': 'original', 'instrumento_oficializado': 'VIOLINO'}]
        self.assertEqual(import_cadastros(self.import_request()).status_code, 200)
        post.assert_not_called()

    @patch('ColorAdminApp.modules.darpe.requests.post')
    @patch('ColorAdminApp.modules.darpe._read_all_rows')
    @patch('ColorAdminApp.modules.darpe._read_imported_cadastros')
    def test_import_keeps_sam_spelling_when_sheet_uses_equivalent_variant(self, read, current, post):
        row = {'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL', 'instrumento_oficializado': 'Violoncello'}
        read.return_value = [row]
        current.return_value = [{**row, 'id': 'original', 'instrumento_oficializado': 'VIOLONCELO'}]
        self.assertEqual(import_cadastros(self.import_request()).status_code, 200)
        post.assert_not_called()

    @patch('ColorAdminApp.modules.darpe.requests.delete')
    @patch('ColorAdminApp.modules.darpe.requests.post')
    @patch('ColorAdminApp.modules.darpe._read_all_rows')
    @patch('ColorAdminApp.modules.darpe._read_imported_cadastros')
    def test_failed_write_never_starts_cleanup(self, read, current, post, delete):
        read.return_value = [{'nome_completo': f'Pessoa {i}', 'comum_congregacao': 'CENTRAL'} for i in range(501)]
        current.return_value = [{'id': 'stale', 'nome_completo': 'Antigo', 'comum_congregacao': 'CENTRAL'}]
        post.side_effect = [Mock(status_code=201), requests.Timeout()]
        self.assertEqual(import_cadastros(self.import_request()).status_code, 502)
        delete.assert_not_called()

    @patch('ColorAdminApp.modules.darpe.requests.delete')
    @patch('ColorAdminApp.modules.darpe.requests.post')
    @patch('ColorAdminApp.modules.darpe._read_all_rows')
    @patch('ColorAdminApp.modules.darpe._read_imported_cadastros')
    def test_cleanup_is_batched_and_keeps_canonical_id(self, read, current, post, delete):
        read.return_value = [{'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL'}]
        current.return_value = [{'id': 'keep', **read.return_value[0]}] + [
            {'id': f'stale-{i}', 'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL'} for i in range(205)]
        self.assertEqual(import_cadastros(self.import_request()).status_code, 200)
        post.assert_not_called()
        self.assertEqual(delete.call_count, 3)
        self.assertTrue(all('keep' not in call.kwargs['params']['id'] for call in delete.call_args_list))

    def test_import_batch_adapts_to_legacy_schema(self):
        writer = Mock(return_value=Mock(status_code=201))
        rows = [{'nome_completo': 'Pessoa', 'darpe_musico': False,
                 'darpe_organista': True, 'darpe_musico_organista': True}]
        _write_cadastro(writer, reference={'darpe_musico_organista': False}, json=rows)
        self.assertEqual(writer.call_args.kwargs['json'], [
            {'nome_completo': 'Pessoa', 'darpe_musico_organista': True}])
        self.assertIn('darpe_organista', rows[0])

    def test_new_schema_keeps_separate_roles_without_legacy_column(self):
        writer = Mock(return_value=Mock(status_code=200))
        _write_cadastro(writer, reference={'darpe_musico': False, 'darpe_organista': False},
                        json={'darpe_musico': False, 'darpe_organista': True,
                              'darpe_musico_organista': True})
        self.assertEqual(writer.call_args.kwargs['json'], {'darpe_musico': False, 'darpe_organista': True})

    @patch('ColorAdminApp.modules.darpe.requests.patch')
    @patch('ColorAdminApp.modules.darpe.requests.get')
    def test_edit_legacy_schema_clears_optional_fields_and_saves_musical_role(self, get, update):
        get.return_value.json.return_value = [{'id': 'test', 'nome_completo': 'Pessoa',
            'comum_congregacao': 'CENTRAL', 'cidade': 'ITAPEVI', 'darpe_musico_organista': False}]
        update.return_value.status_code = 200
        update.return_value.json.return_value = [{'id': 'test'}]
        response = api_resource(request_for('/darpe/api/cadastros/test/', 'patch', {
            'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL', 'cidade': 'ITAPEVI',
            'cpf': '', 'cep': '', 'uf': '', 'data_nascimento': '', 'data_batismo': '',
            'declaracao_data': '', 'darpe_musico': True, 'darpe_organista': False,
        }), 'cadastros', 'test')
        self.assertEqual(response.status_code, 200)
        saved = update.call_args.kwargs['json']
        for field in ('cpf', 'cep', 'uf', 'data_nascimento', 'data_batismo', 'declaracao_data'):
            self.assertIsNone(saved[field])
        self.assertTrue(saved['darpe_musico_organista'])
        self.assertNotIn('darpe_musico', saved)
        self.assertNotIn('darpe_organista', saved)

    @patch('ColorAdminApp.modules.darpe.requests.post')
    def test_create_retries_only_missing_musical_columns(self, post):
        missing = Mock(status_code=400)
        missing.json.return_value = {'code': 'PGRST204', 'message': "Could not find the 'darpe_musico' column"}
        saved = Mock(status_code=201)
        saved.json.return_value = [{'id': 'test'}]
        post.side_effect = [missing, saved]
        response = api_resource(request_for('/darpe/api/cadastros/', 'post', {
            'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL', 'cidade': 'ITAPEVI',
            'darpe_musico': True, 'darpe_organista': False,
        }), 'cadastros')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(post.call_count, 2)
        self.assertTrue(post.call_args.kwargs['json']['darpe_musico_organista'])
        self.assertNotIn('darpe_musico', post.call_args.kwargs['json'])

    @patch('ColorAdminApp.modules.darpe.requests.post')
    def test_invalid_dates_and_documents_fail_before_database_write(self, post):
        for field, value in [('data_nascimento', '2026-02-30'), ('cpf', '123'), ('cep', '123'), ('uf', 'SPP')]:
            with self.subTest(field=field):
                response = api_resource(request_for('/darpe/api/cadastros/', 'post', {
                    'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL', field: value,
                }), 'cadastros')
                self.assertEqual(response.status_code, 400)
        post.assert_not_called()

    @patch('ColorAdminApp.modules.darpe.requests.post')
    def test_duplicate_is_reported_as_conflict_without_retry(self, post):
        failure = Mock(status_code=409)
        failure.json.return_value = {'code': '23505'}
        failure.raise_for_status.side_effect = requests.HTTPError(response=failure)
        post.return_value = failure
        response = api_resource(request_for('/darpe/api/cadastros/', 'post', {
            'nome_completo': 'Pessoa', 'comum_congregacao': 'CENTRAL',
        }), 'cadastros')
        self.assertEqual(response.status_code, 409)
        self.assertIn('CPF', response.content.decode())
        post.assert_called_once()

    def test_import_resolves_abbreviations_and_city_in_common_name(self):
        catalog = [
            {"comum": "BR-22-3061 - JARDIM RUTH", "cidade": "ITAPEVI"},
            {"comum": "BR-22-0673 - VILA DOUTOR CARDOSO", "cidade": "ITAPEVI"},
            {"comum": "BR-22-0683 - JARDIM ROSEMARY", "cidade": "ITAPEVI"},
            {"comum": "BR-22-3050 - JARDIM ROSEMARY II", "cidade": "ITAPEVI"},
            {"comum": "BR-22-0748 - CENTRAL DE JANDIRA", "cidade": "JANDIRA"},
        ]
        resolver = CommonResolver(catalog)
        for common, index in [("JD RUTH", 0), ("VL DR CARDOSO - ITAPEVI", 1),
                              ("JD ROSEMARY I - ITAPEVI", 2), ("JANDIRA - CENTRAL", 4)]:
            with self.subTest(common=common):
                self.assertEqual(_resolve_import_common(resolver, common, None), catalog[index])
        self.assertIsNone(_resolve_import_common(resolver, "JD RUTH", "JANDIRA"))
        self.assertIsNone(_resolve_import_common(resolver, "BR-99-9999 - JARDIM RUTH", None))

    def test_import_rejects_ambiguous_names(self):
        resolver = CommonResolver([
            {"comum": "BR-22-0001 - JARDIM RUTH", "cidade": "ITAPEVI"},
            {"comum": "BR-22-0002 - JARDIM RUTH", "cidade": "JANDIRA"},
        ])
        self.assertIsNone(_resolve_import_common(resolver, "JD RUTH", None))

    @patch("ColorAdminApp.modules.darpe.common_catalog", return_value=[])
    @patch("ColorAdminApp.modules.darpe.load_workbook")
    def test_import_error_identifies_unresolved_line(self, load, _catalog):
        values = [None] * 29
        values[1], values[7] = "Pessoa Teste", "COMUM DESCONHECIDA"
        load.return_value.__getitem__.return_value.iter_rows.return_value = [values]
        with self.assertRaisesRegex(ValueError, "linha 4: COMUM DESCONHECIDA"):
            _read_imported_cadastros(Mock())

    def test_only_regional_municipalities_are_listed(self):
        rows = [{"cidade": "COTIA"}, {"cidade": "OSASCO"}, {"cidade": "SANTANA DE PARNAÍBA"}]
        self.assertEqual([row["cidade"] for row in _regional_rows(rows)], ["COTIA", "SANTANA DE PARNAIBA"])

    def test_city_grouping_uses_one_official_spelling(self):
        rows = [{"cidade": "Cotia", "nome_completo": "A"}, {"cidade": "COTIA", "nome_completo": "B"}]
        self.assertEqual([row["cidade"] for row in _regional_rows(rows)], ["COTIA", "COTIA"])

    @patch("ColorAdminApp.modules.darpe.common_catalog", return_value=[{
        "comum": "BR-22-0673 - VILA DOUTOR CARDOSO", "cidade": "ITAPEVI",
    }])
    def test_common_uses_regional_br_code_without_changing_legacy_value(self, _catalog):
        result = _with_canonical_commons([{
            "comum_congregacao": "VILA DOUTOR CARDOSO", "cidade": "ITAPEVI",
        }])[0]
        self.assertEqual(result["comum_congregacao"], "VILA DOUTOR CARDOSO")
        self.assertEqual(result["comum_exibicao"], "BR-22-0673 - VILA DOUTOR CARDOSO")

    def test_sector_has_own_module(self):
        self.assertEqual(primary_module({"sector": "D.A.R.P.E"}), MODULE_DARPE)
        with patch("ColorAdminApp.module_access.explicit_modules", return_value=set()):
            self.assertEqual(allowed_modules(request_for("/darpe/")), {MODULE_DARPE})

    def test_page_renders_cadastros(self):
        response = page(request_for("/darpe/cadastros/"), "cadastros")
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Dados do D.A.R.P.E")
        self.assertContains(response, "Importar")
        self.assertContains(response, "select2.full.min.js")

    def test_import_requires_workbook(self):
        response = import_cadastros(request_for("/darpe/api/cadastros/importar/", "post"))
        self.assertEqual(response.status_code, 400)
        self.assertIn("planilha oficial", response.content.decode("utf-8"))

    @patch("ColorAdminApp.modules.darpe.requests.get")
    def test_excel_uses_institutional_header_and_phone_format(self, get):
        from io import BytesIO
        from openpyxl import load_workbook
        response_mock = Mock()
        response_mock.raise_for_status.return_value = None
        response_mock.json.return_value = [{"id": "x", "nome_completo": "Pessoa Teste",
            "comum_congregacao": "CENTRAL", "cidade": "ITAPEVI", "ddd": "11",
            "telefone": "912345678", "status": "ATIVO", "darpe_atendente": True}]
        get.return_value = response_mock
        response = export_cadastros_excel(request_for("/darpe/api/cadastros/exportar-excel/"))
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet["A1"].value, "CONGREGAÇÃO CRISTÃ NO BRASIL")
        self.assertEqual(sheet["A3"].value, "D.A.R.P.E")
        self.assertEqual(sheet["D7"].value, "(11) 9 1234-5678")
        self.assertEqual(sheet.freeze_panes, "A7")
        self.assertFalse(sheet.sheet_view.showGridLines)

    @patch("ColorAdminApp.modules.darpe.requests.post")
    def test_cadastro_records_server_side_authorship(self, post):
        response_mock = Mock()
        response_mock.raise_for_status.return_value = None
        response_mock.json.return_value = [{"id": "saved", "nome_completo": "Pessoa Teste"}]
        post.return_value = response_mock
        response = api_resource(request_for("/darpe/api/cadastros/", "post", {
            "nome_completo": "Pessoa Teste", "comum_congregacao": "CENTRAL DE ITAPEVI", "status": "ATIVO",
            "created_by_name": "Nome forjado",
        }), "cadastros")
        self.assertEqual(response.status_code, 201)
        saved = post.call_args.kwargs["json"]
        self.assertEqual(saved["created_by_name"], "Operador DARPE")
        self.assertEqual(saved["updated_by_user_id"], "user-1")

    def test_spreadsheet_fields_are_covered(self):
        expected = {"cpf", "nome_completo", "instrumento_oficializado", "rg", "data_nascimento", "sexo",
                    "estado_civil", "comum_congregacao", "data_batismo", "logradouro", "numero", "complemento",
                    "bairro", "cidade", "uf", "cep", "ddd", "telefone", "darpe_ministerio", "darpe_atendente",
                    "darpe_colaborador", "darpe_encarregado", "darpe_musico_organista", "darpe_grupo_canto",
                    "declaracao_local", "declaracao_data", "nome_ministerio", "nome_encarregado_orquestra",
                    "nome_voluntario"}
        self.assertTrue(expected.issubset(RESOURCES["cadastros"]["fields"]))
