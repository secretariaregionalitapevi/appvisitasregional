from datetime import datetime
from io import BytesIO
from unittest.mock import patch
from zoneinfo import ZoneInfo

from django.http import QueryDict
from django.test import SimpleTestCase
from openpyxl import load_workbook

from ColorAdminApp.modules.darpe_sam import filtered, reconcile
from ColorAdminApp.modules.darpe_sam_exports import build_excel, build_pdf, export, report_scope
from ColorAdminApp.tests.tests_darpe import request_for
from ColorAdminApp.tests.tests_darpe_sam import CATALOG, SCOPE, source, target


def example_rows(count=1):
    row = reconcile([target()], [source()], CATALOG, SCOPE)[0][0]
    return [{**row, 'nome': f'Pessoa Teste {number:03d}', 'cidade': 'ITAPEVI' if number % 2 else 'COTIA'}
            for number in range(count)]


class DarpeSamReportTests(SimpleTestCase):
    def test_multiple_municipalities_and_active_filter(self):
        rows = example_rows(3) + [{**example_rows()[0], 'cidade': 'JANDIRA'}]
        rows[1]['darpe_status'] = 'INATIVO'
        params = QueryDict('cidade=ITAPEVI&cidade=COTIA&darpe_status=ATIVO')
        result = filtered(rows, params)
        self.assertEqual(len(result), 2)
        self.assertEqual({r['cidade'] for r in result}, {'COTIA'})
        self.assertIn('ITAPEVI, COTIA', report_scope(params))

    def test_excel_institutional_header_complete_rows_and_print_settings(self):
        now = datetime(2026, 10, 21, 14, 30, tzinfo=ZoneInfo('America/Sao_Paulo'))
        rows = example_rows(125)
        data = build_excel(rows, 'ITAPEVI e COTIA', 'Operador Teste', now)
        sheet = load_workbook(BytesIO(data)).active
        self.assertEqual(sheet['A1'].value, 'CONGREGAÇÃO CRISTÃ NO BRASIL')
        self.assertEqual(sheet['A2'].value, 'Regional Itapevi - São Paulo')
        self.assertIn('D.A.R.P.E', sheet['A3'].value)
        self.assertIn('21/10/2026 14:30', sheet['F4'].value)
        self.assertIn('Operador Teste', sheet['F4'].value)
        self.assertEqual(sheet.max_row, 131)
        self.assertEqual(sheet['A131'].value, 'Pessoa Teste 124')
        self.assertEqual(sheet['A6'].fill.fgColor.rgb[-6:], '1E4B7A')
        self.assertEqual(sheet.freeze_panes, 'A7')
        self.assertEqual(sheet.auto_filter.ref, 'A6:I131')
        self.assertFalse(sheet.sheet_view.showGridLines)
        self.assertEqual(sheet.page_setup.fitToWidth, 1)
        self.assertEqual(sheet.print_title_rows, '$1:$6')
        self.assertNotIn('GEM', ' '.join(str(c.value or '') for row in sheet for c in row))

    def test_excel_preserves_names_as_text(self):
        rows = example_rows()
        rows[0]['nome'] = '=SUM(1,2)'
        data = build_excel(rows, 'Todos', 'Teste', datetime.now())
        cell = load_workbook(BytesIO(data)).active['A7']
        self.assertEqual(cell.value, '=SUM(1,2)')
        self.assertEqual(cell.data_type, 's')

    def test_pdf_generates_multiple_pages_in_landscape(self):
        data = build_pdf(example_rows(125), 'ITAPEVI e COTIA', 'Operador Teste', datetime.now())
        self.assertTrue(data.startswith(b'%PDF-'))
        self.assertGreater(data.count(b'/Type /Page'), 2)

    @patch('ColorAdminApp.modules.darpe_sam_exports.read_comparison')
    def test_export_ignores_screen_pagination_and_respects_all_filters(self, read):
        read.return_value = (example_rows(125), None)
        request = request_for('/darpe/api/conferencia-sam/exportar-excel/?pagina=2&cidade=ITAPEVI&cidade=COTIA')
        response = export(request, 'xlsx')
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 131)
        self.assertIn('ITAPEVI, COTIA', sheet['A4'].value)
        self.assertIn('Conferencia_DARPE_SAM_', response['Content-Disposition'])
        self.assertIn('no-store', response['Cache-Control'])

    @patch('ColorAdminApp.modules.darpe_sam_exports.read_comparison', return_value=([], None))
    def test_empty_export_has_clear_message(self, read):
        response = export(request_for('/darpe/api/conferencia-sam/exportar-pdf/'), 'pdf')
        self.assertEqual(response.status_code, 400)

    @patch('ColorAdminApp.modules.darpe_sam_exports.can_access_module', return_value=False)
    @patch('ColorAdminApp.modules.darpe_sam_exports.read_comparison')
    def test_export_checks_access_before_reading(self, read, _allowed):
        self.assertEqual(export(request_for('/darpe/api/conferencia-sam/exportar-excel/'), 'xlsx').status_code, 403)
        read.assert_not_called()

    def test_page_shows_badges_multiselect_exports_and_clear_sam_reference(self):
        from ColorAdminApp.modules.darpe_sam import page
        response = page(request_for('/darpe/conferencia-sam/'))
        self.assertContains(response, 'multiple data-placeholder="Todos os municípios"')
        self.assertContains(response, 'ds-status-active')
        self.assertContains(response, 'id="ds-excel"')
        self.assertContains(response, 'id="ds-pdf"')
        self.assertContains(response, 'Sistema Administrativo da Música')
        self.assertNotContains(response, 'Cadastro SAM no GEM')
        self.assertNotContains(response, 'vínculo único no SAM e no DARPE')
