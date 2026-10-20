import io
import json
from datetime import datetime
from unittest.mock import Mock,patch
from django.test import SimpleTestCase,RequestFactory,override_settings
from openpyxl import load_workbook
from pypdf import PdfReader
from ColorAdminApp.modules.music_events_conductors import summarize
from ColorAdminApp.modules.music_events_dashboard import aggregate
from ColorAdminApp.modules.music_events_reports import build_excel,build_pdf
from ColorAdminApp.modules import music_events


class ConductorSummaryTests(SimpleTestCase):
    def fixture(self):
        events=[{'id':'e','titulo':'Ensaio Regional','data_evento':'2026-09-20','tipo':'ENSAIO_REGIONAL','familia':'GERAL','status':'ABERTO'}]
        sites=[{'id':'l','evento_id':'e','nome':'CAUCAIA DO ALTO'}]
        records=[{'evento_id':'e','local_id':'l','nome':'Encarregado de teste','comum_informada':'BR-22-1 - A','municipio_informado':'COTIA','cargo':'ENCARREGADO LOCAL','categoria':'MUSICO'},
                 {'evento_id':'e','local_id':'l','comum_informada':'BR-22-2 - B','municipio_informado':'COTIA','categoria':'MUSICO'}]
        return events,records,sites

    def test_infers_present_location_but_never_nj(self):
        e,r,l=self.fixture();data=summarize(e,r,l,[],{})
        self.assertEqual(data['rows'][0]['registro'],'CAUCAIA DO ALTO')
        self.assertEqual(data['rows'][1]['situacao'],'NAO_INFORMADO')
        self.assertEqual(data['totals'].get('NAO_JUSTIFICOU',0),0)

    def test_registered_leader_in_nominal_list_is_present_even_without_imported_role(self):
        e,r,l=self.fixture();r[0]['cargo']='MÚSICO'
        leaders=[{'nome_aluno':'Encarregado de teste','cargo_ministerio':'ENCARREGADO LOCAL',
                  'comum_congregacao':'BR-22-1 - A','municipio':'COTIA','status':'Ativo'}]
        data=summarize(e,r,l,[],{},catalog=[{'comum':'BR-22-1 - A','cidade':'COTIA'}],leaders=leaders)
        row=next(x for x in data['rows'] if x['comum_chave']=='BR-22-1')
        self.assertEqual(row['situacao'],'PRESENTE')
        self.assertEqual(row['registro'],'CAUCAIA DO ALTO')

    def test_registered_name_resolves_unrecognized_itapevi_common(self):
        e,r,l=self.fixture();r[0].update(cargo='MÚSICO',comum_informada='ITACOLMY',municipio_informado='ITAPEVI',nome='Claudio')
        catalog=[{'comum':'BR-22-3270 - JARDIM ITACOLOMI','cidade':'ITAPEVI'}]
        leaders=[{'nome_aluno':'Claudio','cargo_ministerio':'ENCARREGADO LOCAL',
                  'comum_congregacao':'BR-22-3270 - JARDIM ITACOLOMI','municipio':'ITAPEVI','status':'Ativo'}]
        data=summarize(e,r,l,[],{},catalog=catalog,leaders=leaders)
        row=next(x for x in data['rows'] if x['comum_chave']=='BR-22-3270')
        self.assertEqual(row['situacao'],'PRESENTE')

    def test_manual_status_does_not_erase_nominal_presence(self):
        e,r,l=self.fixture()
        review={'evento_id':'e','comum_chave':'BR-22-1','comum_nome':'A','municipio':'COTIA',
                'situacao':'NAO_INFORMADO','observacoes':'','local_id':None}
        data=summarize(e,r,l,[review],{})
        self.assertEqual(data['rows'][0]['situacao'],'PRESENTE')

    def test_manual_justification_can_include_common_with_no_attendance(self):
        e,r,l=self.fixture();review={'evento_id':'e','comum_chave':'BR-22-3','comum_nome':'Comum C','municipio':'COTIA','situacao':'JUSTIFICOU','observacoes':'Trabalho','local_id':None}
        data=summarize(e,r,l,[review],{})
        self.assertEqual(data['totals']['JUSTIFICOU'],1)
        self.assertEqual(data['municipalities'][0]['JUSTIFICOU'],1)
        self.assertEqual(next(x for x in data['rows'] if x['comum_chave']=='BR-22-3')['registro'],'J · Justificou')

    def test_reports_reference_site_and_presence_at_another_rehearsal(self):
        e,r,l=self.fixture();l.append({'id':'other','evento_id':'e','nome':'ITAPEVI'});r[0]['local_id']='other'
        frequency=[{'evento_id':'e','comum_codigo':'BR-22-1','local_referencia_id':'l'}]
        data=summarize(e,r,l,[],{},frequency=frequency)
        row=next(x for x in data['rows'] if x['comum_chave']=='BR-22-1')
        self.assertEqual(row['local_referencia'],'CAUCAIA DO ALTO')
        self.assertTrue(row['em_outro_local'])
        self.assertTrue(row['fora_municipio'])
        self.assertEqual(row['participacao'],'FORA DO MUNICÍPIO')
        self.assertEqual(row['registro'],'Outro local · ITAPEVI')
        self.assertEqual(data['municipalities'][0]['PRESENTE_OUTRO'],1)

    def test_every_convoked_common_appears_even_without_attendance(self):
        e,r,l=self.fixture()
        frequency=[{'evento_id':'e','comum_codigo':'BR-22-3','comum':'Comum C','municipio':'COTIA','local_referencia_id':'l'}]
        data=summarize(e,[],l,[],{},frequency=frequency)
        self.assertEqual(len(data['rows']),1)
        self.assertEqual(data['rows'][0]['situacao'],'NAO_INFORMADO')
        self.assertEqual(data['rows'][0]['local_referencia'],'CAUCAIA DO ALTO')

    def test_period_counts_common_per_event_not_unique_people(self):
        e,r,l=self.fixture();e.append(dict(e[0],id='second'));r.append(dict(r[0],evento_id='second'))
        d=summarize(e,r,l,[],{})
        self.assertEqual(d['totals']['PRESENTE'],2)

    def test_conductor_report_contains_presence_and_site_only(self):
        e,r,l=self.fixture();review={'evento_id':'e','comum_chave':'BR-22-2','comum_nome':'B','municipio':'COTIA','situacao':'JUSTIFICOU','observacoes':'Trabalho','local_id':None}
        data=aggregate(e,r,l,[],{});data['conductors']=summarize(e,r,l,[review],{});data['report']='encarregados'
        book=load_workbook(io.BytesIO(build_excel(data,'Operador',datetime(2026,10,19,12,0))))
        self.assertNotIn('Encarregados por município',book.sheetnames)
        ws=book['Encarregados por comum'];statuses={ws.cell(row,4).value:ws.cell(row,5).value for row in range(10,ws.max_row+1)}
        self.assertEqual(statuses['PRESENTE'],'FORA DO MUNICÍPIO')
        self.assertIn('JUSTIFICOU',statuses)
        text=''.join(p.extract_text() for p in PdfReader(io.BytesIO(build_pdf(data,'Operador',datetime(2026,10,19,12,0)))).pages)
        self.assertNotIn('Trabalho',text);self.assertIn('CAUCAIA DO ALTO',text);self.assertIn('PRESENTE',text);self.assertIn('JUSTIFICOU',text)


@override_settings(SUPABASE_URL='https://db.example',SUPABASE_SERVICE_ROLE_KEY='test')
class ConductorApiTests(SimpleTestCase):
    def request(self,status='JUSTIFICOU',notes='Trabalho'):
        r=RequestFactory().post('/musica/eventos/api/eventos/e/encarregado/',data=json.dumps({'comum_chave':'BR-22-1','situacao':status,'observacoes':notes,'created_by_name':'Forjado'}),content_type='application/json')
        r.session={'authenticated':True,'user_id':'actor','user_profile':{'role_id':1,'full_name':'Operador'}};return r

    @patch.object(music_events,'_allowed',return_value=True)
    @patch.object(music_events,'_visible_events',return_value=[{'id':'e','status':'ABERTO'}])
    @patch.object(music_events,'common_catalog',return_value=[{'comum':'BR-22-1 - A','cidade':'COTIA'}])
    @patch.object(music_events,'_read',return_value=[])
    @patch.object(music_events.requests,'post')
    def test_authenticated_actor_and_justification(self,post,read,catalog,events,allowed):
        post.return_value=Mock();post.return_value.json.return_value=[{'id':'saved'}]
        response=music_events.api_conductor(self.request(),'e')
        self.assertEqual(response.status_code,200)
        self.assertEqual(post.call_args.kwargs['json']['created_by_name'],'Operador')
        self.assertEqual(post.call_args.kwargs['json']['local_id'],None)

    @patch.object(music_events,'_allowed',return_value=True)
    @patch.object(music_events,'_visible_events',return_value=[{'id':'e','status':'ABERTO'}])
    @patch.object(music_events,'common_catalog',return_value=[{'comum':'BR-22-1 - A','cidade':'COTIA'}])
    @patch.object(music_events.requests,'post')
    def test_justification_requires_note(self,post,catalog,events,allowed):
        self.assertEqual(music_events.api_conductor(self.request(notes=''),'e').status_code,400);post.assert_not_called()

    @patch.object(music_events,'_allowed',return_value=True)
    @patch.object(music_events,'_visible_events',return_value=[{'id':'e','status':'ENCERRADO'}])
    @patch.object(music_events.requests,'post')
    def test_closed_event_is_immutable(self,post,events,allowed):
        self.assertEqual(music_events.api_conductor(self.request(),'e').status_code,409);post.assert_not_called()
