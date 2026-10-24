import io
from datetime import datetime
from unittest.mock import patch

from django.test import SimpleTestCase
from openpyxl import load_workbook
from pypdf import PdfReader

from ColorAdminApp.modules.music_events_dashboard import aggregate
from ColorAdminApp.modules.music_events_reports import build_excel, build_pdf


class DashboardTests(SimpleTestCase):
    def fixture(self):
        events=[{'id':'e1','titulo':'Ensaio regional','data_evento':'2026-09-20','tipo':'ENSAIO_REGIONAL','familia':'GERAL','base_id':None,'status':'ABERTO'},
                {'id':'e2','titulo':'RT de Cordas','data_evento':'2026-08-29','tipo':'REUNIAO_TECNICA','familia':'CORDAS','base_id':None,'status':'ABERTO'}]
        sites=[{'id':'l1','evento_id':'e1','nome':'ITAPEVI'},{'id':'l2','evento_id':'e1','nome':'PIRAPORA'},{'id':'l3','evento_id':'e2','nome':'COTIA'}]
        def row(event,local,city,common,category):
            return {'evento_id':event,'local_id':local,'municipio_informado':city,'comum_informada':common,
                    'categoria':category,'situacao':'PRESENTE','familia':'CORDAS','instrumento':'VIOLINO'}
        rows=[row('e1','l1','ITAPEVI','BR-22-1 - Comum A','MUSICO'),
              row('e1','l2','ITAPEVI','BR-22-1 - Comum A','ORGANISTA'),
              row('e1','l2','COTIA','BR-22-2 - Comum B','MUSICO'),
              row('e1','l1','OSASCO','Visitante','IRMANDADE'),
              row('e2','l3','COTIA','BR-22-2 - Comum B','MUSICO')]
        return events,rows,sites

    def test_city_origin_is_independent_of_meeting_location(self):
        events,rows,sites=self.fixture()
        data=aggregate(events,rows,sites,[],{'evento':'e1','municipios':['ITAPEVI'],'local':'l2'})
        self.assertEqual(data['totals']['total'],1)
        self.assertEqual(data['totals']['ORGANISTA'],1)
        self.assertEqual(data['locations'][0]['nome'],'PIRAPORA')
        self.assertFalse(data['base_complete'])
        self.assertIsNone(data['rows'][0]['percentual_musicos'])

    def test_common_report_identifies_attendance_at_another_regional_site(self):
        events,rows,sites=self.fixture()
        frequency=[{'evento_id':'e1','comum_codigo':'BR-22-1','comum':'Comum A','municipio':'ITAPEVI',
                    'local_referencia_id':'l1','base_validada':False,'musicos_base':0,'organistas_base':0}]
        data=aggregate(events,rows,sites,frequency,{'evento':'e1','municipios':['ITAPEVI']})
        common=data['rows'][0]
        self.assertEqual(common['local_referencia'],'ITAPEVI')
        self.assertEqual(common['musicos_outro_local'],0)
        self.assertEqual(common['organistas_outro_local'],1)
        self.assertEqual(common['outros_locais'],'PIRAPORA')

    def test_rt_family_and_period_filters(self):
        e,r,l=self.fixture()
        d=aggregate(e,r,l,[],{'tipo':'RT_CORDAS','ano':'2026','mes':'08'})
        self.assertEqual(d['totals']['total'],1)
        self.assertEqual(d['events'][0]['id'],'e2')

    def test_regional_and_visitors_are_separate(self):
        e,r,l=self.fixture()
        d=aggregate(e,r,l,[],{'evento':'e1','origem':'visitantes'})
        self.assertEqual(d['totals']['total'],1)
        self.assertEqual(d['rows'][0]['municipio'],'OSASCO')

    def test_organists_are_people_not_physical_organs(self):
        e,r,l=self.fixture();r[1]['instrumento']='ÓRGÃO'
        d=aggregate(e,r,l,[],{'evento':'e1'})
        self.assertEqual(d['instruments']['Organistas'],1)
        self.assertNotIn('ÓRGÃO',d['instruments'])
        self.assertEqual(d['totals']['ORGANISTA'],1)
        ws=load_workbook(io.BytesIO(build_excel(d,'Teste',datetime(2026,10,19,12))))['Instrumentos']
        self.assertEqual(ws['B9'].value,'Participantes presentes')

    def test_default_prioritizes_seven_municipalities_including_exports(self):
        e,r,l=self.fixture();d=aggregate(e,r,l,[],{'evento':'e1'})
        self.assertEqual(d['totals']['total'],3)
        self.assertEqual(len(d['regional_overview']),7)
        self.assertNotIn('OSASCO',d['options']['cities'])
        self.assertFalse(any(x['municipio']=='OSASCO' for x in d['rows']))
        self.assertEqual(d['origin_context']['visitantes'],1)
        book=load_workbook(io.BytesIO(build_excel(d,'Teste',datetime(2026,10,19,12))))
        self.assertNotIn('OSASCO',[c.value for row in book['Frequência por comum'] for c in row])
        self.assertEqual(aggregate(e,r,l,[],{'evento':'e1','origem':'todos'})['totals']['total'],4)

    def test_regional_seat_aliases_are_not_visitors(self):
        e,r,l=self.fixture()
        r=[dict(r[0],municipio_informado=city,comum_informada=str(i))
           for i,city in enumerate(['Fazendinha','Caucaia','Parque Laranjeiras II','VGP'])]
        d=aggregate(e,r,l,[],{'evento':'e1'})
        self.assertEqual(d['totals']['total'],4)
        self.assertEqual(d['totals']['municipios'],4)
        self.assertIn('SANTANA DE PARNAÍBA',[x['municipio'] for x in d['rows']])

    def test_ministry_role_overlaps_music_without_double_counting(self):
        e,r,l=self.fixture();r[0]['cargo']='ANCIÃO';r[2]['cargo']='ENCARREGADO LOCAL'
        d=aggregate(e,r,l,[],{'evento':'e1'})
        self.assertEqual(d['totals']['MUSICO'],2)
        self.assertEqual(d['totals']['ministerio_cargos'],1)
        self.assertEqual(d['totals']['encarregados'],1)
        self.assertEqual(d['totals']['total'],3)
        self.assertEqual(next(x for x in d['regional_overview'] if x['municipio']=='ITAPEVI')['ministerio_cargos'],1)

    def test_local_conductor_presence_is_evidenced_without_inventing_absence(self):
        e,r,l=self.fixture();r[0]['cargo']='ENCARREGADO LOCAL'
        d=aggregate(e,r,l,[],{'evento':'e1'})
        self.assertEqual(next(x for x in d['rows'] if x['comum'].startswith('BR-22-1'))['encarregado'],'Presente')
        self.assertEqual(next(x for x in d['rows'] if x['comum'].startswith('BR-22-2'))['encarregado'],'Não informado')

    def test_canonical_common_prevents_city_typo_and_alias_split(self):
        e,r,l=self.fixture();r[0]['municipio_informado']='ERRADO'
        d=aggregate(e,r,l,[],{'municipios':['ITAPEVI']},[{'comum':'BR-22-1 - Comum canônica','cidade':'ITAPEVI'}])
        self.assertEqual(d['totals']['total'],2)
        self.assertEqual(d['rows'][0]['comum'],'BR-22-1 - Comum canônica')

    def test_percentages_use_eligible_matches_and_weighted_population(self):
        e,r,l=self.fixture();e=e[:1];e[0]['base_id']='base'
        r[0]['sam_pessoa_id']='sam1';r[2]['sam_pessoa_id']='sam2'
        f=[{'evento_id':'e1','comum_codigo':'BR-22-1','comum':'A','municipio':'ITAPEVI','base_validada':True,'musicos_base':2,'organistas_base':0},
           {'evento_id':'e1','comum_codigo':'BR-22-2','comum':'B','municipio':'COTIA','base_validada':True,'musicos_base':90,'organistas_base':0}]
        people=[{'evento_id':'e1','sam_pessoa_id':'sam1','comum_codigo':'BR-22-1','categoria':'MUSICO'},
                {'evento_id':'e1','sam_pessoa_id':'sam2','comum_codigo':'BR-22-2','categoria':'MUSICO'}]
        d=aggregate(e,r,l,f,{},eligible=people)
        self.assertEqual(d['percent_m'],2.17) # 2/92, não a média das taxas
        self.assertIsNone(d['percent_o'])

    def test_excel_and_pdf_use_institutional_header(self):
        e,r,l=self.fixture();r[0]['comum_informada']='=HYPERLINK("bad")'
        d=aggregate(e,r,l,[],{'evento':'e1'})
        now=datetime(2026,10,19,12,30)
        book=load_workbook(io.BytesIO(build_excel(d,'Operador de teste',now)))
        ws=book['Frequência por comum']
        self.assertEqual(ws['A1'].value,'CONGREGAÇÃO CRISTÃ NO BRASIL')
        self.assertEqual(ws['A9'].fill.fgColor.rgb[-6:],'1E4B7A')
        self.assertEqual(ws.freeze_panes,'A10')
        self.assertTrue(all(c.data_type!='f' for row in ws for c in row))
        pdf=PdfReader(io.BytesIO(build_pdf(d,'Operador de teste',now)))
        text='\n'.join(p.extract_text() for p in pdf.pages)
        self.assertIn('CONGREGAÇÃO CRISTÃ NO BRASIL',text)
        self.assertIn('19/10/2026 12:30',text)
        self.assertIn('Operador de teste',text)
        self.assertNotIn('2026-09-20',text)

    def test_long_frequency_table_starts_on_first_page(self):
        e,r,l=self.fixture()
        rows=[dict(r[0],comum_informada=f'Comum de teste {i:03}') for i in range(80)]
        d=aggregate(e,rows,l,[],{'evento':'e1'})
        pdf=PdfReader(io.BytesIO(build_pdf(d,'Operador de teste',datetime(2026,10,19,12,30))))
        self.assertGreater(len(pdf.pages),1)
        self.assertIn('Frequência por comum',pdf.pages[0].extract_text())
        self.assertIn('Comum de teste 000',pdf.pages[0].extract_text())

    def test_presence_pdf_paginates_many_commons_from_same_municipality(self):
        e,r,l=self.fixture()
        rows=[dict(r[0],comum_informada=f'Comum de teste {i:03}',municipio_informado='ITAPEVI') for i in range(80)]
        d=aggregate(e,rows,l,[],{'evento':'e1'});d['report']='comuns'
        pdf=PdfReader(io.BytesIO(build_pdf(d,'Operador de teste',datetime(2026,10,20,12,30))))
        self.assertGreater(len(pdf.pages),1)
        self.assertIn('RELATÓRIO UNIFICADO ENSAIO REGIONAL',pdf.pages[0].extract_text())

    def test_executive_pdf_is_one_page_even_with_hundreds_of_commons(self):
        from ColorAdminApp.modules.music_events_dashboard import REGIONAL
        e,r,l=self.fixture()
        records=[dict(r[0],comum_informada=f'Comum de teste {i:03}',municipio_informado=REGIONAL[i%7]) for i in range(350)]
        d=aggregate(e,records,l,[],{'evento':'e1'})
        d['report']='executivo'
        d['report_identity']={'scope':'Global · Itapevi','common':'BR-22-0673 - VILA DOUTOR CARDOSO'}
        pdf=PdfReader(io.BytesIO(build_pdf(d,'Operador de teste',datetime(2026,10,19,12,30))))
        self.assertEqual(len(pdf.pages),1)
        text=pdf.pages[0].extract_text()
        self.assertIn('Resumo executivo de presença',text)
        self.assertIn('TOTAL DO RECORTE',text)
        self.assertIn('350',text)
        self.assertNotIn('Comum de teste 000',text)
        self.assertNotIn('Cadastro atual',text)

    def test_pdf_endpoint_defaults_to_frequency_report(self):
        from django.test import RequestFactory
        from ColorAdminApp.modules.music_events_reports import export_response
        e,r,l=self.fixture();d=aggregate(e,r,l,[],{'evento':'e1'})
        req=RequestFactory().get('/musica/eventos/exportar-pdf/')
        req.session={'user_profile':{'full_name':'Teste','role_id':1}}
        response=export_response(d,req,'pdf')
        self.assertIn('PRESENCA_EnR_COTIA_ITAPEVI_SET26.pdf',response['Content-Disposition'])
        text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
        self.assertIn('RELATÓRIO UNIFICADO ENSAIO REGIONAL',text)
        from urllib.parse import unquote
        d['filters']['municipios'] = ['ITAPEVI']
        d['events'][0]['data_evento'] = '2026-05-17'
        renamed = export_response(d,req,'pdf')
        self.assertIn("filename*=UTF-8''[RELATÓRIO] PRESENÇA_EnR_ITAPEVI_MAI26.pdf", unquote(renamed['Content-Disposition']))
        req=RequestFactory().get('/musica/eventos/exportar-pdf/?relatorio=comuns')
        req.session={'user_profile':{'full_name':'Teste','role_id':1}}
        response=export_response(d,req,'pdf')
        text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
        self.assertIn('RELATÓRIO UNIFICADO ENSAIO REGIONAL',text)
        self.assertNotIn('Instrumentos',text)
        self.assertNotIn('Encarregados por município',text)
