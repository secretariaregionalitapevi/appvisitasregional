import io
from datetime import datetime
from django.test import SimpleTestCase
from pypdf import PdfReader
from ColorAdminApp.modules.music_events_leaders import attach_leaders
from ColorAdminApp.modules.music_events_dashboard import aggregate
from ColorAdminApp.modules.music_events_reports import build_pdf,build_excel
from openpyxl import load_workbook


class RegisteredLeaderTests(SimpleTestCase):
    def data(self):
        catalog=[{'comum':'BR-22-2797 - LAGEADO - COTIA','cidade':'COTIA'}]
        e=[{'id':'e','titulo':'Ensaio Regional','data_evento':'2026-09-20','tipo':'ENSAIO_REGIONAL','familia':'GERAL','status':'ABERTO'}]
        r=[{'evento_id':'e','local_id':'l','comum_informada':'LAJEADO','municipio_informado':'COTIA','categoria':'MUSICO'}]
        d=aggregate(e,r,[{'id':'l','evento_id':'e','nome':'ITAPEVI'}],[],{},catalog)
        person={'nome_aluno':'Responsável cadastrado','cargo_ministerio':'ENCARREGADO LOCAL','comum_congregacao':'LAGEADO','municipio':'COTIA','status':'Concluido'}
        return d,catalog,person

    def test_name_is_identified_without_inventing_attendance(self):
        d,c,p=self.data();attach_leaders(d,[p],c)
        self.assertEqual(d['rows'][0]['encarregado_nome'],'Responsável cadastrado')
        self.assertEqual(d['rows'][0]['encarregado'],'Não informado')

    def test_ambiguous_names_remain_visible_not_arbitrarily_selected(self):
        d,c,p=self.data();attach_leaders(d,[p,dict(p,nome_aluno='Outro responsável')],c)
        self.assertEqual(len(d['rows'][0]['encarregado_nomes']),2)

    def test_regional_or_inactive_role_is_not_local_conductor(self):
        d,c,p=self.data();attach_leaders(d,[dict(p,cargo_ministerio='ENCARREGADO REGIONAL'),dict(p,status='Inativo')],c)
        self.assertEqual(d['rows'][0]['encarregado_nomes'],[])

    def test_irmandade_is_fallback_and_sam_has_priority(self):
        d,c,p=self.data()
        member={'nome':'Cadastro da irmandade','cargo_outros':'Encarregado Local,Grupo de Visitas','comum':'LAGEADO','cidade':'COTIA','status':'Ativo'}
        attach_leaders(d,[],c,[member]);self.assertEqual(d['rows'][0]['encarregado_nome'],'Cadastro da irmandade')
        attach_leaders(d,[p],c,[member]);self.assertEqual(d['rows'][0]['encarregado_nome'],'Responsável cadastrado')

    def test_registered_name_is_in_common_pdf_and_excel(self):
        d,c,p=self.data();attach_leaders(d,[p],c);d['report']='comuns'
        pdf=PdfReader(io.BytesIO(build_pdf(d,'Operador',datetime(2026,10,19,12))))
        self.assertIn('RELATÓRIO UNIFICADO ENSAIO REGIONAL',pdf.pages[0].extract_text())
        ws=load_workbook(io.BytesIO(build_excel(d,'Operador',datetime(2026,10,19,12))))['Frequência por comum']
        self.assertIn('Responsável cadastrado',ws['C10'].value)
