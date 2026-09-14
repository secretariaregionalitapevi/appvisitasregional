from datetime import date
from unittest.mock import patch

from django.test import RequestFactory, SimpleTestCase
from openpyxl import load_workbook

from . import gem_projection


class GemProjectionTests(SimpleTestCase):
    def setUp(self):
        self.request = RequestFactory().get("/gem/projecao/")
        self.request.session = {"user_profile": {"role_id": 1, "full_name": "Responsável Teste"}}

    def test_ready_requires_program_attendance_and_absence_limit(self):
        ready, inconsistent, status = gem_projection._classification(100, 85, 3, "CANDIDATO(A)")
        self.assertTrue(ready)
        self.assertFalse(inconsistent)
        self.assertEqual(status, "APTO PARA AVALIAÇÃO")

    def test_fourth_absence_triggers_priority(self):
        ready, _, status = gem_projection._classification(100, 80, 4, "CANDIDATO(A)")
        self.assertFalse(ready)
        self.assertEqual(status, "ACOMPANHAMENTO PRIORITÁRIO")

    def test_projection_never_exceeds_two_sessions_per_week(self):
        projected = gem_projection._projection_date(30, [date(2026, 1, 1), date(2026, 4, 1)])
        self.assertGreaterEqual((projected - date.today()).days, 15 * 7)

    @patch("ColorAdminApp.gem_projection.can_open_module", return_value=True)
    @patch("ColorAdminApp.gem_projection._filtered", return_value=[])
    def test_excel_uses_institutional_identity_and_report_layout(self, _filtered, _access):
        response = gem_projection.export_excel(self.request)
        workbook = load_workbook(filename=__import__("io").BytesIO(response.content))
        self.assertEqual(workbook.sheetnames, ["Resumo por etapa", "Projeção"])
        self.assertEqual(workbook["Resumo por etapa"]["A1"].value, "CONGREGAÇÃO CRISTÃ NO BRASIL")
        self.assertFalse(workbook["Resumo por etapa"].sheet_view.showGridLines)
        sheet = workbook["Projeção"]
        self.assertEqual(sheet["A1"].value, "CONGREGAÇÃO CRISTÃ NO BRASIL")
        self.assertEqual(sheet["A2"].value, "Regional Itapevi - São Paulo")
        self.assertEqual(sheet.freeze_panes, "A8")
        self.assertFalse(sheet.sheet_view.showGridLines)
        self.assertEqual(sheet.auto_filter.ref, "A7:R7")

    def test_stage_mapping_follows_next_ministerial_step(self):
        self.assertEqual(gem_projection._stage_key({"nivel": "CANDIDATO(A)"}), "entrada_rjm")
        self.assertEqual(gem_projection._stage_key({"nivel": "RJM / ENSAIO"}), "rjm_culto")
        self.assertEqual(gem_projection._stage_key({"nivel": "CULTO OFICIAL"}), "culto_oficializacao")
