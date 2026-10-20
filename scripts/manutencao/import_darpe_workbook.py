"""Importa o cadastro consolidado do D.A.R.P.E para a tabela UUID canônica.

Uso:
  python scripts/manutencao/import_darpe_workbook.py CAMINHO.xlsx
  python scripts/manutencao/import_darpe_workbook.py CAMINHO.xlsx --apply --replace

Por segurança, a carga valida duplicidades, municípios e a presença de registros.
A substituição remove registros antigos somente depois que todas as linhas
oficiais forem gravadas com sucesso.
"""
import argparse
import os
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "ColorAdmin.settings")

import django  # noqa: E402
django.setup()
from django.conf import settings  # noqa: E402
from ColorAdminApp.modules.access_control import _norm, common_catalog, service_headers  # noqa: E402
from ColorAdminApp.modules.music_events_identity import CommonResolver  # noqa: E402
from ColorAdminApp.modules.musicalizacao import REGIONAL_MUNICIPALITIES  # noqa: E402


def clean(value):
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = str(value).strip()
    return text or None


def yes(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char)).strip().upper()
    return text in {"SIM", "S", "TRUE", "1", "X"}


def musical_flags(instrument, sex, selected):
    if not selected:
        return False, False
    organist = _norm(instrument) == "ORGAO" or _norm(sex) == "FEMININO"
    return not organist, organist


def key(row):
    return (str(row.get("nome_completo") or "").strip().upper(),
            str(row.get("comum_congregacao") or "").strip().upper())


def load_rows(path):
    sheet = load_workbook(path, read_only=True, data_only=True)["Cadastro"]
    result = []
    for values in sheet.iter_rows(min_row=4, values_only=True):
        if not clean(values[1]):
            continue
        comum = clean(values[7]) or "NÃO INFORMADA"
        musician, organist = musical_flags(values[2], values[5], yes(values[22]))
        result.append({
            "cpf": clean(values[0]), "nome_completo": clean(values[1]),
            "instrumento_oficializado": clean(values[2]), "rg": clean(values[3]),
            "data_nascimento": clean(values[4]), "sexo": clean(values[5]),
            "estado_civil": clean(values[6]), "comum_congregacao": comum,
            "data_batismo": clean(values[8]), "logradouro": clean(values[9]),
            "numero": clean(values[10]), "complemento": clean(values[11]),
            "bairro": clean(values[12]), "cidade": clean(values[13]),
            "uf": (clean(values[14]) or "").upper() or None, "cep": clean(values[15]),
            "ddd": clean(values[16]), "telefone": clean(values[17]),
            "darpe_ministerio": yes(values[18]), "darpe_atendente": yes(values[19]),
            "darpe_colaborador": yes(values[20]), "darpe_encarregado": yes(values[21]),
            "darpe_musico": musician, "darpe_organista": organist,
            "darpe_musico_organista": musician or organist, "darpe_grupo_canto": yes(values[23]),
            "declaracao_local": clean(values[24]), "declaracao_data": clean(values[25]),
            "nome_ministerio": clean(values[26]), "nome_encarregado_orquestra": clean(values[27]),
            "nome_voluntario": clean(values[28]), "status": "ATIVO",
            "created_by_user_id": "importacao-darpe-planilha",
            "created_by_name": "Importação do mapeamento regional D.A.R.P.E",
            "updated_by_user_id": "importacao-darpe-planilha",
            "updated_by_name": "Importação do mapeamento regional D.A.R.P.E",
        })
    return result


def canonicalize_locations(rows):
    resolver = CommonResolver(common_catalog())
    allowed = {_norm(city) for city in REGIONAL_MUNICIPALITIES}
    unresolved = []
    for row in rows:
        canonical = resolver.resolve(row.get("comum_congregacao"), row.get("cidade"))
        if canonical:
            row["comum_congregacao"] = canonical["comum"]
            row["cidade"] = canonical["cidade"]
        if _norm(row.get("cidade")) not in allowed:
            unresolved.append(row.get("comum_congregacao") or row.get("nome_completo"))
    if unresolved:
        raise RuntimeError(
            f"Carga bloqueada: {len(unresolved)} cadastro(s) sem município regional confirmado. "
            "Nenhum dado foi alterado."
        )
    return rows


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--replace", action="store_true",
                        help="Remove da tabela os registros ausentes da planilha oficial.")
    parser.add_argument("--expected-count", type=int,
                        help="Validação opcional de quantidade, quando houver uma contagem oficial confirmada.")
    args = parser.parse_args()
    rows = load_rows(args.workbook)
    keys = [key(row) for row in rows]
    if not rows:
        print("CARGA BLOQUEADA: a aba Cadastro não possui registros para importar.")
        print("Nenhum dado foi alterado.")
        return 2
    if len(keys) != len(set(keys)):
        print(f"ERRO: a planilha possui {len(keys) - len(set(keys))} cadastro(s) repetido(s).")
        print("Nenhum dado foi alterado.")
        return 2
    if args.expected_count is not None and len(rows) != args.expected_count:
        print(
            f"CARGA BLOQUEADA: a planilha possui {len(rows)} pessoas únicas; "
            f"eram esperadas {args.expected_count}."
        )
        print(f"Faltam {max(0, args.expected_count - len(rows))} cadastro(s) na aba Cadastro.")
        print("Nenhum dado foi alterado.")
        return 2
    rows = canonicalize_locations(rows)
    url = f"{settings.SUPABASE_URL}/rest/v1/darpe_cadastro"
    response = requests.get(url, headers=service_headers(),
                            params={"select": "id,nome_completo,comum_congregacao"}, timeout=30)
    response.raise_for_status()
    current = response.json()
    existing = {}
    for row in current:
        existing.setdefault(key(row), []).append(row)
    pending = [row for row in rows if key(row) not in existing]
    matched = [row for row in rows if key(row) in existing]
    stale = [row for row in current if key(row) not in set(keys)]
    duplicate_count = sum(max(0, len(items) - 1) for items in existing.values())
    print(
        f"Planilha validada: {len(rows)} | atualizar: {len(matched)} | novos: {len(pending)} | "
        f"fora da lista: {len(stale)} | duplicados atuais: {duplicate_count}"
    )
    if not args.apply:
        return
    if not args.replace:
        raise RuntimeError("Para sincronizar a lista oficial, informe também --replace.")

    for start in range(0, len(pending), 50):
        batch = pending[start:start + 50]
        result = requests.post(url, headers=service_headers("return=representation"), json=batch, timeout=60)
        if not result.ok:
            raise RuntimeError(f"Falha no lote {start // 50 + 1}: {result.status_code} {result.text}")
        print(f"Importados: {min(start + len(batch), len(pending))}/{len(pending)}")

    for index, row in enumerate(matched, 1):
        record_id = existing[key(row)][0]["id"]
        result = requests.patch(url, headers=service_headers("return=minimal"),
                                params={"id": f"eq.{record_id}"}, json=row, timeout=30)
        if not result.ok:
            raise RuntimeError(f"Falha ao atualizar {row['nome_completo']}: {result.status_code} {result.text}")
        if index % 25 == 0 or index == len(matched):
            print(f"Atualizados: {index}/{len(matched)}")

    refreshed = requests.get(url, headers=service_headers(),
                             params={"select": "id,nome_completo,comum_congregacao"}, timeout=30)
    refreshed.raise_for_status()
    keep = set()
    remove = []
    official_keys = set(keys)
    for row in refreshed.json():
        row_key = key(row)
        if row_key not in official_keys or row_key in keep:
            remove.append(row["id"])
        else:
            keep.add(row_key)
    for index, record_id in enumerate(remove, 1):
        result = requests.delete(url, headers=service_headers("return=minimal"),
                                 params={"id": f"eq.{record_id}"}, timeout=30)
        if not result.ok:
            raise RuntimeError(f"Falha ao remover registro antigo {record_id}: {result.status_code} {result.text}")
        if index % 25 == 0 or index == len(remove):
            print(f"Removidos: {index}/{len(remove)}")
    print(f"Sincronização concluída: {len(rows)} cadastros oficiais.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
